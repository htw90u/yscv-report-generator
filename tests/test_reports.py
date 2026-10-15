"""塔塔加保育宣導回報單、生態廁所回報單：Word 填寫與畫面流程測試。執行：python -m unittest discover -s tests"""
import gc
import json
import os
import tempfile
import time
import tkinter as tk
from tkinter import ttk
import unittest
from tkinter import filedialog
from unittest import mock

from PIL import Image
from docx import Document
from docx.oxml.ns import qn

from _common import TATAKA_TEMPLATE, ECO_TEMPLATE
import yscv_report_generator as g
import yscv_reports as r

messagebox = g.messagebox


def all_text(doc):
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return '\n'.join(parts)


def cell_text(table, row, col):
    return r._cell(table, row, col).text


def checked_symbols(paragraph_el):
    """段落裡所有方框符號的狀態，依序：True＝打勾。"""
    return [sym.get(qn('w:char')) == r.CHECKED_SYMBOL for sym in paragraph_el.iter(qn('w:sym'))]


nl = chr(10)


TATAKA_DATA = dict(
    year='115', month='10', day='15', volunteer='王小明', time_text='09:00～16:00', weather='陰',
    guide_booked='1', guide_walkin='', center_on='', center_text='櫃檯諮詢', other_on='', other_text='',
    incident='有', incident_text='遊客扭傷腳踝', facility='● 木棧板鬆動\n● 衛生紙不足', visitors='指標不清楚',
    reflection='收穫很多', questions='1. 問：海拔多高？答：約2600公尺。')

ECO_DATA = dict(
    year='115', month='10', day='15', personnel='王小明', start_year='115', start_month='10', start_day='15',
    end_year='115', end_month='10', end_day='17', location='玉山主峰線', nature='● 玉山杜鵑', environment='帶下山垃圾 2 袋',
    reflection='狀況良好')
for _prefix, _ in r.TOILETS:
    for _spec, _value in zip(r.EXCRETA_FIELDS + r.COMPOST_FIELDS,
                             ['12.5', '晴', '有', '1.2', '3', '4', f'{_prefix}備註一', '正常', '55%', '18℃', '0.5', '是',
                              f'{_prefix}備註二']):
        ECO_DATA[f'{_prefix}_{_spec[0]}'] = _value


class ReportBuildTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def photos(self, count, size=(400, 200)):
        paths = []
        for i in range(count):
            p = os.path.join(self.tmp.name, f'p{i}.png')
            Image.new('RGB', size if i % 2 == 0 else size[::-1], 'red').save(p)
            paths.append(p)
        return paths

    def tataka(self, data=None, count=0, captions=None):
        out = os.path.join(self.tmp.name, 'tataka.docx')
        paths = self.photos(count)
        caps = {r.TATAKA_PHOTO_CATEGORY: captions or {p: f'名稱{i}\n地點{i}' for i, p in enumerate(paths)}}
        warnings = r.build_tataka_report(TATAKA_TEMPLATE, data or TATAKA_DATA, {r.TATAKA_PHOTO_CATEGORY: paths},
                                         caps, out)
        return Document(out), warnings

    def eco(self, data=None, count=0):
        out = os.path.join(self.tmp.name, 'eco.docx')
        paths = self.photos(count)
        caps = {r.ECO_PHOTO_CATEGORY: {p: f'照片說明{i}' for i, p in enumerate(paths)}}
        warnings = r.build_ecotoilet_report(ECO_TEMPLATE, data or ECO_DATA, {r.ECO_PHOTO_CATEGORY: paths}, caps, out)
        return Document(out), warnings

    # ---- 塔塔加
    def test_tataka_header_cells_and_weather_box(self):
        doc, warnings = self.tataka()
        self.assertEqual(warnings, [])
        table = doc.tables[0]
        self.assertEqual(cell_text(table, 1, 0), '王小明')
        self.assertEqual(cell_text(table, 1, 1).replace('\n', ''), '115年10月15日')
        self.assertEqual(cell_text(table, 1, 2), '09:00～16:00')
        weather = r._cell(table, 1, 3)._tc.findall(qn('w:p'))[0]
        self.assertEqual(checked_symbols(weather), [False, True, False])   # 晴 陰 雨：只有「陰」打勾
        self.assertEqual(r._cell(table, 1, 3).text, '晴陰雨')

    def test_split_names_detects_common_separators(self):
        expected = ['王小明', '李小華', '陳大文']
        for text in ('王小明 李小華 陳大文', '王小明、李小華、陳大文', '王小明,李小華,陳大文', '王小明，李小華；陳大文',
                     '王小明' + chr(0x3000) + '李小華/陳大文', '  王小明 ,  、李小華　　陳大文 ', '王小明' + chr(10) + '李小華' + chr(9) + '陳大文'):
            self.assertEqual(r.split_names(text), expected, repr(text))
        self.assertEqual(r.split_names('王小明'), ['王小明'])
        # 外文姓名：中間的空白不拆；中文姓名之間的空白照常拆
        self.assertEqual(r.split_names('John Smith'), ['John Smith'])
        self.assertEqual(r.split_names('John Smith, Mary Jane Watson'), ['John Smith', 'Mary Jane Watson'])
        self.assertEqual(r.split_names('王小明 John Smith 李小華'), ['王小明', 'John Smith', '李小華'])
        self.assertEqual(r.split_names('王小明 李小華 John Smith'), ['王小明', '李小華', 'John Smith'])
        self.assertEqual(r.split_names('John Smith、王小明' + chr(0x3000) + 'Mary Jane'), ['John Smith', '王小明', 'Mary Jane'])
        self.assertEqual(r.split_names('김 철수 田中 太郎'), ['김 철수', '田中', '太郎'])   # 韓文不拆，漢字姓名照拆
        self.assertEqual(r.split_names('約翰·史密斯 王小明'), ['約翰·史密斯', '王小明'])
        self.assertEqual(r.split_names('John  Smith' + chr(10) + 'Mary Jane'), ['John Smith', 'Mary Jane'])   # 換行一定分隔
        self.assertEqual(r.split_names('  ， 、 '), [])
        self.assertEqual(r.split_names(''), [])

    def test_tataka_each_volunteer_on_its_own_centered_line(self):
        doc, _ = self.tataka(dict(TATAKA_DATA, volunteer='王小明、李小華 陳大文,林小美'))
        cell = r._cell(doc.tables[0], 1, 0)
        paragraphs = cell._tc.findall(qn('w:p'))
        self.assertEqual([g._paragraph_text(p) for p in paragraphs], ['王小明', '李小華', '陳大文', '林小美'])
        for p in paragraphs:
            self.assertEqual(p.find(qn('w:pPr')).find(qn('w:jc')).get(qn('w:val')), 'center')
            self.assertEqual([c.tag.split('}')[1] for c in p.find(qn('w:pPr'))][-1], 'rPr')   # rPr 放最後（Word 的順序）
        height = doc.tables[0]._tbl.tr_lst[1].trPr.find(qn('w:trHeight'))
        self.assertNotEqual(height.get(qn('w:hRule')), 'exact')            # 人多時列會長高，不會被切掉

    def test_tataka_foreign_name_stays_on_one_line(self):
        doc, _ = self.tataka(dict(TATAKA_DATA, volunteer='王小明 John Smith 李小華'))
        paragraphs = r._cell(doc.tables[0], 1, 0)._tc.findall(qn('w:p'))
        self.assertEqual([g._paragraph_text(p) for p in paragraphs], ['王小明', 'John Smith', '李小華'])

    def test_tataka_service_items_and_incident(self):
        doc, _ = self.tataka()
        table = doc.tables[0]
        paragraphs = r._cell(table, 2, 1)._tc.findall(qn('w:p'))
        self.assertEqual(checked_symbols(paragraphs[0]), [True, False])    # 預約打勾、散客沒有
        self.assertEqual(checked_symbols(paragraphs[2]), [True])           # 有填遊客中心的文字，自動打勾
        self.assertIn('櫃檯諮詢', g._paragraph_text(paragraphs[2]))
        self.assertEqual(checked_symbols(paragraphs[3]), [False])          # 其他：沒填
        incident = r._cell(table, 3, 1)._tc.findall(qn('w:p'))[0]
        self.assertEqual(checked_symbols(incident), [False, True])         # 無／有：選「有」
        self.assertIn('遊客扭傷腳踝', g._paragraph_text(incident))

    def test_tataka_several_incidents_are_numbered_under_the_yes_box(self):
        records = nl.join(['遊客扭傷腳踝', '山屋停水', '有人違規露營'])
        doc, _ = self.tataka(dict(TATAKA_DATA, incident_text=records))
        paragraphs = r._cell(doc.tables[0], 3, 1)._tc.findall(qn('w:p'))
        self.assertEqual(len(paragraphs), 4)                              # 「有：」那行＋三筆
        self.assertEqual(checked_symbols(paragraphs[0]), [False, True])
        self.assertNotIn('遊客扭傷腳踝', g._paragraph_text(paragraphs[0]))
        texts = [g._paragraph_text(p).replace(' ', '').replace(chr(0x3000), '') for p in paragraphs[1:]]
        self.assertEqual(texts, ['1.遊客扭傷腳踝', '2.山屋停水', '3.有人違規露營'])
        height = doc.tables[0]._tbl.tr_lst[3].trPr.find(qn('w:trHeight'))
        self.assertNotEqual(height.get(qn('w:hRule')), 'exact')

    def test_tataka_incident_records_ignored_unless_yes(self):
        doc, _ = self.tataka(dict(TATAKA_DATA, incident='無', incident_text=nl.join(['甲', '乙'])))
        paragraphs = r._cell(doc.tables[0], 3, 1)._tc.findall(qn('w:p'))
        self.assertEqual(len(paragraphs), 1)
        self.assertNotIn('甲', g._paragraph_text(paragraphs[0]))

    def test_tataka_no_incident_checks_none_box(self):
        doc, _ = self.tataka(dict(TATAKA_DATA, incident='無', incident_text=''))
        incident = r._cell(doc.tables[0], 3, 1)._tc.findall(qn('w:p'))[0]
        self.assertEqual(checked_symbols(incident), [True, False])

    def test_tataka_text_rows_are_filled_and_not_clipped(self):
        doc, _ = self.tataka()
        table = doc.tables[0]
        self.assertIn('木棧板鬆動', cell_text(table, 4, 1))
        self.assertIn('衛生紙不足', cell_text(table, 4, 1))
        self.assertEqual(cell_text(table, 5, 1), '指標不清楚')
        self.assertEqual(cell_text(table, 6, 1), '收穫很多')
        for tr in table._tbl.tr_lst[1:5]:
            height = tr.trPr.find(qn('w:trHeight'))
            self.assertNotEqual(height.get(qn('w:hRule')), 'exact')        # 固定高度會切掉長文字

    def test_tataka_question_sample_line_removed(self):
        doc, _ = self.tataka()
        text = cell_text(doc.tables[0], 7, 1)
        self.assertNotIn('因為', text)
        self.assertIn('海拔多高', text)
        self.assertIn('將服勤期間相關疑問', text)             # 說明文字保留
        doc, _ = self.tataka(dict(TATAKA_DATA, questions=''))
        self.assertNotIn('因為', cell_text(doc.tables[0], 7, 1))   # 沒填也要拿掉範例

    def test_tataka_one_row_per_photo_with_name_and_place(self):
        doc, _ = self.tataka(count=3)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 9 + 3)                  # 範本 10 列：最後的空白列換成 3 張照片
        self.assertEqual(len(doc.inline_shapes), 3)
        self.assertEqual(cell_text(table, 9, 0), '名稱0\n地點0')
        self.assertEqual(cell_text(table, 11, 0), '名稱2\n地點2')
        for shape in doc.inline_shapes:                           # 照片放得進儲存格，不超過高度上限
            self.assertLessEqual(shape.height / 914400, r.PHOTO_MAX_HEIGHT_IN + 0.01)
            self.assertLess(shape.width / 914400, 4.5)

    def test_tataka_name_only_caption_and_no_photos(self):
        paths = self.photos(1)
        doc, _ = self.tataka(count=1, captions={paths[0]: '只有名稱'})
        self.assertEqual(cell_text(doc.tables[0], 9, 0), '只有名稱')
        doc, _ = self.tataka(count=0)
        self.assertEqual(len(doc.tables[0].rows), 10)             # 沒照片：維持範本的空白列
        self.assertEqual(len(doc.inline_shapes), 0)

    def test_tataka_bad_template_is_reported(self):
        out = os.path.join(self.tmp.name, 'bad.docx')
        with self.assertRaises(ValueError):
            r.build_tataka_report(ECO_TEMPLATE, TATAKA_DATA, {}, {}, out)

    # ---- 生態廁所
    def test_eco_header_and_dates(self):
        doc, warnings = self.eco()
        self.assertEqual(warnings, [])
        header = doc.tables[0]
        self.assertIn('115 年 10 月 15 日', cell_text(header, 1, 0))
        self.assertEqual(cell_text(header, 2, 1), '王小明')
        self.assertEqual(cell_text(header, 3, 1), '115年10月15日　至　115年10月17日')
        self.assertEqual(cell_text(header, 4, 1), '玉山主峰線')

    def test_eco_records_go_into_the_right_toilet_row(self):
        doc, _ = self.eco()
        excreta, compost = doc.tables[1], doc.tables[2]
        for row, prefix in ((3, 'm'), (7, 'b')):
            values = [c.text for c in excreta.rows[row].cells]
            self.assertEqual(values, ['12.5', '晴', '有', '1.2', '3', '4', f'{prefix}備註一'])
            values = [c.text for c in compost.rows[row].cells]
            self.assertEqual(values, ['正常', '55%', '18℃', '0.5', '是', f'{prefix}備註二'])
        self.assertEqual(cell_text(excreta, 2, 0), 'kg')         # 單位列維持範本原樣

    def test_eco_sections_fill_after_headings(self):
        doc, _ = self.eco()
        texts = [p.text.strip() for p in doc.paragraphs]
        for heading, content in (('二、自然資源狀況', '● 玉山杜鵑'), ('三、環境清理及遊客宣導情形', '帶下山垃圾 2 袋'),
                                 ('四、協勤心得與建議', '狀況良好')):
            index = next(i for i, t in enumerate(texts) if t.startswith(heading))
            following = [t for t in texts[index + 1:index + 4] if t]
            self.assertIn(content, following, heading)
        # 二、三的標題有兩行：內容放在第二行之後，不能插在標題中間
        nature = next(i for i, t in enumerate(texts) if t.startswith('二、自然資源狀況'))
        self.assertTrue(texts[nature + 1].endswith('等）'))
        self.assertEqual(texts[nature + 2], '● 玉山杜鵑')

    def test_eco_headings_stay_with_tables(self):
        doc, _ = self.eco()
        for p in doc.paragraphs:
            if p.text.strip().startswith(('（一）', '（二）', '（三）')):
                self.assertTrue(p.paragraph_format.keep_with_next, p.text)

    def test_eco_photo_grid_grows_and_shrinks(self):
        for count, rows in ((0, 2), (1, 2), (4, 4), (5, 6), (8, 8)):
            doc, _ = self.eco(count=count)
            grid = doc.tables[3]
            self.assertEqual(len(grid.rows), rows, count)
            self.assertEqual(len(doc.inline_shapes), count)
            self.assertNotIn('照片圖檔', all_text(doc))
            self.assertNotIn('照片文字', all_text(doc))
        doc, _ = self.eco(count=3)
        grid = doc.tables[3]
        self.assertEqual([cell_text(grid, 1, 0), cell_text(grid, 1, 1)], ['照片說明0', '照片說明1'])
        self.assertEqual([cell_text(grid, 3, 0), cell_text(grid, 3, 1)], ['照片說明2', ''])   # 單數張：最後一格空白

    def test_eco_photos_fit_the_cell(self):
        doc, _ = self.eco(count=4)
        for shape in doc.inline_shapes:
            self.assertLessEqual(shape.width / 914400, 3.4)
            self.assertLessEqual(shape.height / 914400, r.PHOTO_MAX_HEIGHT_IN + 0.01)

    def test_eco_bad_template_is_reported(self):
        out = os.path.join(self.tmp.name, 'bad.docx')
        with self.assertRaises(ValueError):
            r.build_ecotoilet_report(TATAKA_TEMPLATE, ECO_DATA, {}, {}, out)


class FormGuiTest(unittest.TestCase):
    app_class = None

    def setUp(self):
        if self.app_class is None:
            self.skipTest('基底類別')
        self.tmp = tempfile.TemporaryDirectory()
        self.root = tk.Tk()
        self.autosave = os.path.join(self.tmp.name, 'autosave.json')
        self.app = self.app_class(self.root, autosave_path=self.autosave, restore_on_start=False,
                                  settings_path=os.path.join(self.tmp.name, 'settings.json'))
        self.shown, self.pdf_calls, self.confirm_calls = [], [], []

        def fake_confirm(title, message, **kwargs):
            self.confirm_calls.append(message)
            return True

        def fake_pdf(docx, pdf):
            self.pdf_calls.append((docx, pdf))
            with open(pdf, 'wb') as f:
                f.write(b'%PDF-fake')

        for p in (mock.patch.object(messagebox, 'showerror', lambda t, m: self.shown.append(('err', m))),
                  mock.patch.object(messagebox, 'showinfo', lambda t, m: self.shown.append(('info', m))),
                  mock.patch.object(g, 'export_pdf', fake_pdf),
                  mock.patch.object(g, 'open_file', lambda path: None),
                  mock.patch.object(messagebox, 'askyesno', fake_confirm)):
            p.start()
            self.addCleanup(p.stop)

    def tearDown(self):
        if self.app_class is None:
            return
        self.root.destroy()
        self.app = self.root = None
        gc.collect()
        self.tmp.cleanup()

    def make_photos(self, names=('a.jpg', 'b.jpg')):
        category = self.app.PHOTO_CATEGORIES[0]
        paths = []
        for n in names:
            p = os.path.join(self.tmp.name, n)
            Image.new('RGB', (400, 200), 'red').save(p)
            paths.append(p)
        self.app._add_photo_files(category, paths)
        return category, self.app.listboxes[category], paths

    def wait_done(self, timeout=20):
        t0 = time.time()
        while time.time() - t0 < timeout and str(self.app.generate_btn['state']) == 'disabled':
            self.root.update()
            time.sleep(0.05)
        self.root.update()

    def generate(self):
        out = os.path.join(self.tmp.name, 'out.docx')
        captured = {}

        def ask(**kw):
            captured.update(kw)
            return out

        with mock.patch.object(filedialog, 'asksaveasfilename', ask):
            self.app.generate_report()
            self.wait_done()
        return out, captured

    def fill(self, name, value):
        g.widget_set(getattr(self.app, name), value)

    def set_text(self, key, value):
        self.app.text_fields[key].delete('1.0', tk.END)
        self.app.text_fields[key].insert('1.0', value)


class TatakaGuiTest(FormGuiTest):
    app_class = r.TatakaReportGUI

    def fill_all(self):
        values = dict(year='115', month='10', day='15', volunteer='王小明', time_text='09:00～16:00', weather='晴',
                      guide_booked='1', center_text='櫃檯諮詢', incident='無')
        for name, value in values.items():
            self.fill(name, value)
        for key in ('facility', 'visitors', 'reflection', 'questions'):
            self.set_text(key, f'{key}內容')

    def test_window_title_names_the_report(self):
        self.assertEqual(self.root.title(), f'{g.APP_TITLE}－塔塔加保育宣導回報單　{g.version_label()}')

    def test_tabs_and_default_date(self):
        titles = [self.app.notebook.tab(t, 'text').replace(' ' + g.MARK_MISSING, '') for t in self.app.notebook.tabs()]
        self.assertEqual(titles, ['基本資料', '服務項目', '設施與遊客意見', '心得與疑問解答', '相片記錄'])
        self.assertEqual(self.app.year.get(), g.roc_date_parts()[0])      # 值勤日期預設今天

    def test_blank_fields_and_tab_marks(self):
        blanks = self.app.find_blank_fields({})
        self.assertEqual(list(blanks), ['基本資料', '服務項目', '設施與遊客意見', '心得與疑問解答', '相片記錄'])
        self.assertIn('志工姓名', blanks['基本資料'])
        self.assertIn('服務項目（至少選一項）', blanks['服務項目'])
        self.fill_all()
        self.assertEqual(self.app.find_blank_fields({}), {'相片記錄': ['相片記錄（尚未加入照片）']})
        self.make_photos(('a.jpg',))
        self.assertEqual(self.app.find_blank_fields({}), {})
        self.assertIn('相片記錄', self.app._tabs_with_blanks())          # 照片還沒填資料：分頁仍標 ●
        self.fill('incident', '有')
        self.assertIn('偶發事件的說明', self.app.find_blank_fields({})['服務項目'])

    def test_two_caption_fields_are_joined_and_split(self):
        category, table, (a, b) = self.make_photos()
        table.set_selection([0])
        self.app._on_photo_select(category, table)
        self.app.caption_vars[0].set(' 塔塔加遊客中心 ')
        self.app.caption_vars[1].set('入口大廳')
        self.app._commit_caption()
        self.assertEqual(self.app.photo_captions[category][a], '塔塔加遊客中心\n入口大廳')
        self.assertEqual(table.item('0', 'values')[1], '塔塔加遊客中心／入口大廳')
        table.set_selection([1])
        self.app._on_photo_select(category, table)
        self.assertEqual([v.get() for v in self.app.caption_vars], ['', ''])
        table.set_selection([0])
        self.app._on_photo_select(category, table)
        self.assertEqual([v.get() for v in self.app.caption_vars], ['塔塔加遊客中心', '入口大廳'])
        self.app.caption_vars[0].set('')
        self.app.caption_vars[1].set('')
        self.app._commit_caption()
        self.assertNotIn(a, self.app.photo_captions[category])         # 兩欄都空＝沒填

    def test_enter_moves_to_next_field_then_next_photo(self):
        category, table, (a, b) = self.make_photos()
        self.app._select_photo(category, 0)
        self.app.caption_vars[0].set('甲')
        self.assertEqual(self.app._caption_return(None, 0), 'break')
        self.assertEqual(self.app.photo_captions[category].get(a), None)   # 還在第一張，還沒儲存
        self.app.caption_vars[1].set('地點')
        self.app._caption_return(None, 1)
        self.assertEqual(self.app.photo_captions[category][a], '甲\n地點')
        self.assertEqual(list(table.selected_indexes()), [1])           # 跳到下一張

    def test_collect_data_validates_date(self):
        self.fill_all()
        self.fill('month', '2')
        self.fill('day', '30')
        data, error = self.app.collect_data()
        self.assertIsNone(data)
        self.assertIn('值勤日期', error)

    def test_draft_roundtrip_keeps_checkboxes_and_choices(self):
        self.fill_all()
        self.fill('guide_walkin', '1')
        category, table, (a, b) = self.make_photos()
        self.app.photo_captions[category][a] = '名稱\n地點'
        draft = json.loads(json.dumps(self.app.build_draft()))
        self.assertEqual(draft['report'], 'tataka')
        self.assertEqual(draft['fields']['guide_booked'], '1')
        self.app._clear_widgets()
        self.assertEqual(self.app.guide_booked.get(), '')
        self.assertEqual(self.app.weather.get(), '')
        self.assertEqual(self.app.apply_draft(draft), [])
        self.assertEqual((self.app.guide_booked.get(), self.app.guide_walkin.get(), self.app.weather.get(),
                          self.app.incident.get()), ('1', '1', '晴', '無'))
        self.assertEqual(self.app.photo_captions[category][a], '名稱\n地點')
        self.assertEqual(self.app.text_fields['facility'].get('1.0', tk.END).strip(), 'facility內容')

    def test_plus_button_adds_incident_records(self):
        box = self.app.incident_text
        self.assertEqual(len(box.rows), 1)
        self.assertTrue(box.rows[0][2].instate(['disabled']))             # 只有一筆時不能刪
        box.add_button.invoke()
        box.add_button.invoke()
        self.assertEqual(len(box.rows), 3)
        self.assertFalse(box.rows[0][2].instate(['disabled']))
        for (frame, entry, button), text in zip(box.rows, ('甲', '', '丙')):
            entry.insert(0, text)
        self.assertEqual(box.get(), '甲' + nl + '丙')                       # 空白筆不算
        before = self.app._signature()
        box.rows[1][2].invoke()                                           # 刪掉中間（空白）那筆
        self.assertEqual(len(box.rows), 2)
        box.rows[0][2].invoke()
        self.assertEqual(box.get(), '丙')
        self.assertNotEqual(self.app._signature(), before)
        self.assertTrue(box.rows[0][2].instate(['disabled']))             # 只剩一筆：刪除鈕停用
        box.remove_row(box.rows[0])                                       # 硬要刪就只是清空文字
        self.assertEqual((len(box.rows), box.get()), (1, ''))

    def test_incident_records_survive_draft_and_mark_the_tab(self):
        self.fill('incident', '有')
        self.assertIn('偶發事件的說明', self.app.find_blank_fields({})['服務項目'])
        self.fill('incident_text', '甲' + nl + '乙')
        self.assertEqual(len(self.app.incident_text.rows), 2)
        self.assertNotIn('偶發事件的說明', self.app.find_blank_fields({}).get('服務項目', []))
        draft = json.loads(json.dumps(self.app.build_draft()))
        self.assertEqual(draft['fields']['incident_text'], '甲' + nl + '乙')
        self.app._clear_widgets()
        self.assertEqual((len(self.app.incident_text.rows), self.app.incident_text.get()), (1, ''))
        self.app.apply_draft(draft)
        self.assertEqual([e.get() for _, e, _ in self.app.incident_text.rows], ['甲', '乙'])
        data, _ = self.app.collect_data()
        self.assertEqual(data['incident_text'], '甲' + nl + '乙')

    def test_enter_in_incident_entry_adds_next_record(self):
        box = self.app.incident_text
        box.rows[0][1].insert(0, '甲')
        self.app.notebook.select(self.app._tab_by_title(self.app.TAB_SERVICE))
        self.root.deiconify()
        self.root.update()
        box.rows[0][1].focus_force()
        self.root.update()
        box.rows[0][1].event_generate('<Return>')
        self.root.update()
        self.assertEqual(len(box.rows), 2)

    def test_checkbox_change_marks_the_tab_and_dirty(self):
        before = self.app._signature()
        self.assertFalse(self.app.is_dirty())
        self.app.guide_booked.invoke()
        self.assertEqual(self.app.guide_booked.get(), '1')
        self.assertNotEqual(self.app._signature(), before)
        self.assertTrue(self.app.is_dirty())

    def test_draft_of_another_report_is_rejected(self):
        patrol_draft = {'version': g.DRAFT_VERSION, 'fields': {}}              # 舊草稿沒有種類欄位＝保育巡查回報單
        with self.assertRaises(ValueError) as caught:
            self.app.apply_draft(patrol_draft)
        self.assertIn('保育巡查回報單', str(caught.exception))
        with self.assertRaises(ValueError):
            self.app.apply_draft(dict(patrol_draft, report='ecotoilet'))
        self.app.apply_draft(dict(patrol_draft, report='tataka'))

    def test_file_names(self):
        self.fill_all()
        data, _ = self.app.collect_data()
        self.assertEqual(self.app.output_name(data), '1151015_塔塔加保育宣導回報單_王小明')
        data['volunteer'] = '王小明 李小華,John Smith'                       # 多位：只取第一位加「等」
        self.assertEqual(self.app.output_name(data), '1151015_塔塔加保育宣導回報單_王小明等')
        data['volunteer'] = 'John Smith'
        self.assertEqual(self.app.output_name(data), '1151015_塔塔加保育宣導回報單_John Smith')
        data['volunteer'] = '王小明、'                                         # 只有一位（多餘的分隔符號不算）
        self.assertEqual(self.app.output_name(data), '1151015_塔塔加保育宣導回報單_王小明')
        name = self.app.draft_name({'year': '115', 'month': '10', 'day': '15', 'volunteer': '王小:明'},
                                   g.datetime(2026, 10, 15, 9, 30))
        self.assertEqual(name, '1151015_塔塔加保育宣導回報單草稿_王小明_1015_0930')
        self.assertEqual(self.app.draft_name({}, g.datetime(2026, 10, 15, 9, 30)),
                         '塔塔加保育宣導回報單草稿_1015_0930')

    def test_generate_report_end_to_end(self):
        self.fill_all()
        category, table, (a, b) = self.make_photos()
        self.app.photo_captions[category] = {a: '甲\n地點甲', b: '乙'}
        out, captured = self.generate()
        self.assertEqual(captured['initialfile'], '1151015_塔塔加保育宣導回報單_王小明.docx')
        self.assertEqual(self.confirm_calls, [])                        # 都填了：不會問「有欄位未填」
        self.assertEqual(self.shown, [])
        self.assertEqual(self.pdf_calls, [(out, out[:-5] + '.pdf')])
        doc = Document(out)
        self.assertEqual(len(doc.inline_shapes), 2)
        self.assertEqual(cell_text(doc.tables[0], 9, 0), '甲\n地點甲')

    def test_generate_asks_when_something_is_blank(self):
        self.fill('volunteer', '王小明')
        self.generate()
        self.assertEqual(len(self.confirm_calls), 1)
        self.assertIn('服務項目', self.confirm_calls[0])
        self.assertNotIn('志工姓名', self.confirm_calls[0])      # 姓名已經填了


class EcoToiletGuiTest(FormGuiTest):
    app_class = r.EcoToiletReportGUI

    def fill_all(self):
        for name, value in dict(year='115', month='10', day='15', personnel='王小明', start_year='115',
                                start_month='10', start_day='15', end_year='115', end_month='10', end_day='17',
                                location='玉山主峰線').items():
            self.fill(name, value)
        for key in r.TOILET_FIELD_NAMES:
            self.fill(key, ECO_DATA[key])
        for key in ('nature', 'environment', 'reflection'):
            self.set_text(key, f'{key}內容')

    def test_window_title_and_tabs(self):
        self.assertEqual(self.root.title(), f'{g.APP_TITLE}－生態廁所回報單　{g.version_label()}')
        titles = [self.app.notebook.tab(t, 'text').replace(' ' + g.MARK_MISSING, '') for t in self.app.notebook.tabs()]
        self.assertEqual(titles, ['基本資料', '孟祿亭', '白木林', '自然資源/環境清理', '協勤心得與建議', '監測及維護照片'])

    @unittest.skipUnless(os.path.exists(g.EMOJI_FONT), '需要 Windows 的彩色表情符號字型')
    def test_toilet_tabs_have_a_toilet_icon(self):
        nb = self.app.notebook
        for _, name in r.TOILETS:
            self.assertTrue(str(nb.tab(self.app._tab_by_title(name), 'image')), name)
        for title in ('基本資料', self.app.TAB_NATURE, self.app.TAB_REFLECT, self.app.TAB_PHOTO):
            self.assertEqual(str(nb.tab(self.app._tab_by_title(title), 'image')), '', title)
        # 標籤文字（含未填的 ●）不受圖示影響，仍能用名稱找到分頁
        self.assertEqual(nb.tab(self.app._tab_by_title('孟祿亭'), 'text'), '孟祿亭 ' + g.MARK_MISSING)

    def test_toilet_tabs_still_work_without_the_emoji_font(self):
        self.root.destroy()
        self.root = tk.Tk()
        with mock.patch.object(g, 'EMOJI_FONT', os.path.join(self.tmp.name, 'missing.ttf')):
            app = r.EcoToiletReportGUI(self.root, autosave_path=self.autosave, restore_on_start=False,
                                       settings_path=os.path.join(self.tmp.name, 's2.json'))
        self.assertEqual(str(app.notebook.tab(app._tab_by_title('白木林'), 'image')), '')
        self.assertEqual(app._tab_icons, [])

    def test_location_defaults_to_both_toilets_and_can_be_edited(self):
        self.assertEqual(self.app.location.get(), '孟祿亭、白木林生態廁所')
        self.assertEqual(self.app.find_blank_fields({}).get('基本資料', []).count('協勤地點'), 0)
        self.assertFalse(self.app.is_dirty())                                # 預設值不算使用者改過
        self.fill('location', '玉山主峰線')
        self.assertTrue(self.app.is_dirty())
        self.app._clear_widgets()
        self.assertEqual(self.app.location.get(), '')
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: True):
            self.app.clear_all()                                             # 清除全部：回到預設值
        self.assertEqual(self.app.location.get(), '孟祿亭、白木林生態廁所')
        self.assertFalse(self.app.is_dirty())

    def test_default_location_alone_does_not_count_as_content_to_autosave(self):
        self.assertFalse(self.app.autosave_now())                            # 只有預設值＝空白，不寫自動存檔
        self.fill('personnel', '王小明')
        self.assertTrue(self.app.autosave_now())

    def test_draft_keeps_a_changed_location(self):
        self.fill('location', '玉山主峰線')
        draft = json.loads(json.dumps(self.app.build_draft()))
        self.app._clear_widgets()
        self.app.apply_draft(draft)
        self.assertEqual(self.app.location.get(), '玉山主峰線')

    def type_start(self, year, month, day):
        for name, value in (('start_year', year), ('start_month', month), ('start_day', day)):
            self.fill(name, value)
        self.app._sync_end_date()

    def end_date(self):
        return tuple(getattr(self.app, f'end_{p}').get() for p in ('year', 'month', 'day'))

    def test_end_date_follows_start_date_by_one_day(self):
        self.type_start('115', '10', '15')
        self.assertEqual(self.end_date(), ('115', '10', '16'))
        self.type_start('115', '10', '31')                                   # 跨月
        self.assertEqual(self.end_date(), ('115', '11', '1'))
        self.type_start('115', '12', '31')                                   # 跨年（民國年）
        self.assertEqual(self.end_date(), ('116', '1', '1'))
        self.type_start('113', '2', '28')                                    # 閏年（民國 113 年 = 2024）
        self.assertEqual(self.end_date(), ('113', '2', '29'))
        self.type_start('2026', '10', '15')                                  # 輸入西元年就維持西元年
        self.assertEqual(self.end_date(), ('2026', '10', '16'))

    def test_end_date_waits_for_a_complete_valid_start_date(self):
        self.fill('start_year', '115')
        self.fill('start_month', '2')
        self.app._sync_end_date()
        self.assertEqual(self.end_date(), ('', '', ''))                       # 日還沒填
        self.type_start('115', '2', '30')                                    # 2/30 不是有效日期
        self.assertEqual(self.end_date(), ('', '', ''))

    def test_end_date_tracks_edits_until_the_user_changes_it(self):
        self.type_start('115', '10', '1')                                    # 輸入「1」的當下就帶入
        self.assertEqual(self.end_date(), ('115', '10', '2'))
        self.type_start('115', '10', '15')                                   # 接著打成 15：迄日跟著更新
        self.assertEqual(self.end_date(), ('115', '10', '16'))
        self.fill('end_day', '20')                                           # 使用者自己改了迄日
        self.type_start('115', '10', '17')
        self.assertEqual(self.end_date(), ('115', '10', '20'))               # 不再蓋掉
        self.app._clear_widgets()                                            # 清除後重新開始自動帶入
        self.type_start('115', '10', '3')
        self.assertEqual(self.end_date(), ('115', '10', '4'))

    def test_typing_in_start_entries_triggers_the_end_date(self):
        self.app.notebook.select(self.app._tab_by_title('基本資料'))
        self.root.deiconify()
        self.root.update()
        for entry, text in ((self.app.start_year, '115'), (self.app.start_month, '10'), (self.app.start_day, '15')):
            entry.focus_force()
            self.root.update()
            entry.insert(0, text)
            entry.event_generate('<KeyRelease>')
            self.root.update()
        self.assertEqual(self.end_date(), ('115', '10', '16'))
        self.assertIn('已自動帶入迄日', self.app.status_label.cget('text'))

    def test_end_date_is_not_changed_by_loading_a_draft(self):
        self.fill_all()
        draft = json.loads(json.dumps(self.app.build_draft()))               # 起 10/15、迄 10/17
        self.app._clear_widgets()
        self.app.apply_draft(draft)
        self.assertEqual(self.end_date(), ('115', '10', '17'))
        self.app._sync_end_date()                                            # 使用者沒動過起日，也不會蓋掉草稿裡的迄日
        self.assertEqual(self.end_date(), ('115', '10', '17'))

    def test_file_names_carry_the_first_name_plus_deng(self):
        data = dict(start_year='115', start_month='10', start_day='15', end_year='115', end_month='10', end_day='16',
                    location='玉山主峰線', personnel='王小明 李小華')
        now = g.datetime(2026, 10, 15, 9, 30)
        self.assertEqual(self.app.output_name(data), '1151015-16_玉山主峰線生態廁所回報單_王小明等')
        self.assertEqual(self.app.draft_name(data, now), '1151015-16_玉山主峰線生態廁所回報單草稿_王小明等_1015_0930')
        data['personnel'] = '王小明'
        self.assertEqual(self.app.output_name(data), '1151015-16_玉山主峰線生態廁所回報單_王小明')

    def test_file_names_do_not_repeat_the_toilet_suffix(self):
        data = dict(start_year='115', start_month='10', start_day='15', end_year='115', end_month='10', end_day='16',
                    location='孟祿亭、白木林生態廁所')
        self.assertEqual(self.app.output_name(data), '1151015-16_孟祿亭、白木林生態廁所回報單')
        self.assertEqual(self.app.draft_name(data, g.datetime(2026, 10, 15, 9, 30)),
                         '1151015-16_孟祿亭、白木林生態廁所回報單草稿_1015_0930')

    def test_blank_fields_are_grouped_by_toilet_and_remarks_are_optional(self):
        blanks = self.app.find_blank_fields({})
        self.assertIn('本次收集桶總量（kg）', blanks['孟祿亭'])
        self.assertIn('本次收集桶總量（kg）', blanks['白木林'])
        self.assertNotIn('備註', blanks['孟祿亭'])
        self.fill_all()
        self.make_photos(('a.jpg',))
        self.app.photo_captions[self.app.PHOTO_CATEGORIES[0]] = {self.app.photo_files[self.app.PHOTO_CATEGORIES[0]][0]: 'x'}
        self.assertEqual(self.app.find_blank_fields({}), {})
        self.fill('b_temp', '')
        self.assertEqual(self.app.find_blank_fields({}), {'白木林': ['堆積溫度（範本：常溫）']})
        self.assertIn('白木林', self.app._tabs_with_blanks())

    def test_collect_data_includes_toilet_fields_and_validates_dates(self):
        self.fill_all()
        data, error = self.app.collect_data()
        self.assertIsNone(error)
        self.assertEqual(data['m_total'], '12.5')
        self.assertEqual(data['b_remark2'], 'b備註二')
        self.assertEqual(data['nature'], 'nature內容')
        self.fill('end_day', '1')
        data, error = self.app.collect_data()
        self.assertIn('起日', error)

    def test_draft_roundtrip_and_report_kind(self):
        self.fill_all()
        draft = json.loads(json.dumps(self.app.build_draft()))
        self.assertEqual(draft['report'], 'ecotoilet')
        self.app._clear_widgets()
        self.assertEqual(self.app.m_weather.get(), '')
        self.app.apply_draft(draft)
        self.assertEqual((self.app.m_weather.get(), self.app.b_turned.get(), self.app.m_total.get()),
                         ('晴', '是', '12.5'))
        with self.assertRaises(ValueError):
            self.app.apply_draft(dict(draft, report='tataka'))

    def test_generate_report_end_to_end(self):
        self.fill_all()
        category, table, (a, b, c) = self.make_photos(('a.jpg', 'b.jpg', 'c.jpg'))
        self.app.photo_captions[category] = {a: '甲', b: '乙', c: '丙'}
        out, captured = self.generate()
        self.assertEqual(captured['initialfile'], '1151015-17_玉山主峰線生態廁所回報單_王小明.docx')
        self.assertEqual(self.confirm_calls, [])
        self.assertEqual(self.shown, [])
        doc = Document(out)
        self.assertEqual(len(doc.inline_shapes), 3)
        self.assertEqual(cell_text(doc.tables[1], 3, 0), '12.5')
        self.assertEqual(cell_text(doc.tables[3], 1, 1), '乙')

    def test_draft_file_name(self):
        name = self.app.draft_name({'start_year': '115', 'start_month': '10', 'start_day': '15', 'end_year': '115',
                                    'end_month': '10', 'end_day': '15', 'location': '玉山主峰線'},
                                   g.datetime(2026, 10, 15, 9, 30))
        self.assertEqual(name, '1151015_玉山主峰線生態廁所回報單草稿_1015_0930')


class DebugLogTest(unittest.TestCase):
    """記錄檔只有在除錯模式（-d、-D、--debug）才會產生。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.log_path = os.path.join(self.tmp.name, 'debug.log')
        root = g.logging.getLogger()
        saved = (list(root.handlers), root.level)

        def restore():
            for handler in list(root.handlers):
                if handler not in saved[0]:
                    handler.close()
            root.handlers[:] = saved[0]
            root.setLevel(saved[1])

        self.addCleanup(restore)

    def test_parse_args(self):
        for flag in ('-d', '-D', '--debug'):
            self.assertTrue(g.parse_args([flag]).debug, flag)
        self.assertFalse(g.parse_args([]).debug)
        args = g.parse_args(['tataka', '-D'])
        self.assertEqual((args.report, args.debug), ('tataka', True))
        args = g.parse_args(['--debug', 'ecotoilet'])
        self.assertEqual((args.report, args.debug), ('ecotoilet', True))
        self.assertIsNone(g.parse_args([]).report)
        with mock.patch('sys.stderr', new=__import__('io').StringIO()), self.assertRaises(SystemExit):
            g.parse_args(['不存在的回報單'])

    def test_no_log_file_without_debug(self):
        with mock.patch.object(g, 'LOG_FILE', self.log_path):
            g.setup_logging(False)
            g.logging.info('平常不寫記錄')
            g.logging.error('錯誤只印在命令列')
        self.assertFalse(os.path.exists(self.log_path))
        root = g.logging.getLogger()
        self.assertFalse(any(isinstance(h, g.logging.FileHandler) for h in root.handlers))
        self.assertEqual(root.level, g.logging.WARNING)

    def test_log_file_is_written_in_debug_mode(self):
        with mock.patch.object(g, 'LOG_FILE', self.log_path):
            g.setup_logging(True)
            g.logging.info('除錯時要寫記錄')
            for handler in g.logging.getLogger().handlers:
                handler.flush()
        self.assertTrue(os.path.exists(self.log_path))
        text = open(self.log_path, encoding='utf-8').read()
        self.assertIn('除錯時要寫記錄', text)
        self.assertIn('除錯模式', text)

    def test_importing_the_program_does_not_create_a_log_file(self):
        import subprocess
        import sys
        copy = os.path.join(self.tmp.name, 'copy')
        os.makedirs(copy)
        for name in ('yscv_report_generator.py',):
            with open(os.path.join(g.BASE_DIR, name), 'rb') as src, open(os.path.join(copy, name), 'wb') as dst:
                dst.write(src.read())
        subprocess.run([sys.executable, '-c', 'import yscv_report_generator as g; g.logging.info("x"); g.logging.warning("w")'],
                       cwd=copy, check=True, capture_output=True)
        self.assertEqual([n for n in os.listdir(copy) if n.endswith('.log')], [])


class FakeFeedbackServer:
    """本機假的 Apps Script 端點（只在 127.0.0.1，不連外網）：記下收到的 JSON，依 mode 回應。
    mode：ok、fail（ok=false）、redirect（先 302 再回結果，和真的 Apps Script 一樣）、http500、garbage。"""

    def __init__(self, mode='ok'):
        import http.server
        import threading
        outer = self
        self.mode = mode
        self.received = []
        self.methods = []

        class Handler(http.server.BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def _reply(self, code, payload, ctype='application/json'):
                data = payload if isinstance(payload, bytes) else json.dumps(payload).encode('utf-8')
                self.send_response(code)
                self.send_header('Content-Type', ctype)
                self.send_header('Content-Length', str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            def do_POST(self):
                outer.methods.append('POST')
                length = int(self.headers.get('Content-Length', 0))
                outer.received.append(json.loads(self.rfile.read(length).decode('utf-8')))
                if outer.mode == 'redirect':
                    self.send_response(302)
                    self.send_header('Location', f'http://127.0.0.1:{outer.port}/result')
                    self.send_header('Content-Length', '0')
                    self.end_headers()
                elif outer.mode == 'fail':
                    self._reply(200, {'ok': False, 'error': 'Service invoked too many times'})
                elif outer.mode == 'http500':
                    self._reply(500, {'ok': False})
                elif outer.mode == 'garbage':
                    self._reply(200, '<html>不是 JSON</html>'.encode('utf-8'), 'text/html')
                else:
                    self._reply(200, {'ok': True})

            def do_GET(self):
                outer.methods.append('GET')
                self._reply(200, {'ok': True})

        self.server = http.server.HTTPServer(('127.0.0.1', 0), Handler)
        self.port = self.server.server_address[1]
        self.url = f'http://127.0.0.1:{self.port}/exec'
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()

    def close(self):
        self.server.shutdown()
        self.server.server_close()


class FeedbackSendTest(unittest.TestCase):
    """send_feedback：POST 到端點（不管使用者的郵件程式），各種結果都要有清楚的回應。"""

    def server(self, mode='ok'):
        server = FakeFeedbackServer(mode)
        self.addCleanup(server.close)
        return server

    def test_posts_json_with_subject_body_and_app(self):
        server = self.server()
        ok, error = g.send_feedback('[v1] 標題', '內容\n第二行', url=server.url)
        self.assertEqual((ok, error), (True, ''))
        self.assertEqual(server.received, [{'subject': '[v1] 標題', 'body': '內容\n第二行', 'app': g.APP_TITLE,
                                            'reply_to': ''}])
        g.send_feedback('s', 'b', url=server.url, reply_to='  me@example.com ')
        self.assertEqual(server.received[-1]['reply_to'], 'me@example.com')       # 回信地址，端點用來設定 Reply-To

    def test_follows_the_redirect_like_apps_script(self):
        server = self.server('redirect')
        self.assertEqual(g.send_feedback('s', 'b', url=server.url), (True, ''))
        self.assertEqual(server.methods, ['POST', 'GET'])                         # 302 之後用 GET 取結果

    def test_reports_the_error_from_the_endpoint(self):
        ok, error = g.send_feedback('s', 'b', url=self.server('fail').url)
        self.assertFalse(ok)
        self.assertIn('too many times', error)

    def test_http_error_unreachable_and_bad_response(self):
        ok, error = g.send_feedback('s', 'b', url=self.server('http500').url)
        self.assertEqual((ok, error), (False, '伺服器回應錯誤（500）'))
        ok, error = g.send_feedback('s', 'b', url=self.server('garbage').url)
        self.assertEqual((ok, error), (False, '伺服器的回應格式不正確'))
        closed = self.server()
        closed.close()                                                            # 連不上
        ok, error = g.send_feedback('s', 'b', url=closed.url, timeout=3)
        self.assertFalse(ok)
        self.assertIn('無法連線', error)

    def test_no_endpoint_configured(self):
        with mock.patch.object(g, 'feedback_endpoint_url', return_value=''):
            ok, error = g.send_feedback('s', 'b')
        self.assertFalse(ok)
        self.assertIn('沒有設定', error)

    def test_endpoint_url_comes_from_the_ignored_config_module(self):
        import sys
        import types
        fake = types.ModuleType('yscv_feedback_endpoint')
        fake.URL = '  https://script.google.com/macros/s/ABC/exec  '
        with mock.patch.dict(sys.modules, {'yscv_feedback_endpoint': fake}):
            self.assertEqual(g.feedback_endpoint_url(), 'https://script.google.com/macros/s/ABC/exec')
        with mock.patch.dict(sys.modules, {'yscv_feedback_endpoint': None}):      # 沒有這個檔案
            self.assertEqual(g.feedback_endpoint_url(), '')

    def test_no_mail_client_or_recipient_is_involved(self):
        self.assertFalse(hasattr(g, 'open_mailto'))                               # 不再有 mailto 備援
        source = open(os.path.join(g.BASE_DIR, 'yscv_report_generator.py'), encoding='utf-8').read()
        self.assertNotIn('mailto', source)
        self.assertNotRegex(source, r'[\w.]+@[\w.]+\.(io|com|tw|org)')            # 程式碼裡沒有任何信箱


class FeedbackTest(unittest.TestCase):
    """入口選單的「意見反應」視窗（寄送一律用假的，不會真的送出）。"""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = tk.Tk()
        self.launcher = g.ReportLauncher(self.root, settings_path=os.path.join(self.tmp.name, 's.json'))
        self.sent, self.shown, self.reply_tos = [], [], []
        self.gate = None                                    # 讓假的寄送卡住，測試「寄送中」的畫面
        self.result = (True, '')

        def fake_send(subject, body, reply_to=''):
            if self.gate:
                self.gate.wait(5)
            self.sent.append((subject, body))
            self.reply_tos.append(reply_to)
            return self.result

        for p in (mock.patch.object(g, 'send_feedback', fake_send),
                  mock.patch.object(messagebox, 'showinfo', lambda t, m: self.shown.append(('info', t, m))),
                  mock.patch.object(messagebox, 'showerror', lambda t, m: self.shown.append(('err', t, m)))):
            p.start()
            self.addCleanup(p.stop)
        self.addCleanup(lambda: messagebox.attach(None))
        self.addCleanup(gc.collect)
        self.addCleanup(self.root.destroy)

    def dialog(self, title='', content='', email=''):
        dialog = g.FeedbackDialog(self.root, self.launcher.palette, self.launcher.font_family, 'v9.9.3-abc1234')
        self.addCleanup(lambda: dialog.winfo_exists() and dialog.destroy())
        dialog.title_entry.insert(0, title)
        dialog.content_text.insert('1.0', content)
        dialog.email_entry.insert(0, email)
        return dialog

    def finish(self, dialog, timeout=5):
        t0 = time.time()
        while time.time() - t0 < timeout and dialog.winfo_exists() and dialog.sending:
            self.root.update()
            time.sleep(0.02)
        self.root.update()

    def test_feedback_button_is_at_the_bottom_right(self):
        self.root.update()
        button = self.launcher.feedback_button
        self.assertEqual(button.cget('text'), '意見反應')
        cards = list(self.launcher.buttons.values())
        self.assertGreater(button.winfo_rooty(), max(card.winfo_rooty() + card.winfo_height() - 1 for card in cards))
        self.assertGreaterEqual(button.winfo_rootx() + button.winfo_width(),
                                max(card.winfo_rootx() + card.winfo_width() for card in cards) - 12)   # 靠右對齊卡片的右緣
        self.assertGreater(button.winfo_rootx(), self.root.winfo_rootx() + self.root.winfo_width() // 2)

    def test_message_carries_the_version_in_subject_and_body(self):
        subject, body = g.feedback_message('希望新增功能', '請加上九宮格', 'v0.5.16-ca25953')
        self.assertEqual(subject, '[回報單生成器 v0.5.16-ca25953] 希望新增功能')
        self.assertTrue(body.startswith('請加上九宮格'))
        self.assertIn('v0.5.16-ca25953', body)
        self.assertIn('聯絡 Email：（未填）', body)                                    # 沒填聯絡 Email
        self.assertIn('聯絡 Email：me@example.com', g.feedback_message('t', 'c', 'v1', ' me@example.com ')[1])
        self.assertEqual(g.feedback_message('  ', '內容', 'v1')[0], '[回報單生成器 v1] 意見反應')        # 沒填標題用預設

    def test_dialog_does_not_show_a_recipient(self):
        dialog = self.dialog()

        def texts(widget):
            found = [str(widget.cget('text'))] if 'text' in widget.keys() else []
            for child in widget.winfo_children():
                found += texts(child)
            return found

        everything = '\n'.join(texts(dialog))
        self.assertNotIn('@', everything)
        self.assertNotIn('收件人', everything)
        self.assertNotIn('v9.9.3-abc1234', everything)                              # 版本自動加進信裡，不用告訴使用者
        self.assertNotIn('版本', everything)
        self.assertIn('聯絡 Email', everything)                                      # 使用者可以留下自己的聯絡信箱
        self.assertNotIn('選填', everything)                                         # 「選填」只是輸入框裡的灰色淡字
        self.assertEqual(ttk.Entry.get(dialog.email_entry), '(選填)')

    def test_sending_posts_in_the_background_and_closes_on_success(self):
        dialog = self.dialog('照片太小', '預覽可以再大一點\n謝謝')
        self.assertTrue(dialog.send())
        self.finish(dialog)
        self.assertEqual(len(self.sent), 1)
        subject, body = self.sent[0]
        self.assertEqual(subject, '[回報單生成器 v9.9.3-abc1234] 照片太小')
        self.assertIn('預覽可以再大一點\n謝謝', body)
        self.assertIn('v9.9.3-abc1234', body)
        self.assertEqual(self.shown[-1], ('info', '意見反應', '已寄出，謝謝你的意見！'))
        self.assertFalse(dialog.winfo_exists())                                   # 成功後關閉視窗

    def test_contact_email_is_optional_and_goes_into_the_mail_and_reply_to(self):
        dialog = self.dialog('要回覆', '請回信給我', email=' me@example.com ')
        self.assertTrue(dialog.send())
        self.finish(dialog)
        self.assertIn('聯絡 Email：me@example.com', self.sent[0][1])
        self.assertEqual(self.reply_tos, ['me@example.com'])
        dialog = self.dialog('不用回覆', '匿名意見')                                  # 留空也可以寄
        self.assertTrue(dialog.send())
        self.finish(dialog)
        self.assertIn('聯絡 Email：（未填）', self.sent[1][1])
        self.assertEqual(self.reply_tos[1], '')

    def test_email_field_shows_a_gray_hint_but_has_no_value(self):
        dialog = self.dialog()
        entry = dialog.email_entry
        self.root.deiconify()
        self.root.update()
        self.assertEqual(ttk.Entry.get(entry), '(選填)')                              # 畫面上看到的提示
        self.assertEqual(entry.get(), '')                                          # 但欄位沒有值
        self.assertEqual(str(entry.cget('foreground')), g.PlaceholderEntry.PLACEHOLDER_COLOR)     # 灰色淡字
        self.assertTrue(dialog.send() is False)                                    # 空內容擋下；提示文字不會被當成 Email
        dialog.content_text.insert('1.0', '內容')
        self.assertTrue(dialog.send())                                             # 沒填 Email 照樣能寄
        self.finish(dialog)
        self.assertEqual(self.reply_tos, [''])
        self.assertIn('聯絡 Email：（未填）', self.sent[0][1])                       # 信裡也不會出現「選填」

    def test_hint_disappears_on_focus_and_returns_when_left_empty(self):
        dialog = self.dialog()
        entry = dialog.email_entry
        self.root.deiconify()
        self.root.update()
        entry.focus_force()
        self.root.update()
        self.assertEqual(ttk.Entry.get(entry), '')                                 # 點進去提示消失
        self.assertNotEqual(str(entry.cget('foreground')), g.PlaceholderEntry.PLACEHOLDER_COLOR)
        entry.insert(0, 'me@example.com')
        self.assertEqual(entry.get(), 'me@example.com')
        dialog.title_entry.focus_force()
        self.root.update()
        self.assertEqual(ttk.Entry.get(entry), 'me@example.com')                   # 有填內容，離開後不會被提示蓋掉
        entry.delete(0, tk.END)
        self.assertEqual(ttk.Entry.get(entry), '(選填)')                              # 清空後（沒有在輸入）又顯示提示
        self.assertEqual(entry.get(), '')

    def test_invalid_contact_email_is_refused(self):
        for bad in ('abc', 'a@b', '@example.com', 'me@@example.com', 'me @example.com'):
            dialog = self.dialog('標題', '內容', email=bad)
            self.assertFalse(dialog.send(), bad)
            self.assertIn('Email', self.shown[-1][2])
            self.assertTrue(dialog.winfo_exists())
        self.assertEqual(self.sent, [])
        for good in ('me@example.com', 'first.last+tag@sub.example.co.uk'):
            self.assertTrue(g.EMAIL_PATTERN.match(good), good)

    def test_buttons_are_locked_while_sending(self):
        import threading
        self.gate = threading.Event()
        dialog = self.dialog('慢', '網路很慢')
        self.assertTrue(dialog.send())
        self.root.update()
        self.assertTrue(dialog.sending)
        self.assertTrue(dialog.send_button.instate(['disabled']))
        self.assertEqual(dialog.send_button.cget('text'), '寄送中…')
        self.assertTrue(dialog.cancel_button.instate(['disabled']))
        self.assertFalse(dialog.send())                                           # 重複按不會寄第二封
        dialog.cancel()                                                           # 寄送中取消無效
        self.assertTrue(dialog.winfo_exists())
        self.gate.set()
        self.finish(dialog)
        self.assertEqual(len(self.sent), 1)
        self.assertFalse(dialog.winfo_exists())

    def test_empty_content_is_refused(self):
        dialog = self.dialog('只有標題', '   ')
        self.assertFalse(dialog.send())
        self.assertEqual(self.sent, [])
        self.assertEqual(self.shown[-1], ('err', '意見反應', '請先填寫內容'))
        self.assertTrue(dialog.winfo_exists())

    def test_failure_keeps_the_text_and_allows_a_retry(self):
        self.result = (False, '無法連線到伺服器（timed out）')
        dialog = self.dialog('失敗', '網路不通')
        dialog.send()
        self.finish(dialog)
        kind, title, message = self.shown[-1]
        self.assertEqual((kind, title), ('err', '寄送失敗'))
        self.assertIn('無法連線到伺服器', message)
        self.assertTrue(dialog.winfo_exists())
        self.assertEqual(dialog.content_text.get('1.0', tk.END).strip(), '網路不通')          # 內容還在
        self.assertFalse(dialog.send_button.instate(['disabled']))
        self.result = (True, '')                                                  # 網路恢復後可以重寄
        self.assertTrue(dialog.send())
        self.finish(dialog)
        self.assertEqual(len(self.sent), 2)
        self.assertFalse(dialog.winfo_exists())

    def test_end_to_end_with_a_local_server(self):
        server = FakeFeedbackServer('redirect')
        self.addCleanup(server.close)
        mock.patch.stopall()                                                        # 這個測試用真的 send_feedback（只連本機）
        dialog = self.dialog('完整流程', '從畫面一路送到假伺服器', email='me@example.com')
        with mock.patch.object(g, 'feedback_endpoint_url', return_value=server.url), \
                mock.patch.object(messagebox, 'showinfo', lambda t, m: self.shown.append(('info', t, m))):
            dialog.send()
            self.finish(dialog, timeout=10)
        self.assertEqual(len(server.received), 1)
        self.assertEqual(server.received[0]['subject'], '[回報單生成器 v9.9.3-abc1234] 完整流程')
        self.assertIn('從畫面一路送到假伺服器', server.received[0]['body'])
        self.assertEqual(server.received[0]['reply_to'], 'me@example.com')
        self.assertEqual(self.shown[-1][0], 'info')

    def test_open_feedback_uses_the_current_version(self):
        with mock.patch.object(g.FeedbackDialog, 'run', lambda self: None):
            dialog = self.launcher.open_feedback()
        try:
            self.assertEqual(dialog.version, g.version_label())
        finally:
            dialog.destroy()

    def test_launcher_binds_the_themed_message_boxes(self):
        self.assertIs(messagebox.app, self.launcher)


class LauncherTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = tk.Tk()

    def tearDown(self):
        try:
            self.root.destroy()
        except tk.TclError:
            pass
        gc.collect()

    def launcher(self):
        return g.ReportLauncher(self.root, settings_path=os.path.join(self.tmp.name, 's.json'))

    def test_lists_the_three_reports_in_order(self):
        launcher = self.launcher()
        self.assertEqual(list(launcher.buttons), ['patrol', 'tataka', 'ecotoilet'])
        self.assertEqual([c[1] for c in g.REPORT_CHOICES], ['保育巡查回報單', '塔塔加保育宣導回報單', '生態廁所回報單'])
        self.assertIn(g.version_label(), self.root.title())
        self.assertTrue(self.root.title().startswith(g.APP_TITLE))

    def test_clicking_a_button_chooses_and_closes(self):
        launcher = self.launcher()
        launcher.buttons['tataka'].invoke()
        self.assertEqual(launcher.choice, 'tataka')
        with self.assertRaises(tk.TclError):                      # 選完就關掉選單，換成回報單視窗
            self.root.winfo_exists()

    def test_card_names_are_bold_and_descriptions_are_not(self):
        import tkinter.font as tkfont
        card = self.launcher().buttons['tataka']
        _, name, description = card.parts
        self.assertEqual(name.cget('text'), '塔塔加保育宣導回報單')
        self.assertEqual(tkfont.Font(root=self.root, font=name.cget('font')).actual('weight'), 'bold')
        self.assertEqual(tkfont.Font(root=self.root, font=description.cget('font')).actual('weight'), 'normal')
        size = lambda w: tkfont.Font(root=self.root, font=w.cget('font')).actual('size')
        self.assertGreater(size(name), size(description))              # 名稱比說明大
        self.assertNotEqual(name.cget('foreground'), description.cget('foreground'))

    @unittest.skipUnless(os.path.exists(g.EMOJI_FONT), '需要 Windows 的彩色表情符號字型')
    def test_each_card_has_a_color_icon(self):
        launcher = self.launcher()
        self.assertEqual(set(g.REPORT_ICONS), {'patrol', 'tataka', 'ecotoilet'})
        for report_id, card in launcher.buttons.items():
            self.assertIsNotNone(card.icon_label, report_id)
        icon = g.emoji_image(g.REPORT_ICONS['ecotoilet'], 44)
        self.assertEqual(max(icon.size), 44)
        colors = {pixel[:3] for pixel in icon.getdata() if pixel[3] > 200}
        self.assertGreater(len(colors), 3)                               # 是彩色圖，不是單色字形

    @unittest.skipUnless(os.path.exists(g.EMOJI_FONT), '需要 Windows 的彩色表情符號字型')
    def test_icons_are_not_clipped(self):
        # 廁所標誌是正方形：被切掉上緣的話，高度會明顯小於寬度
        icon = g.emoji_image(g.REPORT_ICONS['ecotoilet'], 100)
        self.assertLessEqual(abs(icon.width - icon.height), 4, icon.size)
        # 圖示四邊都有透明邊或完整的圖案邊界（不是貼著圖片邊緣被切斷）：用大字級畫出完整圖，高度應等於寬度
        for report_id, char in g.REPORT_ICONS.items():
            big = g.emoji_image(char, 200)
            self.assertEqual(max(big.size), 200, report_id)
            self.assertGreaterEqual(min(big.size), 140, report_id)       # 最扁的圖示（健行靴）也不會被切成一條

    @unittest.skipUnless(os.path.exists(g.EMOJI_FONT), '需要 Windows 的彩色表情符號字型')
    def test_window_icon_is_a_form_with_a_mountain(self):
        icon = g.report_icon(64)
        self.assertEqual(icon.size, (64, 64))
        pixels = [p for p in icon.getdata() if p[3] > 200]
        orange = sum(1 for r_, g_, b, _ in pixels if r_ > 200 and 130 < g_ < 190 and b < 110)       # 夾板的橘色
        green = sum(1 for r_, g_, b, _ in pixels if g_ > 180 and r_ < 60 and b < 130)                # 山下的綠地
        self.assertGreater(orange, 100)
        self.assertGreater(green, 30)
        corner = [icon.getpixel((x, y)) for x in range(54, 64) for y in range(54, 64)]
        self.assertTrue(any(p[3] > 200 for p in corner))                                              # 右下角有山

    def test_every_window_gets_the_icon_and_missing_font_falls_back(self):
        self.assertTrue(g.set_window_icon(self.root))
        self.assertEqual(len(self.root._window_icons), len(g.WINDOW_ICON_SIZES))
        self.assertTrue(hasattr(self.launcher(), 'buttons'))                 # 選單視窗也設定了（不會出錯）
        g._icon_cache.clear()
        try:
            with mock.patch.object(g, 'EMOJI_FONT', os.path.join(self.tmp.name, 'missing.ttf')):
                self.assertFalse(g.set_window_icon(self.root))               # 沒有字型：維持預設圖示，不丟例外
                self.assertEqual(self.root._window_icons, [])
        finally:
            g._icon_cache.clear()

    def test_launcher_still_works_without_the_emoji_font(self):
        with mock.patch.object(g, 'EMOJI_FONT', os.path.join(self.tmp.name, 'missing.ttf')):
            launcher = self.launcher()
        for card in launcher.buttons.values():
            self.assertIsNone(card.icon_label)
        launcher.buttons['tataka'].invoke()
        self.assertEqual(launcher.choice, 'tataka')
        self.assertIsNone(g.emoji_image('', 44))

    def test_enter_on_the_focused_card_chooses_it(self):
        launcher = self.launcher()
        self.root.update()
        launcher.buttons['patrol'].focus_force()
        self.root.update()
        launcher.buttons['patrol'].event_generate('<Return>')
        self.root.update()
        self.assertEqual(launcher.choice, 'patrol')

    def test_clicking_a_card_label_chooses_it(self):
        launcher = self.launcher()
        self.root.update()
        launcher.buttons['ecotoilet'].parts[1].event_generate('<Button-1>')
        self.root.update()
        self.assertEqual(launcher.choice, 'ecotoilet')

    def test_number_keys_choose(self):
        launcher = self.launcher()
        self.root.update()
        self.root.focus_force()
        self.root.update()
        self.root.event_generate('<Key-3>')
        self.root.update()
        self.assertEqual(launcher.choice, 'ecotoilet')

    def test_escape_closes_without_choice(self):
        launcher = self.launcher()
        self.root.update()
        self.root.event_generate('<Escape>')
        self.root.update()
        self.assertIsNone(launcher.choice)

    def test_create_app_builds_each_report(self):
        self.root.destroy()
        for report_id, cls_name in (('patrol', 'ReportGeneratorGUI'), ('tataka', 'TatakaReportGUI'),
                                    ('ecotoilet', 'EcoToiletReportGUI')):
            root = tk.Tk()
            try:
                app = g.create_app(report_id, root, autosave_path=os.path.join(self.tmp.name, 'a.json'),
                                   restore_on_start=False, settings_path=os.path.join(self.tmp.name, 's.json'))
                self.assertEqual(type(app).__name__, cls_name)
                self.assertEqual(app.REPORT_ID, report_id)
                self.assertTrue(os.path.exists(app.template_path), app.template_path)
            finally:
                root.destroy()
                gc.collect()
        self.root = tk.Tk()

    def test_each_report_has_its_own_autosave_file(self):
        names = {cls.DEFAULT_AUTOSAVE_PATH for cls in (g.ReportGeneratorGUI, r.TatakaReportGUI, r.EcoToiletReportGUI)}
        self.assertEqual(len(names), 3)


if __name__ == '__main__':
    unittest.main()
