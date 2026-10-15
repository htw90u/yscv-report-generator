"""塔塔加保育宣導回報單、生態廁所回報單：Word 填寫邏輯與畫面。

畫面繼承 yscv_report_generator.ReportGeneratorGUI（照片拖放／旋轉／裁切、草稿、自動存檔、PDF 轉檔都共用），
這裡只定義各自的欄位、分頁與填進範本的方式。
"""
import copy
import math
import os
import re
import tkinter as tk
from tkinter import ttk

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import _Cell

import yscv_report_generator as g

# ---------------------------------------------------------------- 共用：填 Word 的小工具

CHECKED_SYMBOL = 'F052'    # Wingdings 2：打勾的方框
UNCHECKED_SYMBOL = 'F0A3'  # Wingdings 2：空的方框（和範本原本的方框相同）
UNDERLINE_MIN_WIDTH = 12   # 填空線至少有幾個半形字寬
PHOTO_MAX_HEIGHT_IN = 3.0  # 塔塔加／生態廁所：單張照片高度上限（吋）
CheckField = g.CheckField


def roc_date_text(year, month, day):
    return f"{year}年{month}月{day}日"


def _cell(table, row_index, col_index):
    tr = table._tbl.tr_lst[row_index]
    return _Cell(tr.tc_lst[col_index], table)


def _clear_runs(paragraph_el):
    """清掉段落裡的文字（保留段落格式）。"""
    for child in list(paragraph_el):
        if child.tag != qn('w:pPr'):
            paragraph_el.remove(child)


def _rpr(size=28, underline=False):
    rpr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn(f'w:{attr}'), g.KAI)
    rpr.append(fonts)
    for tag in ('w:sz', 'w:szCs'):
        el = OxmlElement(tag)
        el.set(qn('w:val'), str(size))
        rpr.append(el)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rpr.append(u)
    return rpr


def _text_run(text, size=28, underline=False):
    run = OxmlElement('w:r')
    run.append(_rpr(size, underline))
    t = OxmlElement('w:t')
    t.text = text
    t.set(qn('xml:space'), 'preserve')
    run.append(t)
    return run


def _checkbox_run(checked, size=32):
    run = OxmlElement('w:r')
    run.append(_rpr(size))
    sym = OxmlElement('w:sym')
    sym.set(qn('w:font'), 'Wingdings 2')
    sym.set(qn('w:char'), CHECKED_SYMBOL if checked else UNCHECKED_SYMBOL)
    run.append(sym)
    return run


def _blank_line(text):
    """底線填空：有內容就把內容放在底線上（至少 UNDERLINE_MIN_WIDTH 個半形字寬），沒內容就是一條空白底線。"""
    width = sum(2 if ord(ch) > 255 else 1 for ch in text)
    return _text_run(f" {text}" + ' ' * max(1, UNDERLINE_MIN_WIDTH - width), underline=True)


def _rebuild(paragraph_el, *runs):
    _clear_runs(paragraph_el)
    for run in runs:
        paragraph_el.append(run)


def _relax_height(table, row_index):
    """範本有些列是「固定高度」，文字一多就被切掉；改成「最小高度」讓列自動長高。"""
    trpr = table._tbl.tr_lst[row_index].trPr
    height = trpr.find(qn('w:trHeight')) if trpr is not None else None
    if height is not None and height.get(qn('w:hRule')) == 'exact':
        height.set(qn('w:hRule'), 'atLeast')


def _content_paragraph(text, shift=480):
    """內容段落（標楷體 12pt、固定行高 22pt），縮排比 g._build_content_paragraph 少 shift（沒有標題文字的儲存格不需要縮排）。"""
    p = g._build_content_paragraph(text)
    ind = p.find(qn('w:pPr')).find(qn('w:ind'))
    left = max(0, int(ind.get(qn('w:left'))) - shift)
    ind.set(qn('w:left'), str(left))
    if ind.get(qn('w:hanging')):
        ind.set(qn('w:hanging'), str(min(int(ind.get(qn('w:hanging'))), left)))
    return p


def fill_cell_content(cell, text, shift=480):
    """把多行內容填進儲存格（取代原有的空白段落）。沒有內容就維持範本原樣。"""
    lines = g.content_lines(text)
    if not text or not text.strip():
        return
    paragraphs = cell._tc.findall(qn('w:p'))
    anchor = paragraphs[0]
    for extra in paragraphs[1:]:
        cell._tc.remove(extra)
    for line in lines:
        new = _content_paragraph(line, shift)
        anchor.addnext(new)
        anchor = new
    cell._tc.remove(paragraphs[0])


def insert_content_after(paragraph, text, shift=480, is_placeholder=None, trailing_blank=False):
    """在 paragraph 後面插入內容，並移除其後的空白段落（以及 is_placeholder 判定的範例文字）。沒有內容就不動。"""
    if not text or not text.strip():
        return
    anchor = paragraph._p
    nxt = anchor.getnext()
    while nxt is not None and nxt.tag == qn('w:p') and (
            g._is_empty_paragraph(nxt) or (is_placeholder and is_placeholder(g._paragraph_text(nxt)))):
        anchor.getparent().remove(nxt)
        nxt = anchor.getnext()
    lines = g.content_lines(text) + ([''] if trailing_blank else [])
    for line in lines:
        new = _content_paragraph(line, shift)
        anchor.addnext(new)
        anchor = new


def _picture_into_cell(cell, item, warnings=None):
    """儲存格放一張照片（置中）。照片載入失敗就放文字說明。"""
    paragraphs = cell._tc.findall(qn('w:p'))
    for extra in paragraphs[1:]:
        cell._tc.remove(extra)
    _clear_runs(paragraphs[0])
    paragraph = cell.paragraphs[0]
    g._tighten(paragraph, keep_with_next=True)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if item.source is not None:
        paragraph.add_run().add_picture(item.source, width=Inches(item.img_w), height=Inches(item.img_h))
        g.logging.info(f"✓ SUCCESS: 插入照片 {os.path.basename(item.path)}")
    else:
        g.set_font_kai(paragraph.add_run(f"[照片無法載入：{os.path.basename(item.path)}]"))


def _page_width_in(doc):
    section = doc.sections[-1]
    return (section.page_width - section.left_margin - section.right_margin) / 914400


def _grid_widths(table):
    return [int(c.get(qn('w:w'))) for c in table._tbl.tblGrid.findall(qn('w:gridCol'))]


# ---------------------------------------------------------------- 塔塔加保育宣導回報單

TATAKA_PHOTO_CATEGORY = '相片記錄'
TATAKA_WEATHERS = ('晴', '陰', '雨')
NEWLINE = chr(10)
ECO_DEFAULT_LOCATION = '孟祿亭、白木林生態廁所'
DATE_PARTS = ('year', 'month', 'day')
TAB_ICON_PX = 22        # 分頁標籤上的圖示大小（96 dpi 時的像素）
split_names = g.split_names




def _tataka_service_paragraphs(data):
    """服務項目、偶發事件各段的內容（方框打勾與填空）。"""
    guide = (_text_run('帶團解說：'), _checkbox_run(bool(data['guide_booked'])), _text_run('預約    '),
             _checkbox_run(bool(data['guide_walkin'])), _text_run('散客'))
    center_on = bool(data['center_on'] or data['center_text'])
    other_on = bool(data['other_on'] or data['other_text'])
    center = (_checkbox_run(center_on), _text_run('塔塔加遊客中心：'), _blank_line(data['center_text']))
    other = (_checkbox_run(other_on), _text_run('其他：'), _blank_line(data['other_text']))
    has_incident = data['incident'] == '有'
    records = [line.strip() for line in data['incident_text'].split(NEWLINE) if line.strip()] if has_incident else []
    incident = [_checkbox_run(data['incident'] == '無'), _text_run('無  '), _checkbox_run(has_incident),
                _text_run('有：')]
    if len(records) == 1:  # 只有一筆：和範本一樣寫在「有：」後面
        incident.append(_blank_line(records[0]))
    return guide, center, other, incident, (records if len(records) > 1 else [])


def _add_incident_records(paragraph_el, records):
    """偶發事件有多筆時，接在「有：」下面一筆一行，加上編號（1. 2. 3. …）。"""
    anchor = paragraph_el
    for number, text in enumerate(records, start=1):
        new = copy.deepcopy(paragraph_el)
        _rebuild(new, _text_run(f"{number}. "), _blank_line(text))
        ppr = new.find(qn('w:pPr'))
        for old in ppr.findall(qn('w:ind')):
            ppr.remove(old)
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '560')
        ind.set(qn('w:hanging'), '560')
        jc = ppr.find(qn('w:jc'))
        ppr.insert(list(ppr).index(jc) if jc is not None else len(ppr), ind)
        anchor.addnext(new)
        anchor = new


def build_tataka_report(template_path, data, photo_files, captions, output_path, rotations=None, crops=None):
    """依塔塔加保育宣導回報單範本產生報告並存檔，回傳警告訊息清單。"""
    data, captions = g.xml_safe_inputs(data, captions)
    doc = Document(template_path)
    if not doc.tables or len(doc.tables[0].rows) < 10 or len(doc.tables[0].columns) < 4:
        raise ValueError("模板格式錯誤：塔塔加保育宣導回報單範本的表格需要 10 列、4 欄")
    table = doc.tables[0]
    warnings = []

    for row in range(1, 5):
        _relax_height(table, row)

    _set_cell_lines(_cell(table, 1, 0), split_names(data['volunteer']))   # 每位志工一行
    g._set_cell_value(_cell(table, 1, 1), f"{data['year']}年\n{data['month']}月{data['day']}日", center=True)
    g._set_cell_value(_cell(table, 1, 2), data['time_text'], center=True)
    weather = _cell(table, 1, 3)
    _rebuild(weather._tc.findall(qn('w:p'))[0], *[
        part for name in TATAKA_WEATHERS
        for part in (_checkbox_run(data['weather'] == name), _text_run(name))])

    guide, center, other, incident, more_records = _tataka_service_paragraphs(data)
    service_cell = _cell(table, 2, 1)
    p0, _, p2, p3 = service_cell._tc.findall(qn('w:p'))[:4]
    _rebuild(p0, *guide)
    _rebuild(p2, *center)
    _rebuild(p3, *other)
    incident_p = _cell(table, 3, 1)._tc.findall(qn('w:p'))[0]
    _rebuild(incident_p, *incident)
    _add_incident_records(incident_p, more_records)

    fill_cell_content(_cell(table, 4, 1), data['facility'], shift=480)
    fill_cell_content(_cell(table, 5, 1), data['visitors'], shift=480)
    fill_cell_content(_cell(table, 6, 1), data['reflection'], shift=480)

    # 相關疑問與解答：範例文字「1.相關疑問？答：因為……。」一律拿掉
    question_cell = _cell(table, 7, 1)
    head = question_cell.paragraphs[0]
    insert_content_after(head, data['questions'], shift=480, is_placeholder=g.is_question_placeholder)
    if not data['questions'].strip():
        for p in question_cell.paragraphs[1:]:
            if g.is_question_placeholder(p.text):
                p._p.getparent().remove(p._p)

    _add_tataka_photos(doc, table, photo_files, captions, warnings, rotations, crops)
    doc.save(output_path)
    return warnings


def _add_tataka_photos(doc, table, photo_files, captions, warnings, rotations, crops):
    widths = _grid_widths(table)
    right_w = _page_width_in(doc) * sum(widths[1:]) / sum(widths)
    fit_width = right_w - g.CELL_SIDE_PAD - 0.1
    items = g._load_photos(photo_files, captions, fit_width, PHOTO_MAX_HEIGHT_IN, warnings, rotations, crops)
    photos = items.get(TATAKA_PHOTO_CATEGORY, [])
    # 「名稱／詳細地點」「相片記錄」標題列和第一張照片不要分在兩頁
    for p in _cell(table, 8, 0).paragraphs + _cell(table, 8, 1).paragraphs:
        p.paragraph_format.keep_with_next = True
    if not photos:
        return
    template_tr = table._tbl.tr_lst[9]
    anchor = template_tr
    for item in photos:
        tr = copy.deepcopy(template_tr)
        anchor.addnext(tr)
        anchor = tr
        g._mark_row(_RowProxy(tr), 'w:cantSplit')
        name, _, place = item.caption.partition(chr(10))
        left, right = (_Cell(tc, table) for tc in tr.tc_lst[:2])
        _set_cell_lines(left, [line for line in (name.strip(), place.strip()) if line])
        _picture_into_cell(right, item)
    table._tbl.remove(template_tr)


def _set_cell_lines(cell, lines):
    """儲存格填入幾行置中的文字（標楷體 14pt，段落格式沿用儲存格第一段）；沒有文字就清成空白。"""
    paragraphs = cell._tc.findall(qn('w:p'))
    for extra in paragraphs[1:]:
        cell._tc.remove(extra)
    first = paragraphs[0]
    _clear_runs(first)
    template_ppr = first.find(qn('w:pPr'))
    anchor = first
    for i, line in enumerate(lines or ['']):
        paragraph = first
        if i:
            paragraph = OxmlElement('w:p')
            if template_ppr is not None:
                paragraph.append(copy.deepcopy(template_ppr))
            anchor.addnext(paragraph)
            anchor = paragraph
        ppr = paragraph.find(qn('w:pPr'))
        if ppr is None:
            ppr = OxmlElement('w:pPr')
            paragraph.insert(0, ppr)
        for old in ppr.findall(qn('w:jc')):
            ppr.remove(old)
        jc = OxmlElement('w:jc')
        jc.set(qn('w:val'), 'center')
        mark = ppr.find(qn('w:rPr'))   # 依 Word 的元素順序，段落標記的 rPr 要放在最後
        ppr.insert(list(ppr).index(mark) if mark is not None else len(ppr), jc)
        if line:
            paragraph.append(_text_run(line))


class _RowProxy:
    """只提供 _tr 屬性，讓 g._mark_row 可以直接處理 XML 列。"""

    def __init__(self, tr):
        self._tr = tr


# ---------------------------------------------------------------- 生態廁所回報單

TOILETS = (('m', '孟祿亭'), ('b', '白木林'))
ECO_PHOTO_CATEGORY = '監測及維護照片'
# (欄位後綴, 畫面標籤, 型態, 下拉選項, 是否必填)；前 7 項是（一）排泄物處理記錄，後 6 項是（二）堆積場處理記錄
EXCRETA_FIELDS = [
    ('total', '本次收集桶總量（kg）', 'entry', None, True),
    ('weather', '天候', 'combo', ('', '晴', '雨'), True),
    ('sludge', '汙泥桶清理', 'combo', ('', '有', '否'), True),
    ('nonrot', '從收集桶取出不能腐化垃圾量（kg）', 'entry', None, True),
    ('humus_new', '清空收集桶後添加新腐植土入桶量（kg）', 'entry', None, True),
    ('humus_spare', '收集腐植土備用量（kg）', 'entry', None, True),
    ('remark1', '備註', 'entry', None, False),
]
COMPOST_FIELDS = [
    ('obs', '生態觀察', 'entry', None, True),
    ('moisture', '溼度處理（範本：正常 60%）', 'entry', None, True),
    ('temp', '堆積溫度（範本：常溫）', 'entry', None, True),
    ('odor', '異味處理添加量（kg）', 'entry', None, True),
    ('turned', '翻堆', 'combo', ('', '是', '否'), True),
    ('remark2', '備註', 'entry', None, False),
]
TOILET_FIELD_NAMES = [f'{prefix}_{spec[0]}' for prefix, _ in TOILETS for spec in EXCRETA_FIELDS + COMPOST_FIELDS]


def build_ecotoilet_report(template_path, data, photo_files, captions, output_path, rotations=None, crops=None):
    """依生態廁所回報單範本產生報告並存檔，回傳警告訊息清單。"""
    data, captions = g.xml_safe_inputs(data, captions)
    doc = Document(template_path)
    if len(doc.tables) < 4:
        raise ValueError("模板格式錯誤：生態廁所回報單範本需要 4 個表格（基本資料、兩份記錄表、照片）")
    header, excreta, compost, photos_table = doc.tables[:4]
    warnings = []

    # 基本資料：第 2 列是「回報單 ○年○月○日」，之後是姓名、協勤日期、地點
    date_p = header._tbl.tr_lst[1].tc_lst[0].findall(qn('w:p'))[0]
    runs = date_p.findall(qn('w:r'))
    for extra in runs[1:]:
        date_p.remove(extra)
    for proof in date_p.findall(qn('w:proofErr')):
        date_p.remove(proof)
    date_run = _text_run(f"　　{data['year']} 年 {data['month']} 月 {data['day']} 日")
    if runs and runs[0].find(qn('w:rPr')) is not None:
        date_run.replace(date_run.find(qn('w:rPr')), copy.deepcopy(runs[0].find(qn('w:rPr'))))
        fonts = date_run.find(qn('w:rPr')).find(qn('w:rFonts'))
        if fonts is not None:
            for attr in ('ascii', 'hAnsi', 'eastAsia'):
                fonts.set(qn(f'w:{attr}'), g.KAI)
    date_p.append(date_run)
    g._set_cell_value(_cell(header, 2, 1), data['personnel'])
    g._set_cell_value(_cell(header, 3, 1),
                      f"{roc_date_text(data['start_year'], data['start_month'], data['start_day'])}　至　"
                      f"{roc_date_text(data['end_year'], data['end_month'], data['end_day'])}", left=True)
    g._set_cell_value(_cell(header, 4, 1), data['location'])

    # 兩份記錄表：每座廁所 4 列，最後一列是填寫列
    for table, fields in ((excreta, EXCRETA_FIELDS), (compost, COMPOST_FIELDS)):
        for index, (prefix, _) in enumerate(TOILETS):
            row = index * 4 + 3
            if row >= len(table._tbl.tr_lst):
                raise ValueError("模板格式錯誤：記錄表的列數與預期不符")
            tcs = table._tbl.tr_lst[row].tc_lst
            if len(tcs) < len(fields):
                raise ValueError("模板格式錯誤：記錄表的欄數與預期不符")
            for tc, spec in zip(tcs, fields):
                g._set_cell_value(_Cell(tc, table), data.get(f'{prefix}_{spec[0]}', ''), center=True)

    _tidy_ecotoilet_layout(doc, excreta)
    for table in (excreta, compost):
        _keep_toilet_blocks_together(table)
    _fill_ecotoilet_sections(doc, data)
    _add_ecotoilet_photos(doc, photos_table, photo_files, captions, warnings, rotations, crops)
    doc.save(output_path)
    return warnings


def _keep_toilet_blocks_together(table):
    """每座廁所的 4 列（標題、欄位名稱、單位、填寫列）不要被分到兩頁。"""
    trs = table._tbl.tr_lst
    for start in range(0, len(trs), 4):
        for tr in trs[start:start + 3]:
            for p in tr.iter(qn('w:p')):
                ppr = p.find(qn('w:pPr'))
                if ppr is None:
                    ppr = OxmlElement('w:pPr')
                    p.insert(0, ppr)
                if ppr.find(qn('w:keepNext')) is None:
                    style = ppr.find(qn('w:pStyle'))  # 依 Word 的元素順序，keepNext 要放在 pStyle 之後
                    ppr.insert(0 if style is None else list(ppr).index(style) + 1, OxmlElement('w:keepNext'))
        for tr in trs[start:start + 4]:
            g._mark_row(_RowProxy(tr), 'w:cantSplit')


def _tidy_ecotoilet_layout(doc, excreta_table):
    """範本靠一整排空白段落把（二）推到下一頁；填入資料後列高改變，標題可能被留在上一頁、表格卻在下一頁。
    改成：（一）～（三）的小標題一定和後面的表格在同一頁，並把表格與（二）之間的空白段落收成一行。"""
    for paragraph in doc.paragraphs:
        if re.match(r'^[（(][一二三][）)]', paragraph.text.strip()):
            paragraph.paragraph_format.keep_with_next = True
    blank = excreta_table._tbl.getnext()
    kept = False
    while blank is not None and blank.tag == qn('w:p') and g._is_empty_paragraph(blank):
        nxt = blank.getnext()
        if kept:
            blank.getparent().remove(blank)
        kept = True
        blank = nxt


SECTION_HEADING_RE = re.compile(r'^[一二三四五六七八九十]+、')


def _fill_ecotoilet_sections(doc, data):
    """二、自然資源狀況　三、環境清理及遊客宣導情形　四、協勤心得與建議：內容放在各標題（含接續的第二行）之後。"""
    sections = (('二、自然資源狀況', 'nature', True), ('三、環境清理及遊客宣導情形', 'environment', True),
                ('四、協勤心得與建議', 'reflection', False))
    for start, key, blank_after in sections:
        paragraphs = doc.paragraphs
        index = next((i for i, p in enumerate(paragraphs) if p.text.strip().startswith(start)), None)
        if index is None:
            continue
        last = index
        while (last + 1 < len(paragraphs) and paragraphs[last + 1].text.strip()
               and not SECTION_HEADING_RE.match(paragraphs[last + 1].text.strip())):
            last += 1  # 標題太長時接續在下一個段落
        insert_content_after(paragraphs[last], data.get(key, ''), trailing_blank=blank_after)


def _add_ecotoilet_photos(doc, table, photo_files, captions, warnings, rotations, crops):
    rows = table._tbl.tr_lst
    if len(rows) < 2 or len(rows[0].tc_lst) < 2:
        raise ValueError("模板格式錯誤：照片表格需要「照片圖檔／照片文字」兩列、兩欄")
    widths = _grid_widths(table)
    col_w = _page_width_in(doc) * widths[0] / sum(widths)
    fit_width = col_w - g.CELL_SIDE_PAD
    photos = g._load_photos(photo_files, captions, fit_width, PHOTO_MAX_HEIGHT_IN, warnings, rotations,
                            crops).get(ECO_PHOTO_CATEGORY, [])
    pair_count = max(1, math.ceil(len(photos) / 2))
    pairs = [(rows[i], rows[i + 1]) for i in range(0, len(rows) - 1, 2)]
    template_pair = pairs[0]
    while len(pairs) < pair_count:  # 範本只有 2 組，不夠就複製
        last = pairs[-1][1]
        new_image, new_text = copy.deepcopy(template_pair[0]), copy.deepcopy(template_pair[1])
        last.addnext(new_image)
        new_image.addnext(new_text)
        pairs.append((new_image, new_text))
    for image_row, text_row in pairs[pair_count:]:  # 照片不多就刪掉用不到的列
        table._tbl.remove(image_row)
        table._tbl.remove(text_row)
    pairs = pairs[:pair_count]
    for pair_index, (image_row, text_row) in enumerate(pairs):
        g._mark_row(_RowProxy(image_row), 'w:cantSplit')
        g._mark_row(_RowProxy(text_row), 'w:cantSplit')
        for col in range(2):
            image_cell = _Cell(image_row.tc_lst[col], table)
            text_cell = _Cell(text_row.tc_lst[col], table)
            index = pair_index * 2 + col
            if index < len(photos):
                _picture_into_cell(image_cell, photos[index])
                g._set_cell_value(text_cell, photos[index].caption)
            else:  # 沒有照片的格子：清掉範本的「照片圖檔」「照片文字」
                g._set_cell_value(image_cell, '')
                g._set_cell_value(text_cell, '')


# ---------------------------------------------------------------- 畫面

class IncidentList(ttk.Frame):
    """可以按「+」增加筆數的記錄清單：每筆一個輸入框，多於一筆時每筆右邊有「－」可以刪除。
    和輸入框一樣有 get()：回傳各筆（不含空白筆）以換行串起來；set_value() 依換行還原成多筆。"""

    def __init__(self, master):
        super().__init__(master)
        self.columnconfigure(0, weight=1)
        self.rows = []              # [(框架, 輸入框, 刪除鈕)]
        self._callback = None
        self.add_button = ttk.Button(self, text="+ 新增一筆", command=self.add_row)
        self.add_row(focus=False)

    def on_change(self, callback):
        self._callback = callback

    def _changed(self, _event=None):
        if self._callback:
            self._callback()

    def add_row(self, text='', focus=True):
        frame = ttk.Frame(self)
        frame.columnconfigure(0, weight=1)
        entry = ttk.Entry(frame)
        entry.grid(row=0, column=0, sticky=(tk.W, tk.E))
        entry.insert(0, text)
        for sequence in ('<KeyRelease>', '<FocusOut>', '<<Paste>>', '<<Cut>>'):
            entry.bind(sequence, self._changed, add='+')
        entry.bind('<Return>', lambda e: (self.add_row(), 'break')[1])   # 在輸入框按 Enter：新增下一筆
        item = [frame, entry, None]
        item[2] = ttk.Button(frame, text="－", width=3, command=lambda: self.remove_row(item))
        item[2].grid(row=0, column=1, padx=(6, 0))
        self.rows.append(item)
        self._layout()
        if focus:
            entry.focus_set()
            self._changed()

    def remove_row(self, item):
        if len(self.rows) <= 1:       # 只剩一筆：清空文字就好
            item[1].delete(0, tk.END)
        else:
            self.rows.remove(item)
            item[0].destroy()
        self._layout()
        self._changed()

    def _layout(self):
        only_one = len(self.rows) == 1
        for index, (frame, _, button) in enumerate(self.rows):
            frame.grid(row=index, column=0, sticky=(tk.W, tk.E), pady=2)
            button.state(['disabled'] if only_one else ['!disabled'])
        self.add_button.grid(row=len(self.rows), column=0, sticky=tk.W, pady=(4, 0))

    def get(self):
        return NEWLINE.join(text for text in (entry.get().strip() for _, entry, _ in self.rows) if text)

    def set_value(self, value):
        lines = [line.strip() for line in str(value).split(NEWLINE) if line.strip()] or ['']
        for frame, _, _ in self.rows:
            frame.destroy()
        self.rows = []
        for line in lines:
            self.add_row(line, focus=False)


def _combo(parent, values, width=10):
    return ttk.Combobox(parent, values=list(values), state='readonly', width=width, justify=tk.CENTER)


class _FormGUI(g.ReportGeneratorGUI):
    """塔塔加、生態廁所共用的部分：未填欄位用 blank_fields 統一判斷。"""

    def blank_fields(self, with_captions=False):
        """沒填的欄位 [(分頁, 欄位名稱)]；with_captions=True 時，沒填說明的照片也算（用來標分頁的 ●）。"""
        raise NotImplementedError

    def _photo_blanks(self, with_captions):
        found = []
        for category in self.PHOTO_CATEGORIES:
            if not self.photo_files[category]:
                found.append((self.PHOTOS_TAB, f"{category}（尚未加入照片）"))
            elif with_captions and self.missing_caption_count(category):
                found.append((self.PHOTOS_TAB, f"{category}（有照片沒填資料）"))
        return found

    def _text_blanks(self):
        return [(self.FIELD_TAB[key], self.FIELD_NAMES[key]) for key, widget in self.text_fields.items()
                if not widget.get('1.0', tk.END).strip()]

    def _tabs_with_blanks(self):
        return {tab for tab, _ in self.blank_fields(with_captions=True)}

    def find_blank_fields(self, data):
        blanks = {}
        for tab, label in self.blank_fields():
            blanks.setdefault(tab, []).append(label)
        return {tab: blanks[tab] for tab in self.TAB_ORDER if tab in blanks}

    def _date_entries(self, parent, prefix=''):
        for key, unit, width in (('year', '年', 8), ('month', '月', 5), ('day', '日', 5)):
            entry = ttk.Entry(parent, width=width, justify=tk.CENTER)
            entry.pack(side=tk.LEFT)
            ttk.Label(parent, text=unit).pack(side=tk.LEFT, padx=g.DATE_UNIT_PADX)
            setattr(self, f'{prefix}{key}', entry)

    def _valid_date(self, prefix=''):
        try:
            g.datetime(int(getattr(self, f'{prefix}year').get()), int(getattr(self, f'{prefix}month').get()),
                       int(getattr(self, f'{prefix}day').get()))
        except ValueError:
            return False
        return True


class TatakaReportGUI(_FormGUI):
    REPORT_ID = 'tataka'
    REPORT_NAME = '塔塔加保育宣導回報單'
    DEFAULT_TEMPLATE_PATH = os.path.join(g.BASE_DIR, '塔塔加保育宣導回報單.docx')
    DEFAULT_AUTOSAVE_PATH = os.path.join(g.DATA_DIR, 'yscv_tataka_autosave.json')
    TAB_SERVICE = '服務項目'
    TAB_OPINION = '設施與遊客意見'
    TAB_REFLECT = '心得與疑問解答'
    TAB_PHOTO = '相片記錄'
    PHOTO_CATEGORIES = [TATAKA_PHOTO_CATEGORY]
    PHOTO_SHORT_NAMES = {TATAKA_PHOTO_CATEGORY: '相片記錄'}
    LEGACY_PHOTO_CATEGORIES = {}
    PHOTOS_TAB = TAB_PHOTO
    PHOTO_HEADING = '相片記錄'
    PHOTO_PANEL_MIN_WIDTH = 560
    CAPTION_LABELS = ('名稱：', '詳細地點：')
    ENTRY_FIELDS = ['year', 'month', 'day', 'volunteer', 'time_text', 'weather', 'guide_booked', 'guide_walkin',
                    'center_on', 'center_text', 'other_on', 'other_text', 'incident', 'incident_text']
    REPORT_DATE_FIELDS = ('year', 'month', 'day')
    FIELD_TAB = {'facility': TAB_OPINION, 'visitors': TAB_OPINION, 'reflection': TAB_REFLECT,
                 'questions': TAB_REFLECT}
    FIELD_NAMES = {'facility': '園區設施現況及改善意見', 'visitors': '遊客意見', 'reflection': '值勤心得',
                   'questions': '相關疑問與解答'}
    TAB_ORDER = [g.TAB_BASIC, TAB_SERVICE, TAB_OPINION, TAB_REFLECT, TAB_PHOTO]

    def build_tabs(self, notebook):
        self._basic_tab(notebook)
        self._service_tab(notebook)
        self.create_text_tab(notebook, self.TAB_OPINION, [
            ("園區設施現況及改善意見：", 'facility'),
            ("遊客意見：", 'visitors'),
        ], height=6)
        self.create_text_tab(notebook, self.TAB_REFLECT, [
            ("值勤心得：", 'reflection'),
            ("相關疑問與解答：（將服勤期間相關疑問，以及透過職員或資深志工尋得之答案記錄下來）", 'questions'),
        ], height=6)
        self.create_photo_tab(notebook, self.TAB_PHOTO)

    def _tab_frame(self, notebook, title):
        frame = ttk.Frame(notebook, padding="10")
        notebook.add(frame, text=title)
        self._tab_frames[title] = frame
        frame.columnconfigure(1, weight=1)
        return frame

    def _basic_tab(self, notebook):
        frame = self._tab_frame(notebook, g.TAB_BASIC)
        ttk.Label(frame, text="1. 志工姓名:", style='Field.TLabel').grid(row=0, column=0, sticky=tk.W, pady=5)
        self.volunteer = ttk.Entry(frame, width=50)
        self.volunteer.grid(row=0, column=1, sticky=(tk.W, tk.E), padx=5)
        ttk.Label(frame, text="多位志工請用空白、頓號（、）或逗號隔開，報告裡會自動一人一行；"
                              "外文姓名中間的空白不會被拆開（多位外文姓名請用逗號隔開）", style='Hint.TLabel').grid(
            row=1, column=1, sticky=tk.W, padx=5)
        ttk.Label(frame, text="2. 值勤日期:", style='Field.TLabel').grid(row=2, column=0, sticky=tk.W, pady=5)
        date_frame = ttk.Frame(frame)
        date_frame.grid(row=2, column=1, sticky=tk.W, padx=5)
        self._date_entries(date_frame)
        ttk.Label(frame, text="3. 時間:", style='Field.TLabel').grid(row=3, column=0, sticky=tk.W, pady=5)
        time_frame = ttk.Frame(frame)
        time_frame.grid(row=3, column=1, sticky=tk.W, padx=5)
        self.time_text = ttk.Entry(time_frame, width=24)
        self.time_text.pack(side=tk.LEFT)
        ttk.Label(time_frame, text="例如 09:00～16:00", style='Hint.TLabel').pack(side=tk.LEFT, padx=8)
        ttk.Label(frame, text="4. 天氣:", style='Field.TLabel').grid(row=4, column=0, sticky=tk.W, pady=5)
        self.weather = _combo(frame, ('',) + TATAKA_WEATHERS, width=8)
        self.weather.grid(row=4, column=1, sticky=tk.W, padx=5)

    def _service_tab(self, notebook):
        frame = self._tab_frame(notebook, self.TAB_SERVICE)
        ttk.Label(frame, text="服務項目", style='Heading.TLabel').grid(row=0, column=0, columnspan=2, sticky=tk.W,
                                                                       pady=(0, 6))
        ttk.Label(frame, text="帶團解說:", style='Field.TLabel').grid(row=1, column=0, sticky=tk.W, pady=5)
        guide = ttk.Frame(frame)
        guide.grid(row=1, column=1, sticky=tk.W, padx=5)
        self.guide_booked = CheckField(guide, '預約')
        self.guide_booked.pack(side=tk.LEFT, padx=(0, 18))
        self.guide_walkin = CheckField(guide, '散客')
        self.guide_walkin.pack(side=tk.LEFT)
        self.center_on = CheckField(frame, '塔塔加遊客中心：')
        self.center_on.grid(row=2, column=0, sticky=tk.W, pady=5)
        self.center_text = ttk.Entry(frame, width=40)
        self.center_text.grid(row=2, column=1, sticky=(tk.W, tk.E), padx=5)
        self.other_on = CheckField(frame, '其他：')
        self.other_on.grid(row=3, column=0, sticky=tk.W, pady=5)
        self.other_text = ttk.Entry(frame, width=40)
        self.other_text.grid(row=3, column=1, sticky=(tk.W, tk.E), padx=5)
        ttk.Label(frame, text="（填了文字會自動打勾）", style='Hint.TLabel').grid(row=4, column=1, sticky=tk.W, padx=5)
        ttk.Label(frame, text="偶發事件", style='Heading.TLabel').grid(row=5, column=0, columnspan=2, sticky=tk.W,
                                                                       pady=(18, 6))
        ttk.Label(frame, text="偶發事件:", style='Field.TLabel').grid(row=6, column=0, sticky=tk.W, pady=5)
        incident = ttk.Frame(frame)
        incident.grid(row=6, column=1, sticky=(tk.W, tk.E), padx=5)
        incident.columnconfigure(2, weight=1)
        self.incident = _combo(incident, ('', '無', '有'), width=6)
        self.incident.grid(row=0, column=0)
        ttk.Label(incident, text="選「有」請在下面逐筆記錄，需要多筆就按「+」", style='Hint.TLabel').grid(
            row=0, column=1, padx=(14, 0))
        ttk.Label(frame, text="事件記錄:", style='Field.TLabel').grid(row=7, column=0, sticky=(tk.W, tk.N), pady=5)
        self.incident_text = IncidentList(frame)
        self.incident_text.grid(row=7, column=1, sticky=(tk.W, tk.E), padx=5, pady=2)

    def blank_fields(self, with_captions=False):
        found = []
        basic = g.TAB_BASIC
        if not self.volunteer.get().strip():
            found.append((basic, '志工姓名'))
        if any(not getattr(self, n).get().strip() for n in ('year', 'month', 'day')):
            found.append((basic, '值勤日期'))
        if not self.time_text.get().strip():
            found.append((basic, '時間'))
        if not self.weather.get():
            found.append((basic, '天氣'))
        services = (self.guide_booked, self.guide_walkin, self.center_on, self.center_text, self.other_on,
                    self.other_text)
        if not any(w.get().strip() for w in services):
            found.append((self.TAB_SERVICE, '服務項目（至少選一項）'))
        if not self.incident.get():
            found.append((self.TAB_SERVICE, '偶發事件（無／有）'))
        elif self.incident.get() == '有' and not self.incident_text.get().strip():
            found.append((self.TAB_SERVICE, '偶發事件的說明'))
        return found + self._text_blanks() + self._photo_blanks(with_captions)

    def collect_data(self):
        self._commit_caption()
        if not self._valid_date():
            return None, "「值勤日期」不是有效的日期，請檢查"
        data = {name: getattr(self, name).get().strip() for name in self.ENTRY_FIELDS}
        for key, widget in self.text_fields.items():
            data[key] = widget.get('1.0', tk.END).strip()
        return data, None

    def _name_parts(self, fields):
        try:
            date = f"{int(fields['year'])}{int(fields['month']):02d}{int(fields['day']):02d}"
        except (KeyError, ValueError):
            date = ''
        return date, g.name_label(fields.get('volunteer', ''))

    def output_name(self, data):
        date, who = self._name_parts(data)
        return '_'.join(part for part in (date, '塔塔加保育宣導回報單', who) if part)

    def draft_name(self, fields, now):
        date, who = self._name_parts(fields)
        parts = [date, '塔塔加保育宣導回報單草稿', who, now.strftime('%m%d_%H%M')]
        return '_'.join(part for part in parts if part)

    def make_report(self, data, photo_files, captions, output_path, rotations=None, crops=None):
        return build_tataka_report(self.template_path, data, photo_files, captions, output_path, rotations, crops)


class EcoToiletReportGUI(_FormGUI):
    REPORT_ID = 'ecotoilet'
    REPORT_NAME = '生態廁所回報單'
    DEFAULT_TEMPLATE_PATH = os.path.join(g.BASE_DIR, '生態廁所回報單.docx')
    DEFAULT_AUTOSAVE_PATH = os.path.join(g.DATA_DIR, 'yscv_ecotoilet_autosave.json')
    TAB_NATURE = '自然資源/環境清理'
    TAB_REFLECT = '協勤心得與建議'
    TAB_PHOTO = '監測及維護照片'
    PHOTO_CATEGORIES = [ECO_PHOTO_CATEGORY]
    PHOTO_SHORT_NAMES = {ECO_PHOTO_CATEGORY: '監測及維護'}
    LEGACY_PHOTO_CATEGORIES = {}
    PHOTOS_TAB = TAB_PHOTO
    PHOTO_HEADING = '監測及維護照片'
    PHOTO_PANEL_MIN_WIDTH = 560
    CAPTION_LABELS = ('照片文字：',)
    ENTRY_FIELDS = list(g.ENTRY_FIELDS) + TOILET_FIELD_NAMES
    REPORT_DATE_FIELDS = g.REPORT_DATE_FIELDS
    DEFAULT_FIELDS = {'location': ECO_DEFAULT_LOCATION}
    FIELD_TAB = {'nature': TAB_NATURE, 'environment': TAB_NATURE, 'reflection': TAB_REFLECT}
    FIELD_NAMES = {'nature': '二、自然資源狀況', 'environment': '三、環境清理及遊客宣導情形',
                   'reflection': '四、協勤心得與建議'}
    TAB_ORDER = [g.TAB_BASIC] + [name for _, name in TOILETS] + [TAB_NATURE, TAB_REFLECT, TAB_PHOTO]

    def build_tabs(self, notebook):
        self.create_basic_info_tab(notebook)
        self._end_auto = None      # 迄日目前是自動帶入的哪一天（使用者自己改過就不再自動更新）
        self._setup_two_day_duty()
        for prefix, name in TOILETS:
            self._toilet_tab(notebook, prefix, name)
        self.create_text_tab(notebook, self.TAB_NATURE, [
            ("二、自然資源狀況（如特殊植物、動物、人文、地質地形、水文等資源之地點、數量及周圍環境紀錄等）：", 'nature'),
            ("三、環境清理及遊客宣導情形（如垃圾清理及帶下山、遊客登山安全、環境保護宣導等情形）：", 'environment'),
        ], height=6)
        self.create_text_tab(notebook, self.TAB_REFLECT, [("四、協勤心得與建議：", 'reflection')], height=10)
        self.create_photo_tab(notebook, self.TAB_PHOTO)
        self._add_toilet_icons(notebook)

    def _setup_two_day_duty(self):
        """主峰線生態廁所的勤務是兩天：起日填好之後，迄日自動帶入隔天（使用者可以再改）。"""
        frame = self._tab_frames[g.TAB_BASIC]
        ttk.Label(frame, text="兩天的勤務：填好協勤日期的起日，迄日會自動帶入隔天（可自行修改）", style='Hint.TLabel').grid(
            row=4, column=1, sticky=tk.W, padx=5)
        for part in DATE_PARTS:
            entry = getattr(self, f'start_{part}')
            for sequence in ('<KeyRelease>', '<FocusOut>', '<<Paste>>', '<<Cut>>'):
                entry.bind(sequence, self._sync_end_date, add='+')

    def _start_date(self):
        """起日轉成 datetime（年可以是民國年或西元年）；還沒填好或不是有效日期回傳 None。"""
        try:
            year, month, day = (int(getattr(self, f'start_{part}').get()) for part in DATE_PARTS)
            return g.datetime(year + (g.ROC_YEAR_OFFSET if year < g.ROC_YEAR_OFFSET else 0), month, day), year
        except ValueError:
            return None

    def _sync_end_date(self, _event=None):
        """起日是有效日期時，迄日是空白、或還是上次自動帶入的值，就改成起日的隔天。"""
        parsed = self._start_date()
        if parsed is None:
            return
        start, typed_year = parsed
        nxt = start + g.timedelta(days=1)
        roc = typed_year < g.ROC_YEAR_OFFSET
        value = (str(nxt.year - g.ROC_YEAR_OFFSET if roc else nxt.year), str(nxt.month), str(nxt.day))
        current = tuple(getattr(self, f'end_{part}').get().strip() for part in DATE_PARTS)
        if current != ('', '', '') and current != self._end_auto:
            return              # 使用者自己填過迄日，不要蓋掉
        for part, text in zip(DATE_PARTS, value):
            g.widget_set(getattr(self, f'end_{part}'), text)
        self._end_auto = value
        self.status_label.config(text=f"已自動帶入迄日：{value[0]}年{value[1]}月{value[2]}日（隔天），需要的話請直接修改")
        self._refresh_tab_marks()

    def _clear_widgets(self):
        super()._clear_widgets()
        self._end_auto = None

    def _add_toilet_icons(self, notebook):
        """兩座廁所的分頁標籤前面放廁所圖示（找不到彩色表情符號字型就不放，標籤文字不變）。"""
        self._tab_icons = []        # 要保留參照，圖片才不會被回收
        pixels = round(self.root.winfo_fpixels('1i') / 96 * TAB_ICON_PX)
        for _, name in TOILETS:
            icon = g.emoji_image(g.REPORT_ICONS['ecotoilet'], pixels)
            if icon is None:
                continue
            photo = g.ImageTk.PhotoImage(icon, master=self.root)
            self._tab_icons.append(photo)
            notebook.tab(self._tab_frames[name], image=photo, compound=tk.LEFT)

    def _toilet_tab(self, notebook, prefix, name):
        inner = self.make_scrollable_tab(notebook, name)
        inner.columnconfigure(0, weight=0)  # 標籤欄只要剛好放得下文字，多出來的寬度給輸入框
        inner.columnconfigure(1, weight=1)
        row = 0
        for title, specs in (('（一）生態廁所排泄物處理記錄表', EXCRETA_FIELDS), ('（二）生態廁所堆積場處理記錄表', COMPOST_FIELDS)):
            ttk.Label(inner, text=f"{name}生態廁所　{title}", style='Heading.TLabel').grid(
                row=row, column=0, columnspan=2, sticky=tk.W, pady=(12 if row else 0, 6))
            row += 1
            for suffix, label, kind, values, _ in specs:
                ttk.Label(inner, text=label + ':', style='Field.TLabel').grid(row=row, column=0, sticky=tk.W,
                                                                               pady=4, padx=(0, 10))
                if kind == 'combo':
                    widget = _combo(inner, values, width=8)
                    widget.grid(row=row, column=1, sticky=tk.W)
                else:
                    widget = ttk.Entry(inner, width=50)
                    widget.grid(row=row, column=1, sticky=(tk.W, tk.E), padx=(0, 6))
                setattr(self, f'{prefix}_{suffix}', widget)
                row += 1

    def blank_fields(self, with_captions=False):
        found = []
        basic = g.TAB_BASIC
        if not self.personnel.get().strip():
            found.append((basic, '協勤人員姓名'))
        if not self.location.get().strip():
            found.append((basic, '協勤地點'))
        for prefix, name in TOILETS:
            for suffix, label, _, _, required in EXCRETA_FIELDS + COMPOST_FIELDS:
                if required and not getattr(self, f'{prefix}_{suffix}').get().strip():
                    found.append((name, label))
        return found + self._text_blanks() + self._photo_blanks(with_captions)

    def collect_data(self):
        data, error = super().collect_data()
        if error:
            return None, error
        for key in TOILET_FIELD_NAMES:
            data[key] = getattr(self, key).get().strip()
        return data, None

    @staticmethod
    def _place_name(fields):
        """檔名裡的地點；地點本身已經以「生態廁所」結尾（預設值）就不再重複加。"""
        location = g._INVALID_FILENAME_CHARS.sub('', fields.get('location', '')).strip()
        return location[:-len('生態廁所')] if location.endswith('生態廁所') else location

    def output_name(self, data):
        parts = [g.roc_date_range(data), f"{self._place_name(data)}生態廁所回報單", g.name_label(data.get('personnel', ''))]
        return '_'.join(part for part in parts if part)

    def draft_name(self, fields, now):
        try:
            date_part = g.roc_date_range(fields)
        except (KeyError, ValueError):
            date_part = ''
        parts = [date_part, f"{self._place_name(fields)}生態廁所回報單草稿", g.name_label(fields.get('personnel', '')),
                 now.strftime('%m%d_%H%M')]
        return '_'.join(part for part in parts if part)

    def make_report(self, data, photo_files, captions, output_path, rotations=None, crops=None):
        return build_ecotoilet_report(self.template_path, data, photo_files, captions, output_path, rotations, crops)
