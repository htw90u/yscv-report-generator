"""GUI 流程測試（需要可用的顯示環境）。執行：python -m unittest discover -s tests"""
import gc
import os
import sys
import tempfile
import time
import tkinter as tk
import unittest
from tkinter import ttk, filedialog
from unittest import mock

from PIL import Image
from docx import Document

from _common import TEMPLATE
import yscv_report_generator as g
from test_report import unique_cells

messagebox = g.messagebox  # 程式自己的對話框（可以替換成假的）


def style_lookup(root, style_name, option, state=None):
    return ttk.Style(root).lookup(style_name, option, state or [])


_REAL_EXPORT_PDF = g.export_pdf  # setUp 會把 g.export_pdf 換成假的，這裡留一份真的給「沒有 Word」的測試用


def tkfont_linespace(root, widget):
    import tkinter.font as tkfont
    return tkfont.Font(root=root, font=widget.cget('font')).metrics('linespace')

class GuiTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = tk.Tk()
        self.autosave = os.path.join(self.tmp.name, 'autosave.json')
        self.settings = os.path.join(self.tmp.name, 'settings.json')
        self.app = g.ReportGeneratorGUI(self.root, autosave_path=self.autosave, restore_on_start=False,
                                      settings_path=self.settings)
        self.app.template_path = TEMPLATE
        self.shown = []
        self.pdf_calls = []
        self.confirm_calls = []  # 產生前「有欄位未填」的確認視窗；測試預設按「是」

        def fake_confirm(title, message, **kwargs):
            self.confirm_calls.append((title, message, kwargs))
            return True


        def fake_pdf(docx, pdf):
            self.pdf_calls.append((docx, pdf))
            with open(pdf, 'wb') as f:
                f.write(b'%PDF-fake')

        for p in (
            mock.patch.object(messagebox, 'showerror', lambda t, m: self.shown.append(('err', m))),
            mock.patch.object(messagebox, 'showinfo', lambda t, m: self.shown.append(('info', m))),
            mock.patch.object(g, 'export_pdf', fake_pdf),
            mock.patch.object(messagebox, 'askyesno', fake_confirm),
        ):
            p.start()
            self.addCleanup(p.stop)

    def tearDown(self):
        self.root.destroy()
        self.app = self.root = None
        # 在主執行緒先回收 Tk 物件；否則可能在背景執行緒被 GC，造成 Tcl_AsyncDelete 當機
        gc.collect()
        self.tmp.cleanup()

    def fill_dates(self, month='10', day='12', start_day='1', end_day='2'):
        for e, v in [(self.app.year, '115'), (self.app.month, month), (self.app.day, day),
                     (self.app.start_year, '115'), (self.app.start_month, '10'), (self.app.start_day, start_day),
                     (self.app.end_year, '115'), (self.app.end_month, '10'), (self.app.end_day, end_day)]:
            e.delete(0, tk.END)
            e.insert(0, v)

    def test_date_validation(self):
        self.app.generate_report()
        self.assertIn('日期', self.shown[-1][1])
        self.fill_dates(month='2', day='30')
        self.app.generate_report()
        self.assertIn('填表日期', self.shown[-1][1])
        self.fill_dates(start_day='5', end_day='1')
        self.app.generate_report()
        self.assertIn('起日', self.shown[-1][1])

    def test_photo_list_and_caption(self):
        p = os.path.join(self.tmp.name, 'x.jpg')
        Image.new('RGB', (30, 20)).save(p)
        cat = g.PHOTO_CATEGORIES[0]
        table = self.app.listboxes[cat]
        with mock.patch.object(filedialog, 'askopenfilenames', lambda **k: (p,)):
            self.app.select_photos(cat, table)
            self.app.select_photos(cat, table)
        self.assertEqual(table.count(), 1)  # 重複加入同一張不會變兩張
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        self.app.caption_var.set('測試說明')
        self.app._caption_return()
        self.assertEqual(self.app.photo_captions[cat][p], '測試說明')
        self.assertEqual(table.item('0', 'values')[1], '測試說明')
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: True):
            self.app.clear_all()
        self.assertEqual(table.count(), 0)
        self.assertFalse(self.app.photo_files[cat] or self.app.photo_captions[cat])

    def _wait_until_done(self, timeout=20):
        t0 = time.time()
        while time.time() - t0 < timeout and str(self.app.generate_btn['state']) == 'disabled':
            self.root.update()
            time.sleep(0.05)
        self.root.update()

    def test_default_filename_and_pdf_generated(self):
        self.fill_dates()
        self.app.location.insert(0, '排雲山莊')
        out = os.path.join(self.tmp.name, 'o.docx')
        captured = {}

        def ask(**kw):
            captured.update(kw)
            return out

        with mock.patch.object(filedialog, 'asksaveasfilename', ask), \
                mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertEqual(captured['initialfile'], '1151001-02_排雲山莊協勤回報單.docx')
        self.assertEqual(self.pdf_calls, [(out, os.path.join(self.tmp.name, 'o.pdf'))])
        self.assertTrue(os.path.exists(os.path.join(self.tmp.name, 'o.pdf')))

    def test_pdf_failure_is_reported_but_docx_kept(self):
        self.fill_dates()
        out = os.path.join(self.tmp.name, 'o.docx')

        def boom(docx, pdf):
            raise RuntimeError('沒有 Word')

        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out), \
                mock.patch.object(g, 'export_pdf', boom), \
                mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertTrue(os.path.exists(out))
        self.assertTrue(any(k == 'info' and 'PDF 未產生' in m for k, m in self.shown))

    def test_auto_open_is_always_on_and_has_no_option(self):
        self.assertFalse(hasattr(self.app, 'auto_open'))
        texts = [w.cget('text') for w in self._all_widgets(self.root) if isinstance(w, (ttk.Checkbutton, ttk.Button))]
        self.assertNotIn('產生後自動開啟', texts)
        self.fill_dates()
        out = os.path.join(self.tmp.name, 'o.docx')
        opened = []
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out), \
                mock.patch.object(g, 'open_file', lambda path: opened.append(path)):
            self.app.generate_report()
            self._wait_until_done()
        self.assertEqual(opened, [out])

    def test_button_order_and_styles(self):
        frame = self.app.generate_btn.master
        labels = [w.cget('text') for w in frame.winfo_children() if isinstance(w, ttk.Button)]
        self.assertEqual(labels, ['載入草稿(L)', '儲存草稿(S)', '產生報告(G)', '清除所有資料', '離開(E)'])
        self.assertEqual(str(self.app.clear_btn.cget('style')), 'Danger.TButton')
        self.assertEqual(str(self.app.generate_btn.cget('style')), 'Primary.TButton')
        palette = g.THEMES[g.DEFAULT_THEME]
        for style_name, key in (('Primary.TButton', 'primary'), ('Danger.TButton', 'danger')):
            self.assertEqual(style_lookup(self.root, style_name, 'background'), palette[key])
            self.assertEqual(style_lookup(self.root, style_name, 'foreground'), palette[f'{key}_fg'])
            self.assertEqual(style_lookup(self.root, style_name, 'background', ['active']), palette[f'{key}_active'])
        # 離開維持預設灰色，不上色
        leave = [w for w in frame.winfo_children() if isinstance(w, ttk.Button) and w.cget('text') == '離開(E)'][0]
        self.assertIn(str(leave.cget('style')), ('', 'TButton'))

    def test_open_file_failure_does_not_raise(self):
        with mock.patch.object(os, 'startfile', side_effect=OSError('boom'), create=True):
            g.open_file(os.path.join(self.tmp.name, 'nope.docx'))

    def test_generate_in_thread(self):
        self.fill_dates()
        self.app.text_fields['nature'].insert('1.0', '看到山羌\n● 已有bullet')
        out = os.path.join(self.tmp.name, 'o.docx')
        opened = []
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out), \
                mock.patch.object(g, 'open_file', lambda path: opened.append(path)):
            self.app.generate_report()
            self.assertEqual(str(self.app.generate_btn['state']), 'disabled')
            self._wait_until_done()
        self.assertEqual(str(self.app.generate_btn['state']), 'normal')
        self.assertEqual(opened, [out])  # 產生後自動開啟
        self.assertTrue(self.root.winfo_exists())  # 工具沒有被關掉
        text = '\n'.join(c.text for c in unique_cells(Document(out)))
        self.assertEqual(text.count('看到山羌'), 1)
        self.assertNotIn('● ●', text)

    def _add_photo(self, name='d.jpg', caption='說明X'):
        p = os.path.join(self.tmp.name, name)
        Image.new('RGB', (30, 20)).save(p)
        cat = g.PHOTO_CATEGORIES[2]  # 自然資源狀況
        self.app.photo_files[cat].append(p)
        self.app.photo_captions[cat][p] = caption
        self.app._render_list(cat)
        return cat, p

    def _save_draft_to(self, path):
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: path):
            self.app.save_draft()

    def _load_draft_from(self, path, confirm=True):
        with mock.patch.object(filedialog, 'askopenfilename', lambda **k: path), \
                mock.patch.object(messagebox, 'askyesno', lambda *a, **k: confirm):
            self.app.load_draft()

    def test_draft_roundtrip(self):
        self.fill_dates()
        self.app.personnel.insert(0, '王小明')
        self.app.text_fields['nature'].insert('1.0', '看到山羌\n第二行')
        cat, p = self._add_photo()
        draft = os.path.join(self.tmp.name, 'd.json')
        self._save_draft_to(draft)
        self.assertTrue(os.path.exists(draft))

        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: True):
            self.app.clear_all()
        self.assertEqual(self.app.personnel.get(), '')
        self._load_draft_from(draft)

        self.assertEqual(self.app.personnel.get(), '王小明')
        self.assertEqual(self.app.year.get(), '115')
        self.assertEqual(self.app.text_fields['nature'].get('1.0', tk.END).strip(), '看到山羌\n第二行')
        self.assertEqual(self.app.photo_files[cat], [p])
        self.assertEqual(self.app.photo_captions[cat][p], '說明X')
        self.assertEqual(self.app.listboxes[cat].item('0', 'values')[1], '說明X')

    def test_draft_load_missing_photo_warns(self):
        cat, p = self._add_photo()
        draft = os.path.join(self.tmp.name, 'd.json')
        self._save_draft_to(draft)
        os.remove(p)
        self._load_draft_from(draft)
        self.assertEqual(self.app.photo_files[cat], [])
        self.assertEqual(self.app.listboxes[cat].count(), 0)
        self.assertTrue(any(k == 'info' and p in m for k, m in self.shown))

    def test_draft_invalid_file_rejected_and_data_kept(self):
        self.app.personnel.insert(0, '保留我')
        bad = os.path.join(self.tmp.name, 'bad.json')
        with open(bad, 'w', encoding='utf-8') as f:
            f.write('{ not json')
        self._load_draft_from(bad)
        self.assertEqual(self.shown[-1][0], 'err')
        wrong = os.path.join(self.tmp.name, 'wrong.json')
        with open(wrong, 'w', encoding='utf-8') as f:
            f.write('{"version": 99}')
        self._load_draft_from(wrong)
        self.assertEqual(self.shown[-1][0], 'err')
        self.assertEqual(self.app.personnel.get(), '保留我')

    def test_draft_load_declined_keeps_current(self):
        draft = os.path.join(self.tmp.name, 'd.json')
        self._save_draft_to(draft)
        self.app.personnel.insert(0, '現有內容')
        self._load_draft_from(draft, confirm=False)
        self.assertEqual(self.app.personnel.get(), '現有內容')

    # ---- 未存內容提示 / 自動存檔
    def _close_with(self, answer, save_ok=True):
        """模擬按下關閉，回傳 (是否關閉視窗, 提示是否出現)。"""
        asked = []
        destroyed = []

        def ask(*a, **k):
            asked.append(True)
            return answer

        with mock.patch.object(self.app.root, 'destroy', lambda: destroyed.append(True)),                 mock.patch.object(self.app, 'save_draft', lambda: save_ok),                 mock.patch.object(messagebox, 'askyesnocancel', ask):
            self.app.on_close()
        return bool(destroyed), bool(asked)

    def test_close_clean_no_prompt(self):
        closed, asked = self._close_with(None)
        self.assertTrue(closed)
        self.assertFalse(asked)

    def test_close_dirty_cancel_stays(self):
        self.app.personnel.insert(0, 'x')
        self.assertTrue(self.app.is_dirty())
        closed, asked = self._close_with(None)
        self.assertTrue(asked)
        self.assertFalse(closed)

    def test_close_dirty_discard_keeps_autosave_with_latest_content(self):
        self.app.personnel.insert(0, 'x')
        self.assertTrue(self.app.autosave_now())
        self.app.personnel.insert(tk.END, 'y')  # 自動存檔之後又改了內容
        closed, _ = self._close_with(False)
        self.assertTrue(closed)
        self.assertTrue(os.path.exists(self.autosave))  # 保留
        import json
        with open(self.autosave, encoding='utf-8') as f:
            self.assertEqual(json.load(f)['fields']['personnel'], 'xy')  # 離開前補存了最新內容

    def test_close_dirty_discard_with_blank_screen_removes_stale_autosave(self):
        self.app.personnel.insert(0, 'x')
        self.app.autosave_now()
        self.app._saved_signature = 'something else'  # 假裝上次存檔的內容和現在（空白）不同 -> 算有改動
        self.app.personnel.delete(0, tk.END)
        self.assertTrue(self.app.is_dirty())
        closed, _ = self._close_with(False)
        self.assertTrue(closed)
        self.assertFalse(os.path.exists(self.autosave))

    def test_close_clean_removes_autosave(self):
        with open(self.autosave, 'w', encoding='utf-8') as f:
            f.write('{}')
        closed, asked = self._close_with(None)
        self.assertTrue(closed)
        self.assertFalse(asked)
        self.assertFalse(os.path.exists(self.autosave))

    def test_autosave_interval_is_30_seconds(self):
        self.assertEqual(g.AUTOSAVE_INTERVAL_MS, 30_000)

    def test_close_dirty_save_ok_closes_but_save_cancelled_stays(self):
        self.app.personnel.insert(0, 'x')
        closed, _ = self._close_with(True, save_ok=True)
        self.assertTrue(closed)
        self.app.personnel.insert(0, 'y')
        closed, _ = self._close_with(True, save_ok=False)
        self.assertFalse(closed)

    def test_save_draft_marks_clean_and_removes_autosave(self):
        self.app.personnel.insert(0, 'x')
        self.app.autosave_now()
        self._save_draft_to(os.path.join(self.tmp.name, 'd.json'))
        self.assertFalse(self.app.is_dirty())
        self.assertFalse(os.path.exists(self.autosave))
        self.app.personnel.insert(0, 'z')
        self.assertTrue(self.app.is_dirty())

    def test_autosave_skips_blank_clean_and_unchanged(self):
        self.assertFalse(self.app.autosave_now())  # 空白
        self.app.personnel.insert(0, 'x')
        self.assertTrue(self.app.autosave_now())
        self.assertFalse(self.app.autosave_now())  # 沒有新變動
        self.app.personnel.insert(0, 'y')
        self.assertTrue(self.app.autosave_now())
        self.assertFalse(os.path.exists(self.autosave + '.tmp'))

    def test_autosave_restore_roundtrip(self):
        self.app.personnel.insert(0, '王小明')
        self.app.text_fields['nature'].insert('1.0', '山羌')
        self.app.autosave_now()
        # 模擬程式當機後重開：新的 app 實例
        self.root.destroy()
        self.root = tk.Tk()
        app2 = g.ReportGeneratorGUI(self.root, autosave_path=self.autosave, restore_on_start=False,
                                      settings_path=self.settings)
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: True):
            app2.offer_restore()
        self.assertEqual(app2.personnel.get(), '王小明')
        self.assertEqual(app2.text_fields['nature'].get('1.0', tk.END).strip(), '山羌')
        self.assertTrue(app2.is_dirty())  # 還原後仍視為未存
        self.app = app2

    def test_autosave_restore_declined_deletes_file(self):
        self.app.personnel.insert(0, 'x')
        self.app.autosave_now()
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: False):
            self.app.offer_restore()
        self.assertFalse(os.path.exists(self.autosave))
        self.assertEqual(self.app.personnel.get(), 'x')

    def test_corrupt_autosave_is_discarded_silently(self):
        with open(self.autosave, 'w', encoding='utf-8') as f:
            f.write('{ broken')
        self.app.offer_restore()
        self.assertFalse(os.path.exists(self.autosave))
        self.assertEqual(self.shown, [])

    def test_autosave_tick_reschedules(self):
        self.app.personnel.insert(0, 'x')
        with mock.patch.object(self.root, 'after') as after:
            self.app._autosave_tick()
        self.assertTrue(os.path.exists(self.autosave))
        after.assert_called_once()

    def test_draft_does_not_store_template_and_ignores_old_template_path(self):
        draft = os.path.join(self.tmp.name, 'd.json')
        self._save_draft_to(draft)
        import json
        with open(draft, encoding='utf-8') as f:
            data = json.load(f)
        self.assertNotIn('template_path', data)
        data['template_path'] = os.path.join(self.tmp.name, 'old.docx')  # 舊版草稿帶有這個欄位
        with open(draft, 'w', encoding='utf-8') as f:
            json.dump(data, f)
        self._load_draft_from(draft)
        self.assertEqual(self.app.template_path, TEMPLATE)

    def test_template_is_fixed_no_selection_ui(self):
        self.assertEqual(g.TEMPLATE_PATH, TEMPLATE)
        self.assertFalse(hasattr(self.app, 'select_template'))
        self.assertFalse(hasattr(self.app, 'template_label'))
        self.assertTrue(os.path.exists(g.TEMPLATE_PATH))

    def test_missing_template_blocks_generation(self):
        self.app.template_path = os.path.join(self.tmp.name, 'missing.docx')
        self.fill_dates()
        self.app.generate_report()
        self.assertEqual(self.shown[-1][0], 'err')
        self.assertIn('找不到模板', self.shown[-1][1])

    def test_output_same_as_template_rejected(self):
        import shutil
        tpl = os.path.join(self.tmp.name, 'tpl.docx')
        shutil.copy(TEMPLATE, tpl)
        self.app.template_path = tpl
        self.fill_dates()
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: tpl):
            self.app.generate_report()
        self.assertEqual(self.shown[-1][0], 'err')
        self.assertIn('模板', self.shown[-1][1])
        self.assertEqual(str(self.app.generate_btn['state']), 'normal')

    # ---- 編輯區項目符號 / 編號、字體
    def test_bullet_and_number_buttons_on_selection(self):
        w = self.app.text_fields['nature']
        w.insert('1.0', 'a\nb\nc')
        w.tag_add(tk.SEL, '1.0', '2.end')
        self.app.apply_list(w, g.toggle_bullets)
        self.assertEqual(w.get('1.0', 'end-1c'), '● a\n● b\nc')
        self.app.apply_list(w, g.toggle_numbers)
        self.assertEqual(w.get('1.0', 'end-1c'), '1. a\n2. b\nc')
        self.app.apply_list(w, g.toggle_numbers)  # 再按一次取消
        self.assertEqual(w.get('1.0', 'end-1c'), 'a\nb\nc')

    def test_buttons_apply_to_current_line_without_selection(self):
        w = self.app.text_fields['nature']
        w.insert('1.0', 'a\nb')
        w.mark_set(tk.INSERT, '2.1')
        self.app.apply_list(w, g.toggle_bullets)
        self.assertEqual(w.get('1.0', 'end-1c'), 'a\n● b')

    def _bullet_menu(self, widget):
        for w in self._all_widgets(self.root):
            if isinstance(w, ttk.Menubutton) and getattr(w, 'menu', None) is not None:
                # 同一個欄位的選單：以兄弟元件裡的文字區比對
                if widget.master in w.master.master.winfo_children():  # ScrolledText 的文字區外面還包了一層 Frame
                    return w.menu
        self.fail('找不到項目符號選單')

    def _all_widgets(self, parent):
        for child in parent.winfo_children():
            yield child
            yield from self._all_widgets(child)

    def test_bullet_dropdown_lists_all_symbols_and_applies_choice(self):
        w = self.app.text_fields['nature']
        menu = self._bullet_menu(w)
        labels = [menu.entrycget(i, 'label') for i in range(menu.index('end') + 1)
                  if menu.type(i) == 'command']
        self.assertEqual(len(labels), len(g.BULLET_CHOICES) + 2)  # 全部符號 + 編號 + 取消
        self.assertTrue(any('編號' in lbl and lbl.startswith('1.') for lbl in labels))
        self.assertTrue(any(lbl.startswith('◆') for lbl in labels))
        w.insert('1.0', 'a\nb')
        w.tag_add(tk.SEL, '1.0', 'end-1c')
        menu.invoke([i for i in range(menu.index('end') + 1) if menu.type(i) == 'command'
                     and menu.entrycget(i, 'label').startswith('◆')][0])
        self.assertEqual(w.get('1.0', 'end-1c'), '◆ a\n◆ b')
        menu.invoke(menu.index('end'))  # 最後一項：取消項目符號
        self.assertEqual(w.get('1.0', 'end-1c'), 'a\nb')

    def test_numbering_is_in_the_same_dropdown_and_no_separate_button(self):
        w = self.app.text_fields['nature']
        menu = self._bullet_menu(w)
        idx = [i for i in range(menu.index('end') + 1) if menu.type(i) == 'command'
               and menu.entrycget(i, 'label').startswith('1.')][0]
        w.insert('1.0', 'a\nb')
        w.tag_add(tk.SEL, '1.0', 'end-1c')
        menu.invoke(idx)
        self.assertEqual(w.get('1.0', 'end-1c'), '1. a\n2. b')
        menu.invoke(idx)  # 再選一次取消
        self.assertEqual(w.get('1.0', 'end-1c'), 'a\nb')
        buttons = [c.cget('text') for c in self._all_widgets(self.root) if isinstance(c, ttk.Button)]
        self.assertNotIn('1. 編號', buttons)

    # ---- 預設照片說明
    def test_roc_date_parts(self):
        import datetime as dt
        self.assertEqual(g.roc_date_parts(dt.datetime(2026, 10, 14)), ('115', '10', '14'))
        self.assertEqual(g.roc_date_parts(dt.datetime(2027, 1, 5)), ('116', '1', '5'))
        self.assertEqual(g.roc_date_parts(dt.datetime(2012, 12, 31)), ('101', '12', '31'))

    def test_new_form_gets_todays_roc_date(self):
        import datetime as dt
        self.root.destroy()
        self.root = tk.Tk()
        with mock.patch.object(g, 'current_date', lambda: dt.datetime(2026, 10, 14, 9, 30)):
            self.app = g.ReportGeneratorGUI(self.root, autosave_path=self.autosave, restore_on_start=False,
                                      settings_path=self.settings)
        self.assertEqual((self.app.year.get(), self.app.month.get(), self.app.day.get()), ('115', '10', '14'))
        labels = [w.cget('text') for w in self._all_widgets(self.root) if isinstance(w, ttk.Label)]
        self.assertIn('1. 填表日期:', labels)
        self.assertFalse(any('年月日' in t for t in labels))

    def test_auto_date_alone_is_not_unsaved_content(self):
        self.assertTrue(self.app.year.get())
        self.assertFalse(self.app.is_dirty())
        self.assertFalse(self.app.autosave_now())  # 只有自動日期不算內容，不產生自動存檔
        closed, asked = self._close_with(None)
        self.assertTrue(closed)
        self.assertFalse(asked)

    def test_autosave_starts_once_real_content_is_entered(self):
        self.app.personnel.insert(0, 'x')
        self.assertTrue(self.app.is_dirty())
        self.assertTrue(self.app.autosave_now())

    def test_clear_all_regenerates_date(self):
        import datetime as dt
        for name in g.REPORT_DATE_FIELDS:
            getattr(self.app, name).delete(0, tk.END)
        self.app.year.insert(0, '100')
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: True), \
                mock.patch.object(g, 'current_date', lambda: dt.datetime(2026, 10, 20)):
            self.app.clear_all()
        self.assertEqual((self.app.year.get(), self.app.month.get(), self.app.day.get()), ('115', '10', '20'))

    def test_loading_draft_keeps_its_original_fill_date(self):
        import datetime as dt
        draft = os.path.join(self.tmp.name, 'd.json')
        for name, value in zip(g.REPORT_DATE_FIELDS, ('115', '9', '30')):  # 第一次填表是 9/30
            getattr(self.app, name).delete(0, tk.END)
            getattr(self.app, name).insert(0, value)
        self._save_draft_to(draft)
        with mock.patch.object(g, 'current_date', lambda: dt.datetime(2026, 10, 14)):  # 今天已經是 10/14
            self._load_draft_from(draft)
        self.assertEqual((self.app.year.get(), self.app.month.get(), self.app.day.get()), ('115', '9', '30'))

    def test_draft_without_date_gets_todays_date(self):
        import json, datetime as dt
        draft = os.path.join(self.tmp.name, 'nodate.json')
        self._save_draft_to(draft)
        with open(draft, encoding='utf-8') as f:
            data = json.load(f)
        for name in g.REPORT_DATE_FIELDS:
            data['fields'][name] = ''
        with open(draft, 'w', encoding='utf-8') as f:
            json.dump(data, f)
        with mock.patch.object(g, 'current_date', lambda: dt.datetime(2026, 11, 2)):
            self._load_draft_from(draft)
        self.assertEqual((self.app.year.get(), self.app.month.get(), self.app.day.get()), ('115', '11', '2'))

    def test_auto_date_flows_into_report(self):
        import datetime as dt
        self.root.destroy()
        self.root = tk.Tk()
        with mock.patch.object(g, 'current_date', lambda: dt.datetime(2026, 10, 14)):
            self.app = g.ReportGeneratorGUI(self.root, autosave_path=self.autosave, restore_on_start=False,
                                      settings_path=self.settings)
        for e, v in [(self.app.start_year, '115'), (self.app.start_month, '10'), (self.app.start_day, '8'),
                     (self.app.end_year, '115'), (self.app.end_month, '10'), (self.app.end_day, '11')]:
            e.insert(0, v)
        data, err = self.app.collect_data()
        self.assertIsNone(err)
        out = os.path.join(self.tmp.name, 'o.docx')
        g.build_report(TEMPLATE, data, {c: [] for c in self.app.photo_files}, {}, out)
        text = '\n'.join(c.text for c in unique_cells(Document(out)))
        self.assertIn('115年10月14日', text)

    # ---- 清除前的確認
    def _clear_with_answer(self, answer):
        calls = []

        def ask(title, message, **kwargs):
            calls.append((title, message, kwargs))
            return answer

        with mock.patch.object(messagebox, 'askyesno', ask):
            self.app.clear_all()
        return calls

    def test_clear_asks_for_confirmation_with_safe_default(self):
        self.app.personnel.insert(0, '保留')
        calls = self._clear_with_answer(False)
        self.assertEqual(len(calls), 1)
        _, _, kwargs = calls[0]
        self.assertEqual(kwargs.get('default'), 'no')  # 預設「否」，按 Enter 不會誤清除
        self.assertEqual(kwargs.get('icon'), 'warning')
        self.assertEqual(self.app.personnel.get(), '保留')  # 選否：什麼都不動

    def test_clear_confirmed_clears_everything(self):
        self.app.personnel.insert(0, '清掉')
        cat, lb, _ = self._photos()
        self._clear_with_answer(True)
        self.assertEqual(self.app.personnel.get(), '')
        self.assertEqual(lb.count(), 0)

    def test_clear_warning_mentions_scope_and_irreversibility(self):
        _, message, _ = self._clear_with_answer(False)[0]
        for word in ('欄位', '照片', '說明', '無法復原'):
            self.assertIn(word, message)

    def test_clear_warns_extra_when_there_is_unsaved_content(self):
        _, clean_message, _ = self._clear_with_answer(False)[0]
        self.assertNotIn('尚未儲存', clean_message)
        self.app.personnel.insert(0, 'x')
        _, dirty_message, _ = self._clear_with_answer(False)[0]
        self.assertIn('尚未儲存', dirty_message)
        self.assertIn('儲存草稿', dirty_message)
        # 存成草稿之後，又回到一般確認
        self._save_draft_to(os.path.join(self.tmp.name, 'd.json'))
        _, saved_message, _ = self._clear_with_answer(False)[0]
        self.assertNotIn('尚未儲存', saved_message)

    # ---- 外觀主題
    @staticmethod
    def _luminance(color):
        def channel(v):
            v /= 255
            return v / 12.92 if v <= 0.03928 else ((v + 0.055) / 1.055) ** 2.4
        r, gr, b = (int(color[i:i + 2], 16) for i in (1, 3, 5))
        return 0.2126 * channel(r) + 0.7152 * channel(gr) + 0.0722 * channel(b)

    def _contrast(self, fg, bg):
        hi, lo = sorted((self._luminance(fg), self._luminance(bg)), reverse=True)
        return (hi + 0.05) / (lo + 0.05)

    def test_themes_available(self):
        self.assertEqual(set(g.THEMES), {'柔和淺色', '色弱友善', '高對比'})
        self.assertIn(g.DEFAULT_THEME, g.THEMES)
        self.assertEqual(tuple(self.app.theme_box.cget('values')), tuple(g.THEMES))
        self.assertEqual(self.app.theme_var.get(), g.DEFAULT_THEME)

    def test_every_theme_text_is_readable_wcag_aa(self):
        for name, p in g.THEMES.items():
            pairs = [('primary', p['primary_fg'], p['primary']), ('danger', p['danger_fg'], p['danger']),
                     ('heading', p['heading_fg'], p['bg']), ('field label', p['field_fg'], p['bg']),
                     ('active tab', p['tab_active_fg'], p['tab_active']), ('idle tab', p['tab_idle_fg'], p['tab_idle']),
                     ('primary hover', p['primary_fg'], p['primary_active']), ('danger hover', p['danger_fg'], p['danger_active'])]
            for label, fg, bg in pairs:
                with self.subTest(theme=name, item=label):
                    self.assertGreaterEqual(self._contrast(fg, bg), 4.5)

    def test_default_theme_colors_are_light(self):
        p = g.THEMES[g.DEFAULT_THEME]
        for key in ('bg', 'tab_idle', 'tab_active', 'primary', 'danger'):
            with self.subTest(key=key):
                self.assertGreater(self._luminance(p[key]), 0.5)  # 淺色，不是濃重的底色

    def test_colorblind_theme_avoids_red_green_pair(self):
        p = g.THEMES['色弱友善']
        pr = [int(p['primary_border'][i:i + 2], 16) for i in (1, 3, 5)]
        dg = [int(p['danger_border'][i:i + 2], 16) for i in (1, 3, 5)]
        self.assertGreater(pr[2], pr[0])   # 主要按鈕：藍色（B 大於 R）
        self.assertGreater(pr[2], pr[1])
        self.assertGreater(dg[0], dg[2])   # 危險按鈕：橘色（R 大於 B，且不是綠色）
        self.assertGreater(dg[0], dg[1])
        self.assertGreater(dg[1], dg[2])
        self.assertNotEqual(p['primary'], p['danger'])

    def test_apply_theme_changes_styles_and_widgets_and_is_remembered(self):
        import json
        self.app.apply_theme('色弱友善')
        p = g.THEMES['色弱友善']
        self.assertEqual(style_lookup(self.root, 'Primary.TButton', 'background'), p['primary'])
        self.assertEqual(style_lookup(self.root, 'Danger.TButton', 'background'), p['danger'])
        self.assertEqual(style_lookup(self.root, 'TNotebook.Tab', 'background', ['selected']), p['tab_active'])
        self.assertEqual(str(self.root.cget('background')), p['bg'])
        canvases = [w for w in self._all_widgets(self.root) if isinstance(w, tk.Canvas)]
        self.assertTrue(canvases)
        self.assertTrue(all(str(c.cget('background')) == p['bg'] for c in canvases))
        with open(self.settings, encoding='utf-8') as f:
            self.assertEqual(json.load(f), {'theme': '色弱友善'})
        # 重新開啟時沿用上次選的主題
        self.root.destroy()
        self.root = tk.Tk()
        self.app = g.ReportGeneratorGUI(self.root, autosave_path=self.autosave, restore_on_start=False,
                                        settings_path=self.settings)
        self.assertEqual(self.app.theme_name, '色弱友善')
        self.assertEqual(self.app.theme_var.get(), '色弱友善')
        self.assertEqual(style_lookup(self.root, 'Primary.TButton', 'background'), p['primary'])

    def test_theme_combobox_selection_applies_theme(self):
        self.app.theme_var.set('高對比')
        self.app.theme_box.event_generate('<<ComboboxSelected>>')
        self.root.update()
        self.assertEqual(self.app.theme_name, '高對比')
        self.assertEqual(style_lookup(self.root, 'Danger.TButton', 'background'), g.THEMES['高對比']['danger'])

    def test_unknown_or_broken_settings_fall_back_to_default(self):
        for content in ('{ broken', '{"theme": "不存在"}', '[]'):
            with open(self.settings, 'w', encoding='utf-8') as f:
                f.write(content)
            self.assertEqual(self.app._load_theme_setting(), g.DEFAULT_THEME, content)
        os.remove(self.settings)
        self.assertEqual(self.app._load_theme_setting(), g.DEFAULT_THEME)
        self.app.apply_theme('不存在')  # 無效名稱不處理
        self.assertEqual(self.app.theme_name, g.DEFAULT_THEME)

    def test_switching_theme_does_not_make_form_dirty(self):
        self.app.apply_theme('高對比')
        self.assertFalse(self.app.is_dirty())

    # ---- 日期欄位間距
    def test_date_row_spacing_is_symmetric_and_widths_equal(self):
        self.root.update()
        for prefix in ('', 'start_', 'end_'):
            self.assertEqual(self.app.__dict__[f'{prefix}year'].winfo_width(), self.app.year.winfo_width())
            self.assertEqual(self.app.__dict__[f'{prefix}month'].winfo_width(), self.app.month.winfo_width())
            self.assertEqual(self.app.__dict__[f'{prefix}day'].winfo_width(), self.app.day.winfo_width())
        frame = self.app.start_year.master
        for label in [w for w in frame.winfo_children() if isinstance(w, ttk.Label)]:
            padx = label.pack_info()['padx']
            left, right = (padx if isinstance(padx, tuple) else (padx, padx))
            with self.subTest(label=label.cget('text')):
                self.assertEqual(int(left), int(right))  # 單位字前後間距相同
        texts = [w.cget('text') for w in frame.winfo_children() if isinstance(w, ttk.Label)]
        self.assertEqual(texts, ['年', '月', '日', '至', '年', '月', '日'])  # 「至」是獨立的一個字，不再和「日」黏在一起

    # ---- 照片分頁分組與視窗高度
    def test_all_photo_categories_are_in_one_tab_as_sub_tabs_in_word_order(self):
        names = [self.app.photo_notebook.tab(tab, 'text') for tab in self.app.photo_notebook.tabs()]
        expected = [g.PHOTO_SHORT_NAMES[c] for c in g.PHOTO_CATEGORIES]
        self.assertEqual(len(names), 6)
        for name, short in zip(names, expected):
            self.assertTrue(name.startswith(short), (name, short))
        main = [self.app.notebook.tab(tab, 'text') for tab in self.app.notebook.tabs()]
        self.assertEqual(sum(t.startswith('服勤照片') for t in main), 1)  # 只剩一個「服勤照片」分頁
        for category in g.PHOTO_CATEGORIES:  # 每個分類的表格都在這一個分頁裡
            self.assertTrue(str(self.app.listboxes[category]).startswith(self.app._tab_by_title(g.TAB_PHOTOS)))

    def test_word_order_matches_sub_tab_order(self):
        self.assertEqual(list(self.app.photo_files), g.PHOTO_CATEGORIES)
        self.assertEqual(g.PHOTO_CATEGORIES.index('自然資源狀況'), 2)  # 仍在「設施」之後、「環境清理」之前
        self.assertEqual(set(g.PHOTO_SHORT_NAMES), set(g.PHOTO_CATEGORIES))
        self.assertFalse(hasattr(g, 'PHOTO_CATEGORIES_1'))
        self.assertFalse(hasattr(g, 'PHOTO_TAB_TITLES'))

    def test_report_keeps_original_chapter_order(self):
        import glob
        self.fill_dates()
        paths = {}
        for i, cat in enumerate(g.PHOTO_CATEGORIES):
            p = os.path.join(self.tmp.name, f'c{i}.jpg')
            Image.new('RGB', (400, 300), 'red').save(p)
            self.app.photo_files[cat].append(p)
            paths[cat] = p
        data, err = self.app.collect_data()
        self.assertIsNone(err)
        out = os.path.join(self.tmp.name, 'o.docx')
        g.build_report(TEMPLATE, data, {c: list(v) for c, v in self.app.photo_files.items()}, {}, out)
        heads = []
        for table in Document(out).tables[1:]:
            for row in table.rows:
                text = row.cells[0].text.strip()
                if text in g.PHOTO_CATEGORIES and (not heads or heads[-1] != text):
                    heads.append(text)
        self.assertEqual(heads, g.PHOTO_CATEGORIES)

    def test_draft_roundtrip_with_nature_category(self):
        cat = '自然資源狀況'
        table = self.app.listboxes[cat]
        p = os.path.join(self.tmp.name, 'n.jpg')
        Image.new('RGB', (30, 20)).save(p)
        self.app.photo_files[cat].append(p)
        self.app._render_list(cat)
        draft = os.path.join(self.tmp.name, 'd.json')
        self._save_draft_to(draft)
        self.app.clear_photos(cat, table)
        self._load_draft_from(draft)
        self.assertEqual(self.app.photo_files[cat], [p])
        self.assertEqual(table.count(), 1)

    def test_scrollbar_appears_only_when_content_does_not_fit(self):
        nb = self.app.notebook
        canvas, inner, scrollbar = self.app._scroll_tabs[g.TAB_PHOTOS]
        nb.select(self.app._tab_by_title(g.TAB_PHOTOS))
        tables = list(self.app.listboxes.values())
        original = [int(t.cget('height')) for t in tables]
        try:
            for table in tables:  # 內容變少 -> 放得下 -> 隱藏
                table.configure(height=3)
            self.app._refit_scroll_tabs()
            self.root.update()
            self.assertLessEqual(inner.winfo_reqheight(), canvas.winfo_height())
            self.assertEqual(scrollbar.winfo_manager(), '', '內容放得下時要隱藏捲動軸')
            for table in tables:  # 內容變多 -> 放不下 -> 顯示
                table.configure(height=60)
            self.app._refit_scroll_tabs()
            self.root.update()
            self.assertGreater(inner.winfo_reqheight(), canvas.winfo_height())
            self.assertNotEqual(scrollbar.winfo_manager(), '', '內容放不下時要顯示捲動軸')
        finally:
            for table, rows in zip(tables, original):
                table.configure(height=rows)

    def test_short_screen_falls_back_to_compact_photo_lists(self):
        """螢幕不夠高、視窗加高也放不下時，改把照片表格縮成較少的列數，並把視窗限制在螢幕高度內。"""
        for table in self.app.listboxes.values():
            table.configure(height=g.PHOTO_TABLE_ROWS)
        self.root.geometry(f'{self.root.winfo_width()}x560')
        self.root.update()
        before = self.app._tabs_extra_height()
        self.assertGreater(before, 0)
        with mock.patch.object(self.root, 'winfo_screenheight', lambda: 700):  # 可用高度約 610
            self.app._fit_window_height()
        self.root.update()
        self.assertTrue(all(int(t.cget('height')) == g.PHOTO_TABLE_ROWS_COMPACT for t in self.app.listboxes.values()))
        self.assertLessEqual(self.root.winfo_height(), 700 - 90)
        self.assertLess(self.app._tabs_extra_height(), before)  # 比原本好（縮小表格 + 加高視窗）

    def test_lists_stay_normal_size_when_the_window_can_hold_everything(self):
        """視窗放得下所有分頁時，不需要縮小照片表格。"""
        for t in self.app.listboxes.values():
            t.configure(height=g.PHOTO_TABLE_ROWS)
        with mock.patch.object(self.app, '_tabs_extra_height', lambda: 0):
            self.app._fit_window_height()
        self.assertTrue(all(int(t.cget('height')) == g.PHOTO_TABLE_ROWS for t in self.app.listboxes.values()))

    def test_photo_tab_needs_no_horizontal_room_beyond_window(self):
        canvas, inner, _ = self.app._scroll_tabs[g.TAB_PHOTOS]
        self.app.notebook.select(self.app._tab_by_title(g.TAB_PHOTOS))
        self.root.update()
        self.assertLessEqual(inner.winfo_reqwidth(), canvas.winfo_width())

    def _linespace(self, widget):
        import tkinter.font as tkfont
        font = widget.cget('font') or style_lookup(self.root, str(widget.cget('style')), 'font')
        return tkfont.Font(root=self.root, font=font).metrics('linespace')

    def test_long_section_titles_stay_on_one_line(self):
        """三、四、五、八 這些很長的標題，在預設視窗寬度下要一行顯示（不能被固定寬度截斷換行）。"""
        nb = self.app.notebook
        long_titles = ('三、', '四、', '五、', '八、')
        found = []
        for tab_index in (3, 4):  # 自然資源/環境維護/遊客管理、遊客意見/心得建議/疑問解答
            nb.select(tab_index)
            self.root.update()
            tab = nb.nametowidget(nb.tabs()[tab_index])
            for w in self._all_widgets(tab):
                if isinstance(w, ttk.Label) and w.cget('text').startswith(long_titles):
                    found.append(w)
                    with self.subTest(title=w.cget('text')[:8]):
                        self.assertLessEqual(w.winfo_reqheight(), self._linespace(w) + 6,
                                             '標題被換成兩行了：' + w.cget('text')[:12])
        self.assertEqual(len(found), 4)

    def test_section_title_wraps_only_when_window_is_too_narrow_and_dropdown_stays_visible(self):
        nb = self.app.notebook
        nb.select(3)
        self.root.update()
        tab = nb.nametowidget(nb.tabs()[3])
        title = next(w for w in self._all_widgets(tab)
                     if isinstance(w, ttk.Label) and w.cget('text').startswith('三、'))
        button = next(w for w in title.master.winfo_children() if isinstance(w, ttk.Menubutton))
        original = self.root.winfo_height()
        try:
            self.root.geometry(f'620x{original}')
            self.root.update()
            self.assertGreater(title.winfo_reqheight(), self._linespace(title) + 6)  # 太窄時才換行
            header = title.master
            self.assertLessEqual(button.winfo_x() + button.winfo_width(), header.winfo_width())  # 下拉按鈕沒被擠出去
        finally:
            self.root.geometry(f'1050x{original}')
            self.root.update()

    # ---- 文字清晰度（對比）與外觀選單置中
    def test_text_colors_have_high_contrast_in_every_theme(self):
        for name, p in g.THEMES.items():
            for label, fg, bg in (('body text', p['text'], '#ffffff'), ('body text on form', p['text'], p['bg']),
                                  ('hint text', p['muted'], p['bg'])):
                with self.subTest(theme=name, item=label):
                    self.assertGreaterEqual(self._contrast(fg, bg), 7.0)  # WCAG AAA 等級

    def test_default_text_and_hint_use_theme_colors_not_pale_gray(self):
        p = g.THEMES[g.DEFAULT_THEME]
        self.assertEqual(style_lookup(self.root, 'TLabel', 'foreground'), p['text'])
        self.assertEqual(style_lookup(self.root, 'TEntry', 'foreground'), p['text'])
        self.assertEqual(style_lookup(self.root, 'Hint.TLabel', 'foreground'), p['muted'])
        for w in self._all_widgets(self.root):
            if isinstance(w, ttk.Label) and w.cget('text') == g.PHOTO_TAB_HINT:
                self.assertEqual(str(w.cget('style')), 'Hint.TLabel')
            if isinstance(w, (tk.Text, tk.Listbox)):
                self.assertEqual(str(w.cget('foreground')), p['text'])

    def test_theme_change_updates_native_text_colors(self):
        self.app.apply_theme('高對比')
        for w in self._all_widgets(self.root):
            if isinstance(w, (tk.Text, tk.Listbox)):
                self.assertEqual(str(w.cget('foreground')), g.THEMES['高對比']['text'])

    def test_theme_dropdown_text_and_options_are_centered(self):
        self.assertEqual(str(self.app.theme_box.cget('justify')), 'center')  # 選取後顯示的文字置中
        popdown = self.root.tk.eval(f'ttk::combobox::PopdownWindow {self.app.theme_box}')
        self.assertEqual(str(self.root.tk.call(f'{popdown}.f.l', 'cget', '-justify')), 'center')  # 展開的選項置中
        self.assertEqual(tuple(self.app.theme_box.cget('values')), tuple(g.THEMES))

    def test_dpi_awareness_helper_is_safe(self):
        with mock.patch.object(g.sys, 'platform', 'linux'):
            self.assertFalse(g.enable_dpi_awareness())
        self.assertIn(g.enable_dpi_awareness.__name__, dir(g))

    def test_list_dropdown_is_one_point_larger(self):
        import tkinter.font as tkfont
        body = abs(tkfont.nametofont('TkDefaultFont', root=self.root).cget('size'))
        self.assertEqual(g.LIST_MENU_EXTRA_PT, 1)
        # 按鈕本身
        button_font = tkfont.Font(root=self.root, font=style_lookup(self.root, 'List.TMenubutton', 'font'))
        self.assertEqual(abs(button_font.cget('size')), body + 1)
        # 展開的選單項目
        for w in self._all_widgets(self.root):
            if isinstance(w, ttk.Menubutton) and getattr(w, 'menu', None) is not None:
                self.assertEqual(str(w.cget('style')), 'List.TMenubutton')
                menu_font = tkfont.Font(root=self.root, font=w.menu.cget('font'))
                self.assertEqual(abs(menu_font.cget('size')), body + 1)
                self.assertEqual(menu_font.cget('family'), g.FONT_FAMILY)
                found = True
        self.assertTrue(found)

    # ---- 沒有「預設字樣」：沒填說明就是空白
    def test_no_default_caption_setting_anymore(self):
        self.assertFalse(hasattr(self.app, 'default_caption'))
        self.assertFalse(hasattr(g, 'DEFAULT_CAPTION'))
        texts = [w.cget('text') for w in self._all_widgets(self.root) if isinstance(w, ttk.Label)]
        self.assertFalse(any('預設字樣' in x for x in texts))
        self.assertNotIn('default_caption', self.app.build_draft())

    def test_photos_without_caption_have_blank_caption_cells(self):
        self.fill_dates()
        cat, lb, (a, b) = self._photos(names=('a.jpg', 'b.jpg'))
        self.app.photo_captions[cat][a] = '有說明'
        data, err = self.app.collect_data()
        self.assertIsNone(err)
        self.assertNotIn('default_caption', data)
        out = os.path.join(self.tmp.name, 'o.docx')
        g.build_report(TEMPLATE, data, {c: list(p) for c, p in self.app.photo_files.items()},
                       {c: dict(v) for c, v in self.app.photo_captions.items()}, out)
        table = Document(out).tables[-1]
        caption_row = table.rows[2]  # 分類標題列、照片列之後的第一個說明列
        texts = [c.text for c in caption_row.cells]
        self.assertEqual(texts, ['有說明', ''])  # 沒填的就是空白，沒有任何佔位字樣
        self.assertNotIn('[照片說明]', '\n'.join(c.text for c in unique_cells(Document(out))))

    def test_old_draft_with_default_caption_key_still_loads(self):
        import json
        draft = os.path.join(self.tmp.name, 'old.json')
        self._save_draft_to(draft)
        with open(draft, encoding='utf-8') as f:
            data = json.load(f)
        data['default_caption'] = '[照片說明]'  # 舊版草稿留下的欄位
        with open(draft, 'w', encoding='utf-8') as f:
            json.dump(data, f)
        self._load_draft_from(draft)  # 不會出錯，也沒有任何效果
        self.assertFalse(hasattr(self.app, 'default_caption'))

    def test_photo_tab_has_one_hint_line(self):
        hints = [w for w in self._all_widgets(self.root)
                 if isinstance(w, ttk.Label) and w.cget('text') == g.PHOTO_TAB_HINT]
        self.assertEqual(len(hints), 1)
        self.assertIn('在預覽下方輸入說明', g.PHOTO_TAB_HINT)
        self.assertIn('Enter', g.PHOTO_TAB_HINT)

    # ---- 視窗標題列與分頁之間不要有多餘空白
    def test_no_extra_blank_space_above_tabs(self):
        self.root.update()
        nb = self.app.notebook
        geometry = g.tab_geometry(nb, nb.index(nb.select()))
        self.assertIsNotNone(geometry)
        top = nb.winfo_rooty() + geometry[2] - self.root.winfo_rooty()
        self.assertLessEqual(top, 12, f'分頁上方留白 {top}px 太多')
        self.assertGreaterEqual(top, 0)

    # ---- 分頁改名
    def test_tab_names(self):
        titles = [self.app.notebook.tab(tab, 'text').replace(' ●', '') for tab in self.app.notebook.tabs()]
        self.assertEqual(titles, ['基本資料', '道路/步道', '設施', '自然資源/環境維護/遊客管理',
                                  '遊客意見/心得建議/疑問解答', '服勤照片'])
        self.assertEqual((g.TAB_ROAD, g.TAB_NATURE, g.TAB_FEEDBACK, g.TAB_PHOTOS),
                         ('道路/步道', '自然資源/環境維護/遊客管理', '遊客意見/心得建議/疑問解答', '服勤照片'))

    def test_all_tabs_fit_in_the_window_width(self):
        self.root.update()
        nb = self.app.notebook
        geometry = g.tab_geometry(nb, len(nb.tabs()) - 1)
        self.assertIsNotNone(geometry)
        self.assertLessEqual(geometry[1], nb.winfo_width(), '最後一個分頁被切掉了（分頁名稱太長、視窗太窄）')
        # 每個分頁都看得到（量得到位置）
        for i in range(len(nb.tabs())):
            self.assertIsNotNone(g.tab_geometry(nb, i), f'分頁 {i} 被擠出視窗外')

    # ---- 所有分頁預設都不用捲動、按鈕列間距
    def test_every_tab_fits_without_scrollbar_when_the_screen_allows(self):
        nb = self.app.notebook
        if self.root.winfo_height() >= self.root.winfo_screenheight() - 100:
            self.skipTest('螢幕高度不足，視窗已受限於螢幕')
        for title, (canvas, inner, scrollbar) in self.app._scroll_tabs.items():
            nb.select(self.app._tab_by_title(title))
            self.root.update()
            with self.subTest(tab=title):
                self.assertLessEqual(inner.winfo_reqheight(), canvas.winfo_height(), f'{title} 內容放不下')
                self.assertEqual(scrollbar.winfo_manager(), '', f'{title} 出現了捲動軸')

    def test_facility_tab_has_no_scrollbar_by_default(self):
        nb = self.app.notebook
        nb.select(self.app._tab_by_title('設施'))
        self.root.update()
        canvas, inner, scrollbar = self.app._scroll_tabs['設施']
        if self.root.winfo_height() >= self.root.winfo_screenheight() - 100:
            self.skipTest('螢幕高度不足，視窗已受限於螢幕')
        self.assertEqual(scrollbar.winfo_manager(), '')

    def test_text_areas_shrink_only_as_a_last_resort(self):
        """螢幕非常矮時，才會先縮清單、再縮文字區；平常文字區維持原本高度。"""
        base = {k: int(w.cget('height')) for k, w in self.app.text_fields.items()}
        self.assertTrue(all(v >= 5 for v in base.values()) or self.root.winfo_screenheight() < 900)

    def test_small_gap_between_notebook_and_buttons(self):
        self.root.update()
        nb = self.app.notebook
        buttons = self.app.generate_btn.master
        gap = buttons.winfo_rooty() - (nb.winfo_rooty() + nb.winfo_height())
        self.assertLessEqual(gap, 16, f'分頁框與按鈕列之間有 {gap}px 空白')
        self.assertGreaterEqual(gap, 0)

    def test_add_photos_button_label(self):
        labels = [w.cget('text') for w in self._all_widgets(self.root) if isinstance(w, ttk.Button)]
        self.assertEqual(labels.count('加入照片'), len(g.PHOTO_CATEGORIES))  # 每個分類一個
        self.assertNotIn('選擇照片', labels)
        self.assertEqual(labels.count('移除所選照片'), len(g.PHOTO_CATEGORIES))
        self.assertEqual(labels.count('移除全部照片'), len(g.PHOTO_CATEGORIES))
        self.assertNotIn('移除選取', labels)
        self.assertNotIn('清除全部', labels)
        self.assertNotIn('清除', labels)

    def test_add_photos_dialog_title(self):
        titles = []
        cat = g.PHOTO_CATEGORIES[0]
        with mock.patch.object(filedialog, 'askopenfilenames', lambda **k: titles.append(k['title']) or ()):
            self.app.select_photos(cat, self.app.listboxes[cat])
        self.assertEqual(titles, [f'加入照片 - {cat}'])

    # ---- 啟動時量測分頁高度：視窗保持可見，讓使用者知道程式在準備（不是當掉）
    def test_startup_measuring_keeps_window_visible_and_shows_progress(self):
        nb = self.app.notebook
        original = nb.index(nb.select())
        seen = []
        real_select = nb.select

        def spy(*args, **kwargs):
            if args or kwargs:  # 只記錄「切換分頁」的呼叫
                seen.append((float(self.root.attributes('-alpha')), bool(self.root.winfo_viewable()),
                             self.app.status_label.cget('text')))
            return real_select(*args, **kwargs)

        with mock.patch.object(nb, 'select', spy):
            self.app._fit_window_height()
        self.assertTrue(seen, '量測時應該有切換分頁')
        self.assertTrue(all(alpha == 1.0 for alpha, _, _ in seen), '量測期間視窗不可以是透明的')
        self.assertTrue(all(viewable for _, viewable, _ in seen), '量測期間視窗要看得到')
        self.assertTrue(all('準備' in text for _, _, text in seen), '量測期間狀態列要說明正在準備')
        self.assertEqual(self.app.status_label.cget('text'), '準備就緒')   # 量完恢復
        self.assertEqual(nb.index(nb.select()), original)                 # 回到原本的分頁

    def test_status_restored_even_if_measuring_fails(self):
        with mock.patch.object(self.app, '_measure_and_fit', side_effect=RuntimeError('boom')):
            with self.assertRaises(RuntimeError):
                self.app._fit_window_height()
        self.assertEqual(self.app.status_label.cget('text'), '準備就緒')

    def test_no_transparency_helpers_left(self):
        self.assertFalse(hasattr(self.app, '_set_visible'))

    def test_ctrl_e_exits_like_the_exit_button(self):
        calls = []
        with mock.patch.object(self.app, 'on_close', lambda: calls.append('close')):
            self.app._bind_shortcuts()
            self._press(self.app.personnel, 'e')
            self._press(self.app.personnel, 'E')
        self.assertEqual(calls, ['close', 'close'])

    def test_exit_button_shows_shortcut_letter(self):
        labels = {w.cget('text') for w in self._all_widgets(self.root) if isinstance(w, ttk.Button)}
        self.assertIn('離開(E)', labels)
        self.assertNotIn('離開', labels)

    def test_window_is_visible_and_on_first_tab_after_startup(self):
        self.assertEqual(float(self.root.attributes('-alpha')), 1.0)
        self.assertEqual(self.app.notebook.index(self.app.notebook.select()), 0)

    def _press(self, widget, keysym):
        self.root.deiconify()
        self.root.update()
        self.root.focus_force()
        widget.focus_set()
        self.root.update()
        widget.event_generate(f'<Control-{keysym}>')
        self.root.update()

    def test_buttons_show_shortcut_letters(self):
        labels = {w.cget('text') for w in self._all_widgets(self.root) if isinstance(w, ttk.Button)}
        self.assertTrue({'載入草稿(L)', '儲存草稿(S)', '產生報告(G)'} <= labels)
        self.assertTrue({'載入草稿', '儲存草稿', '產生報告'}.isdisjoint(labels))

    def test_ctrl_s_saves_draft_ctrl_l_loads_ctrl_g_generates(self):
        calls = []
        with mock.patch.object(self.app, 'save_draft', lambda: calls.append('save')),                 mock.patch.object(self.app, 'load_draft', lambda: calls.append('load')),                 mock.patch.object(self.app, 'generate_report', lambda: calls.append('generate')):
            self.app._bind_shortcuts()  # 重新綁定，讓綁定使用被替換過的方法
            self._press(self.app.personnel, 's')
            self._press(self.app.personnel, 'l')
            self._press(self.app.personnel, 'g')
        self.assertEqual(calls, ['save', 'load', 'generate'])

    def test_shortcuts_work_with_capital_letters_and_from_text_areas(self):
        calls = []
        with mock.patch.object(self.app, 'save_draft', lambda: calls.append('save')),                 mock.patch.object(self.app, 'load_draft', lambda: calls.append('load')),                 mock.patch.object(self.app, 'generate_report', lambda: calls.append('generate')):
            self.app._bind_shortcuts()
            text = self.app.text_fields['nature']
            self.app.notebook.select(self.app._tab_by_title(g.TAB_NATURE))
            self._press(text, 'S')
            self._press(text, 'L')
            self._press(text, 'G')
        self.assertEqual(calls, ['save', 'load', 'generate'])
        self.assertEqual(text.get('1.0', 'end-1c'), '')  # 快捷鍵不會在文字框裡打出字元

    def test_ctrl_g_ignored_while_generating(self):
        calls = []
        self.app.generate_btn.config(state=tk.DISABLED)
        with mock.patch.object(self.app, 'generate_report', lambda: calls.append('generate')):
            self.app._bind_shortcuts()
            self._press(self.app.personnel, 'g')
        self.assertEqual(calls, [])

    def test_shortcuts_actually_save_a_draft_file(self):
        out = os.path.join(self.tmp.name, 'k.json')
        self.app.personnel.insert(0, '快捷鍵')
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out):
            self._press(self.app.personnel, 's')
        self.assertTrue(os.path.exists(out))

    # ================= 合併後的服勤照片分頁：表格、● 標記、說明輸入、跳到下一張 =================
    def _tab_text(self, title):
        return self.app.notebook.tab(self.app._tab_by_title(title), 'text')

    def _all_main_titles(self):
        return [g.TAB_BASIC, g.TAB_ROAD, g.TAB_FACILITY, g.TAB_NATURE, g.TAB_FEEDBACK, g.TAB_PHOTOS]

    def test_empty_form_shows_warning_mark_on_every_tab(self):
        for title in self._all_main_titles():
            with self.subTest(tab=title):
                self.assertEqual(self._tab_text(title), f'{title} ●')

    def test_tab_marks_never_show_counts(self):
        import re
        for tab in self.app.notebook.tabs():
            self.assertIsNone(re.search(r'●\s*\d', self.app.notebook.tab(tab, 'text')))
        for tab in self.app.photo_notebook.tabs():
            self.assertIsNone(re.search(r'\d', self.app.photo_notebook.tab(tab, 'text')))

    def test_text_tab_mark_disappears_when_all_fields_filled_and_returns_when_one_is_emptied(self):
        keys = ['road_status', 'road_grass', 'road_other']
        for key in keys[:2]:
            self.app.text_fields[key].insert('1.0', '內容')
        self.app._refresh_tab_marks()
        self.assertEqual(self._tab_text(g.TAB_ROAD), f'{g.TAB_ROAD} ●')  # 還有一個沒填
        self.app.text_fields[keys[2]].insert('1.0', '內容')
        self.app._refresh_tab_marks()
        self.assertEqual(self._tab_text(g.TAB_ROAD), g.TAB_ROAD)           # 全部填完：● 消失
        self.app.text_fields[keys[1]].delete('1.0', tk.END)
        self.app._refresh_tab_marks()
        self.assertEqual(self._tab_text(g.TAB_ROAD), f'{g.TAB_ROAD} ●')    # 又清空一個：● 回來

    def test_tab_mark_updates_live_while_typing(self):
        nb = self.app.notebook
        nb.select(self.app._tab_by_title(g.TAB_ROAD))
        self.root.update()
        for key in ('road_status', 'road_grass', 'road_other'):
            widget = self.app.text_fields[key]
            widget.focus_force()
            self.root.update()
            widget.insert('1.0', 'x')
            widget.event_generate('<KeyRelease>', keysym='x')
            self.root.update()
        self.assertEqual(self._tab_text(g.TAB_ROAD), g.TAB_ROAD)

    def test_basic_tab_mark_needs_person_location_and_service_dates(self):
        self.assertTrue(self._tab_text(g.TAB_BASIC).endswith('●'))
        self.fill_dates()
        self.app.personnel.insert(0, '王小明')
        self.app._refresh_tab_marks()
        self.assertTrue(self._tab_text(g.TAB_BASIC).endswith('●'))  # 還缺協勤地點
        self.app.location.insert(0, '排雲山莊')
        self.app._refresh_tab_marks()
        self.assertEqual(self._tab_text(g.TAB_BASIC), g.TAB_BASIC)

    def test_photo_tab_mark_covers_empty_categories_and_missing_captions(self):
        self._fill_everything()
        self.app._refresh_tab_marks()
        self.assertEqual(self._tab_text(g.TAB_PHOTOS), g.TAB_PHOTOS)  # 六個分類都有照片、都有說明
        cat = g.PHOTO_CATEGORIES[3]
        path = self.app.photo_files[cat][0]
        self.app.photo_captions[cat].pop(path)
        self.app._render_list(cat)
        self.assertEqual(self._tab_text(g.TAB_PHOTOS), f'{g.TAB_PHOTOS} ●')  # 一張沒說明
        self.app.photo_captions[cat][path] = '補上'
        self.app._render_list(cat)
        self.assertEqual(self._tab_text(g.TAB_PHOTOS), g.TAB_PHOTOS)
        self.app.clear_photos(g.PHOTO_CATEGORIES[5], self.app.listboxes[g.PHOTO_CATEGORIES[5]])
        self.assertEqual(self._tab_text(g.TAB_PHOTOS), f'{g.TAB_PHOTOS} ●')  # 有分類沒加入照片

    def test_marks_follow_clear_and_draft_load(self):
        self._fill_everything()
        self.app._refresh_tab_marks()
        draft = os.path.join(self.tmp.name, 'd.json')
        self._save_draft_to(draft)
        self.assertTrue(all(self._tab_text(t) == t for t in self._all_main_titles()))
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: True):
            self.app.clear_all()
        self.assertTrue(all(self._tab_text(t) == f'{t} ●' for t in self._all_main_titles()))
        self._load_draft_from(draft)
        self.assertTrue(all(self._tab_text(t) == t for t in self._all_main_titles()))

    def test_sub_tab_marks_and_category_headers(self):
        cat, other = g.PHOTO_CATEGORIES[0], g.PHOTO_CATEGORIES[1]
        short = g.PHOTO_SHORT_NAMES[cat]
        sub = lambda c: self.app.photo_notebook.tab(self.app.photo_pages[c], 'text')
        self.assertEqual(self.app.category_labels[cat].cget('text'), f'{cat}（尚未加入照片）')
        self.assertEqual(sub(cat), f'{short}  ●')
        _, table, (a, b, c) = self._photos()
        self.assertEqual(self.app.category_labels[cat].cget('text'), f'{cat}（3 張，3 張未填說明）')
        self.app.photo_captions[cat][a] = '一'
        self.app._render_list(cat)
        self.assertEqual(self.app.category_labels[cat].cget('text'), f'{cat}（3 張，2 張未填說明）')
        self.assertEqual(sub(cat), f'{short}  ●')
        self.app.photo_captions[cat][b] = '二'
        self.app.photo_captions[cat][c] = '三'
        self.app._render_list(cat)
        self.assertEqual(self.app.category_labels[cat].cget('text'), f'{cat}（3 張，全部已填 ✓）')
        self.assertEqual(sub(cat), f'{short}  ✓')
        self.assertEqual(sub(other), f'{g.PHOTO_SHORT_NAMES[other]}  ●')  # 其他分類不受影響

    def test_table_rows_show_markers_missing_style_and_captions(self):
        cat, table, (a, b) = self._photos(names=('a.jpg', 'b.jpg'))
        self.app.photo_captions[cat][a] = '山屋'
        self.app._render_list(cat)
        self.assertEqual(tuple(table.item('0', 'values')), ('✓ (橫) a.jpg', '山屋'))
        self.assertEqual(tuple(table.item('0', 'tags')), ('done',))
        self.assertEqual(tuple(table.item('1', 'values')), ('● (橫) b.jpg', '（未填說明）'))
        self.assertEqual(tuple(table.item('1', 'tags')), ('missing',))
        palette = g.THEMES[g.DEFAULT_THEME]
        self.assertEqual(str(table.tag_configure('missing', 'background')), palette['missing_bg'])
        self.assertEqual(str(table.tag_configure('missing', 'foreground')), palette['missing_fg'])

    def test_missing_row_colors_follow_the_theme(self):
        cat, table, _ = self._photos(names=('a.jpg',))
        self.app.apply_theme('色弱友善')
        self.assertEqual(str(table.tag_configure('missing', 'background')), g.THEMES['色弱友善']['missing_bg'])

    def test_missing_row_colors_are_readable_in_every_theme(self):
        for name, p in g.THEMES.items():
            with self.subTest(theme=name):
                self.assertGreaterEqual(self._contrast(p['missing_fg'], p['missing_bg']), 4.5)

    def test_table_has_two_columns_with_narrow_filename_column_and_fills_height(self):
        self.app.notebook.select(self.app._tab_by_title(g.TAB_PHOTOS))
        self.root.update()
        table = self.app.listboxes[g.PHOTO_CATEGORIES[0]]
        self.assertEqual(tuple(table.cget('columns')), ('name', 'caption'))
        self.assertEqual([table.heading(c, 'text') for c in ('name', 'caption')], ['檔名', '說明'])
        self.assertLessEqual(abs(int(table.column('name', 'width')) - int(table.column('caption', 'width'))), 2)  # 兩欄對半分
        self.assertGreater(table.winfo_height(), g.PHOTO_TABLE_ROWS * 34, '表格沒有撐滿畫面高度')
        canvas, inner, _ = self.app._scroll_tabs[g.TAB_PHOTOS]
        self.assertGreater(table.winfo_height(), canvas.winfo_height() * 0.5)  # 至少占可視高度的一半以上

    def test_window_is_wide_and_has_no_listbox_left(self):
        expected = min(g.WINDOW_WIDTH, max(900, self.root.winfo_screenwidth() - 40))
        self.assertEqual(self.root.winfo_width(), expected)
        self.assertEqual([w for w in self._all_widgets(self.root) if isinstance(w, tk.Listbox)], [])

    def test_jump_to_next_missing_caption(self):
        cat, table, (a, b, c) = self._photos()
        self.app.photo_captions[cat][a] = '一'
        self.app._render_list(cat)
        jump = self.app.jump_buttons[cat]
        self.assertEqual(jump.cget('text'), '跳到下一張未填說明（2）(N)')
        self.assertTrue(self.app._jump_next_missing(cat))
        self.assertEqual(table.selected_indexes(), (1,))
        self.assertEqual(self.app.previews[cat].current, (cat, b))
        self.assertEqual(self.app.caption_var.get(), '')
        self.assertTrue(self.app._jump_next_missing(cat))
        self.assertEqual(table.selected_indexes(), (2,))
        self.app.photo_captions[cat][b] = '二'
        self.app._render_list(cat)
        self.assertTrue(self.app._jump_next_missing(cat))   # 從最後一張再找：回頭找到沒填的（只剩 c 自己）
        self.assertEqual(table.selected_indexes(), (2,))
        self.app.photo_captions[cat][c] = '三'
        self.app._render_list(cat)
        self.assertFalse(self.app._jump_next_missing(cat))
        self.assertEqual(jump.cget('text'), '全部都已填寫 ✓')
        self.assertEqual(str(jump.cget('state')), 'disabled')
        self.assertIn('都已填寫', self.app.status_label.cget('text'))

    def test_jump_button_states_for_empty_category(self):
        cat = g.PHOTO_CATEGORIES[4]
        self.assertEqual(self.app.jump_buttons[cat].cget('text'), '尚無照片')
        self.assertEqual(str(self.app.jump_buttons[cat].cget('state')), 'disabled')

    def test_n_key_in_table_jumps_to_next_missing_and_ctrl_n_does_not(self):
        cat, table, (a, b, c) = self._photos()
        self.app.photo_captions[cat][a] = '一'
        self.app._render_list(cat)
        self.app.notebook.select(self.app._tab_by_title(g.TAB_PHOTOS))
        self.root.update()
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        table.focus_force()
        self.root.update()
        table.event_generate('<KeyPress-n>', keysym='n')
        self.root.update()
        self.assertEqual(table.selected_indexes(), (1,))              # 按 n：跳到下一張未填說明
        self.assertEqual(self.app.previews[cat].current, (cat, b))
        table.focus_force()
        table.event_generate('<KeyPress-N>', keysym='N')               # 大寫 N 也可以
        self.root.update()
        self.assertEqual(table.selected_indexes(), (2,))
        table.focus_force()
        table.event_generate('<KeyPress-n>', keysym='n', state=4)      # Ctrl+N：不處理
        self.root.update()
        self.assertEqual(table.selected_indexes(), (2,))

    def test_typing_n_in_caption_box_only_types_the_letter(self):
        cat, table, (a, b, c) = self._photos()
        self.app.notebook.select(self.app._tab_by_title(g.TAB_PHOTOS))
        self.root.update()
        table.set_selection([1])
        self.app._on_photo_select(cat, table)
        self.app.caption_entry.focus_force()
        self.root.update()
        self.app.caption_entry.event_generate('<KeyPress-n>', keysym='n')
        self.root.update()
        self.assertEqual(self.app.caption_var.get(), 'n')
        self.assertEqual(table.selected_indexes(), (1,))  # 沒有跳走

    def test_ctrl_n_is_no_longer_a_global_shortcut(self):
        self.assertEqual(self.root.bind_all('<Control-n>'), '')
        self.assertEqual(self.root.bind_all('<Control-N>'), '')
        self.assertFalse(hasattr(self.app, '_shortcut_jump'))
        self.assertIn('按 N', g.PHOTO_TAB_HINT)
        self.assertNotIn('Ctrl', g.PHOTO_TAB_HINT)

    def test_unsaved_caption_text_is_included_in_draft_and_data(self):
        cat, table, (a, b, c) = self._photos()
        table.set_selection([1])
        self.app._on_photo_select(cat, table)
        self.app.caption_var.set('還沒按 Enter')
        self.assertEqual(self.app.build_draft()['captions'][cat][b], '還沒按 Enter')
        self.assertEqual(table.item('1', 'values')[1], '還沒按 Enter')

    def test_caption_is_saved_when_selection_changes_or_focus_leaves(self):
        cat, table, (a, b, c) = self._photos()
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        self.app.caption_var.set('切換前輸入')
        table.set_selection([2])
        self.app._on_photo_select(cat, table)  # 換選別張：前一張的輸入要先存起來
        self.assertEqual(self.app.photo_captions[cat][a], '切換前輸入')
        self.app.caption_var.set('離開輸入框')
        self.app.caption_entry.event_generate('<FocusOut>')
        self.assertEqual(self.app.photo_captions[cat][c], '離開輸入框')

    def test_emptying_the_caption_box_removes_the_caption(self):
        cat, table, (a, b, c) = self._photos()
        self.app.photo_captions[cat][a] = '有'
        self.app._render_list(cat)
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        self.app.caption_var.set('   ')
        self.app._commit_caption()
        self.assertNotIn(a, self.app.photo_captions[cat])
        self.assertEqual(tuple(table.item('0', 'tags')), ('missing',))

    def test_double_click_row_moves_to_caption_box(self):
        cat, table, (a, b, c) = self._photos()
        table.set_selection([1])
        self.app._focus_caption(cat, table)
        self.assertEqual(self.app.previews[cat].current, (cat, b))
        self.assertEqual(str(self.app.caption_entry.cget('state')), 'normal')

    def test_switching_sub_tab_shows_that_categorys_selected_photo_or_clears(self):
        cat0, table0, (a, b, c) = self._photos()
        table0.set_selection([2])
        self.app._on_photo_select(cat0, table0)
        self.app.photo_notebook.select(self.app.photo_pages[g.PHOTO_CATEGORIES[1]])  # 設施：沒有照片
        self.root.update()
        self.assertIsNone(self.app.previews[cat0].current)
        self.assertEqual(str(self.app.caption_entry.cget('state')), 'disabled')
        self.app.photo_notebook.select(self.app.photo_pages[cat0])
        self.root.update()
        self.assertEqual(self.app.previews[cat0].current, (cat0, c))

    # ---- 產生前：也檢查沒填說明的照片
    def test_find_missing_captions_in_word_order(self):
        self._fill_everything()
        for i in (4, 1):
            cat = g.PHOTO_CATEGORIES[i]
            self.app.photo_captions[cat].clear()
            self.app._render_list(cat)
        self.assertEqual(self.app.find_missing_captions(),
                         {g.PHOTO_CATEGORIES[1]: 1, g.PHOTO_CATEGORIES[4]: 1})
        self.assertEqual(list(self.app.find_missing_captions()), [g.PHOTO_CATEGORIES[1], g.PHOTO_CATEGORIES[4]])

    def test_message_lists_missing_captions_after_blank_fields(self):
        message = g.ReportGeneratorGUI.blank_fields_message({'基本資料': ['協勤地點']}, {'其他': 2, '設施狀況及維護情形': 1})
        self.assertIn('以下欄位尚未填寫：', message)
        self.assertIn('另外有 3 張照片沒有填寫說明：', message)
        self.assertIn('    • 其他：2 張', message)
        self.assertTrue(message.endswith('仍要產生報告嗎？'))
        only = g.ReportGeneratorGUI.blank_fields_message({}, {'其他': 2})
        self.assertTrue(only.startswith('有 2 張照片沒有填寫說明：'))
        self.assertNotIn('以下欄位尚未填寫', only)

    def test_generate_warns_when_only_captions_are_missing(self):
        self._fill_everything()
        cat = g.PHOTO_CATEGORIES[0]
        self.app.photo_captions[cat].clear()
        self.app._render_list(cat)
        out = os.path.join(self.tmp.name, 'o.docx')
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out), \
                mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertEqual(len(self.confirm_calls), 1)
        _, message, kwargs = self.confirm_calls[0]
        self.assertEqual(kwargs.get('default'), 'yes')
        self.assertIn('有 1 張照片沒有填寫說明：', message)
        self.assertIn(f'    • {cat}：1 張', message)
        self.assertNotIn('以下欄位尚未填寫', message)
        self.assertTrue(os.path.exists(out))

    def test_answering_no_for_missing_captions_selects_the_first_missing_photo(self):
        self._fill_everything()
        cat = g.PHOTO_CATEGORIES[1]  # 設施
        self.app.photo_captions[cat].clear()
        self.app._render_list(cat)
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: False):
            self.app.generate_report()
        self.assertEqual(self.app.notebook.select(), self.app._tab_by_title(g.TAB_PHOTOS))
        self.assertEqual(self.app.photo_notebook.select(), str(self.app.photo_pages[cat]))
        self.assertEqual(self.app.listboxes[cat].selected_indexes(), (0,))
        self.assertEqual(self.app.previews[cat].current[0], cat)
        self.assertEqual(self.pdf_calls, [])

    def test_unsaved_caption_text_counts_for_the_pre_generation_check(self):
        self._fill_everything()
        cat = g.PHOTO_CATEGORIES[2]
        path = self.app.photo_files[cat][0]
        self.app.photo_captions[cat].pop(path)
        self.app._render_list(cat)
        table = self.app.listboxes[cat]
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        self.app.caption_var.set('剛打的，還沒按 Enter')
        out = os.path.join(self.tmp.name, 'o.docx')
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out), \
                mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertEqual(self.confirm_calls, [])  # 輸入框裡的文字會先存起來，不會誤報「沒填說明」

    # ---- 版面：左邊表格區切齊小分頁列、右邊預覽往左與往下擴大
    def _show_photos_tab(self):
        self.app.notebook.select(self.app._tab_by_title(g.TAB_PHOTOS))
        self.root.update()
        self.root.update_idletasks()
        self.root.update()

    def test_left_frame_right_edge_matches_last_sub_tab(self):
        self._show_photos_tab()
        sub = self.app.photo_notebook
        geometry = g.tab_geometry(sub, len(sub.tabs()) - 1)
        self.assertIsNotNone(geometry)
        self.assertLessEqual(abs(sub.winfo_width() - geometry[1]), 6,
                             f'表格區寬 {sub.winfo_width()}，小分頁列右緣 {geometry[1]}')
        table = self.app.listboxes[g.PHOTO_CATEGORIES[0]]
        self.assertLess(sub.winfo_width(), 800)
        self.assertLessEqual(abs(int(table.column('name', 'width')) - int(table.column('caption', 'width'))), 2)
        self.assertLess(int(table.column('caption', 'width')), 480)

    def test_preview_panel_takes_the_remaining_width_and_the_full_height(self):
        self._show_photos_tab()
        sub, right = self.app.photo_notebook, self.app.photo_right_panel
        preview = self.app.photo_preview
        self.assertGreaterEqual(right.winfo_rootx(), sub.winfo_rootx() + sub.winfo_width())  # 在表格區右邊
        self.assertGreater(right.winfo_width(), g.PREVIEW_SIZE[0] + 150)                      # 比最小尺寸寬很多
        self.assertGreater(preview.box.winfo_height(), g.PREVIEW_SIZE[1] + 100)               # 往下擴大
        sub_bottom = sub.winfo_rooty() + sub.winfo_height()
        right_bottom = right.winfo_rooty() + right.winfo_height()
        self.assertLessEqual(abs(sub_bottom - right_bottom), 12)                              # 底部切齊
        # 右邊的右緣就是視窗內容區的右緣（沒有留一大塊空白）
        canvas, inner, _ = self.app._scroll_tabs[g.TAB_PHOTOS]
        self.assertGreaterEqual(right.winfo_rootx() + right.winfo_width(), canvas.winfo_rootx() + canvas.winfo_width() - 40)

    def test_preview_image_fills_the_box_and_follows_window_size(self):
        cat, table, (a, b, c) = self._photos(size=(1600, 1000))
        self._show_photos_tab()
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        preview = self.app.photo_preview
        self.root.update()
        bw, bh = preview.box_size()
        self.assertLessEqual(preview.image.width(), bw)
        self.assertLessEqual(preview.image.height(), bh)
        self.assertGreater(max(preview.image.width() / bw, preview.image.height() / bh), 0.95)  # 至少有一邊幾乎填滿
        before = preview.image.width()
        width, height = self.root.winfo_width(), self.root.winfo_height()
        try:
            self.root.geometry(f'{width - 300}x{height}')
            self.root.update()
            self.root.update_idletasks()
            self.root.update()
            self.assertLess(preview.image.width(), before)  # 視窗變窄，預覽圖跟著縮小
        finally:
            self.root.geometry(f'{width}x{height}')
            self.root.update()

    def test_preview_never_shrinks_below_minimum_box(self):
        self._show_photos_tab()
        box = self.app.photo_preview.box
        self.assertGreaterEqual(box.winfo_width(), g.PREVIEW_SIZE[0])
        self.assertGreaterEqual(box.winfo_height(), g.PREVIEW_SIZE[1])

    def test_filename_and_caption_columns_split_the_width_in_half(self):
        self._show_photos_tab()
        table = self.app.listboxes[g.PHOTO_CATEGORIES[0]]
        name, caption = int(table.column('name', 'width')), int(table.column('caption', 'width'))
        self.assertLessEqual(abs(name - caption), 2)
        self.assertGreaterEqual(name + caption, table.winfo_width() - 4)  # 兩欄剛好把表格寬度占滿
        self.assertGreater(name, 290)  # 檔名欄比先前（290）寬
        # 視窗變窄，兩欄跟著各縮一半
        w, h = self.root.winfo_width(), self.root.winfo_height()
        try:
            self.root.geometry(f'{w - 240}x{h}')
            self.root.update()
            self.root.update_idletasks()
            self.root.update()
            n2, c2 = int(table.column('name', 'width')), int(table.column('caption', 'width'))
            self.assertLessEqual(abs(n2 - c2), 2)
        finally:
            self.root.geometry(f'{w}x{h}')
            self.root.update()

    def test_clear_everything_button_label(self):
        self.assertEqual(self.app.clear_btn.cget('text'), '清除所有資料')

    # ================= 對話框：使用程式的字型與較大的字級 =================
    def _dialog(self):
        d = g.AppDialogs()
        d.attach(self.app)
        return d

    def _answer_dialog(self, action):
        """對話框開啟後（wait_window 期間）執行 action(對話框視窗)。"""
        def run():
            dialogs = [w for w in self.root.winfo_children() if isinstance(w, tk.Toplevel)]
            self.assertTrue(dialogs, '沒有對話框')
            action(dialogs[-1])
        self.root.after(150, run)

    def _find_button(self, dialog, text):
        for w in self._all_widgets(dialog):
            if isinstance(w, ttk.Button) and w.cget('text') == text:
                return w
        self.fail(f'找不到按鈕 {text}')

    def test_dialog_text_is_larger_than_normal_interface_text(self):
        import tkinter.font as tkfont
        d = self._dialog()
        seen = {}

        def look(dialog):
            text = d.last_text
            font = tkfont.Font(root=self.root, font=text.cget('font'))
            seen['size'] = abs(font.cget('size'))
            seen['family'] = font.cget('family')
            seen['message'] = text.get('1.0', 'end-1c')
            seen['button_font'] = tkfont.Font(root=self.root, font=ttk.Style(self.root).lookup('Dialog.TButton', 'font'))
            self._find_button(dialog, '是(Y)').invoke()

        self._answer_dialog(look)
        self.assertTrue(d.askyesno('確認', '載入草稿會覆蓋目前畫面上的內容，確定嗎？'))
        body = abs(__import__('tkinter.font', fromlist=['x']).nametofont('TkDefaultFont', root=self.root).cget('size'))
        self.assertGreaterEqual(seen['size'], body + 2)             # 比一般介面文字至少大 2 點
        self.assertGreaterEqual(seen['size'], 13)
        self.assertEqual(seen['family'], g.FONT_FAMILY)
        self.assertEqual(seen['message'], '載入草稿會覆蓋目前畫面上的內容，確定嗎？')
        self.assertGreaterEqual(abs(seen['button_font'].cget('size')), 13)  # 按鈕文字也一樣大

    def test_askyesno_default_yes_enter_and_keys(self):
        d = self._dialog()
        focus = {}

        def press_enter(dialog):
            focus['text'] = self.root.focus_get().cget('text') if self.root.focus_get() else None
            dialog.event_generate('<Return>')

        self._answer_dialog(press_enter)
        self.assertTrue(d.askyesno('t', 'm'))                         # 預設「是」，按 Enter = 是
        self.assertEqual(focus['text'], '是(Y)')
        self._answer_dialog(lambda dialog: dialog.event_generate('<KeyPress-n>'))
        self.assertFalse(d.askyesno('t', 'm'))                        # 按 n = 否
        self._answer_dialog(lambda dialog: dialog.event_generate('<KeyPress-Y>'))
        self.assertTrue(d.askyesno('t', 'm', default='no'))           # 大寫 Y 也可以

    def test_askyesno_default_no_and_escape(self):
        d = self._dialog()
        focus = {}

        def press_enter(dialog):
            focus['text'] = self.root.focus_get().cget('text')
            dialog.event_generate('<Return>')

        self._answer_dialog(press_enter)
        self.assertFalse(d.askyesno('t', 'm', icon='warning', default='no'))   # 預設「否」：按 Enter 不會誤清除
        self.assertEqual(focus['text'], '否(N)')
        self._answer_dialog(lambda dialog: dialog.event_generate('<Escape>'))
        self.assertFalse(d.askyesno('t', 'm'))                                   # Esc = 否

    def test_askyesnocancel_values(self):
        d = self._dialog()
        self._answer_dialog(lambda dialog: self._find_button(dialog, '是(Y)').invoke())
        self.assertIs(d.askyesnocancel('t', 'm'), True)
        self._answer_dialog(lambda dialog: self._find_button(dialog, '否(N)').invoke())
        self.assertIs(d.askyesnocancel('t', 'm'), False)
        self._answer_dialog(lambda dialog: self._find_button(dialog, '取消(C)').invoke())
        self.assertIsNone(d.askyesnocancel('t', 'm'))
        self._answer_dialog(lambda dialog: dialog.event_generate('<Escape>'))
        self.assertIsNone(d.askyesnocancel('t', 'm'))                            # Esc = 取消
        self._answer_dialog(lambda dialog: dialog.event_generate('<KeyPress-c>'))
        self.assertIsNone(d.askyesnocancel('t', 'm'))

    def test_info_and_error_dialogs_close_with_enter_or_button(self):
        d = self._dialog()
        self._answer_dialog(lambda dialog: dialog.event_generate('<Return>'))
        self.assertEqual(d.showinfo('t', '完成'), 'ok')
        self._answer_dialog(lambda dialog: self._find_button(dialog, '確定').invoke())
        self.assertEqual(d.showerror('t', '錯誤'), 'ok')

    def test_long_message_gets_scrollbar_and_short_message_does_not(self):
        d = self._dialog()
        info = {}

        def look(dialog):
            scrollbars = [w for w in self._all_widgets(dialog) if isinstance(w, ttk.Scrollbar)]
            info['scroll'] = bool(scrollbars and scrollbars[0].winfo_manager())
            info['height'] = int(d.last_text.cget('height'))
            dialog.event_generate('<Return>')

        self._answer_dialog(look)
        d.showinfo('t', '\n'.join(f'第 {i} 行' for i in range(60)))
        self.assertTrue(info['scroll'])
        self.assertEqual(info['height'], g.DIALOG_MAX_LINES)
        self._answer_dialog(look)
        d.showinfo('t', '短訊息')
        self.assertFalse(info['scroll'])
        self.assertLess(info['height'], g.DIALOG_MAX_LINES)

    def test_dialog_is_modal_and_centered_over_the_window(self):
        d = self._dialog()
        seen = {}

        def look(dialog):
            seen['grab'] = self.root.grab_current() is not None or dialog.grab_current() is not None
            dialog.update_idletasks()
            seen['cx'] = dialog.winfo_rootx() + dialog.winfo_width() // 2
            seen['wx'] = self.root.winfo_rootx() + self.root.winfo_width() // 2
            seen['transient'] = str(dialog.transient())
            dialog.event_generate('<Return>')

        self._answer_dialog(look)
        d.showinfo('t', 'm')
        self.assertTrue(seen['grab'])
        self.assertLessEqual(abs(seen['cx'] - seen['wx']), 40)
        self.assertEqual(seen['transient'], str(self.root))

    def test_dialog_follows_theme_colors(self):
        self.app.apply_theme('高對比')
        d = self._dialog()
        seen = {}

        def look(dialog):
            seen['bg'] = str(d.last_text.cget('background'))
            seen['fg'] = str(d.last_text.cget('foreground'))
            seen['default_style'] = str(self._find_button(dialog, '確定').cget('style'))
            dialog.event_generate('<Return>')

        self._answer_dialog(look)
        d.showinfo('t', 'm')
        self.assertEqual(seen['bg'], g.THEMES['高對比']['bg'])
        self.assertEqual(seen['fg'], g.THEMES['高對比']['text'])
        self.assertEqual(seen['default_style'], 'DialogDefault.TButton')
        self.assertEqual(ttk.Style(self.root).lookup('DialogDefault.TButton', 'background'), g.THEMES['高對比']['primary'])

    def test_app_uses_its_own_dialogs_not_the_system_ones(self):
        self.assertIsInstance(g.messagebox, g.AppDialogs)
        self.assertIs(g.messagebox.app, self.app)
        self.assertFalse(hasattr(g, 'simpledialog'))

    def test_falls_back_to_system_dialogs_when_not_attached(self):
        d = g.AppDialogs()
        calls = []
        with mock.patch.object(g.native_messagebox, 'askyesno', lambda *a, **k: calls.append(('askyesno', a, k)) or True), \
                mock.patch.object(g.native_messagebox, 'showinfo', lambda *a, **k: calls.append(('showinfo', a, k)) or 'ok'):
            self.assertTrue(d.askyesno('標題', '訊息', default='no'))
            self.assertEqual(d.showinfo('標題', '訊息'), 'ok')
        self.assertEqual([c[0] for c in calls], ['askyesno', 'showinfo'])
        self.assertEqual(calls[0][2].get('default'), 'no')

    def test_real_load_draft_confirmation_uses_the_large_dialog(self):
        """實際走「載入草稿」流程：確認訊息顯示在放大的對話框裡。"""
        draft = os.path.join(self.tmp.name, 'd.json')
        self._save_draft_to(draft)
        seen = {}

        def look(dialog):
            seen['message'] = g.messagebox.last_text.get('1.0', 'end-1c')
            seen['size'] = abs(__import__('tkinter.font', fromlist=['x']).Font(root=self.root, font=g.messagebox.last_text.cget('font')).cget('size'))
            self._find_button(dialog, '否(N)').invoke()

        real = g.AppDialogs.askyesno
        with mock.patch.object(g.messagebox, 'askyesno', lambda *a, **k: real(g.messagebox, *a, **k)), \
                mock.patch.object(filedialog, 'askopenfilename', lambda **k: draft):
            self._answer_dialog(look)
            self.app.load_draft()
        self.assertEqual(seen['message'], '載入草稿會覆蓋目前畫面上的內容，確定嗎？')
        self.assertGreaterEqual(seen['size'], 13)

    # ---- 沒有安裝 Word 的電腦：Word 檔照常產生，只是沒有 PDF，且訊息要看得懂
    def test_pdf_failure_message_when_word_is_not_installed(self):
        from docx import Document
        docx = os.path.join(self.tmp.name, 'a.docx')
        Document().save(docx)
        with mock.patch.object(g, '_PDF_SCRIPT', g._PDF_SCRIPT.replace('Word.Application', 'Word.ApplicationNotInstalled')):
            with self.assertRaises(RuntimeError) as ctx:
                _REAL_EXPORT_PDF(docx, os.path.join(self.tmp.name, 'a.pdf'))
        message = str(ctx.exception)
        self.assertIn('找不到 Microsoft Word', message)
        self.assertNotIn('CLIXML', message)
        self.assertNotIn('<Objs', message)
        self.assertTrue(os.path.exists(docx))

    def test_describe_pdf_failure_cleans_powershell_noise(self):
        noisy = b'#< CLIXML\r\n<Objs Version="1.1.0.1"><Obj S="progress"></Obj></Objs>'
        self.assertEqual(g._describe_pdf_failure(noisy), 'Word 轉換 PDF 失敗')
        self.assertIn('找不到 Microsoft Word', g._describe_pdf_failure(b'Retrieving the COM class factory ... 80040154'))
        self.assertIn('boom', g._describe_pdf_failure('boom'))

    def test_without_word_report_still_generated_and_user_is_told_why(self):
        self.fill_dates()
        out = os.path.join(self.tmp.name, 'o.docx')

        def no_word(docx, pdf):
            raise RuntimeError(g._describe_pdf_failure(b'Retrieving the COM class factory ... 80040154'))

        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out),                 mock.patch.object(g, 'export_pdf', no_word),                 mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertTrue(os.path.exists(out))                    # Word 檔照常產生
        self.assertFalse(os.path.exists(out[:-5] + '.pdf'))     # 沒有 PDF
        infos = [m for k, m in self.shown if k == 'info']
        self.assertTrue(any('PDF 未產生' in m and '找不到 Microsoft Word' in m for m in infos), infos)
        self.assertEqual(str(self.app.generate_btn['state']), 'normal')  # 程式沒有卡住

    def test_open_file_without_associated_program_does_not_crash(self):
        with mock.patch.object(os, 'startfile', side_effect=OSError('no association'), create=True):
            g.open_file(os.path.join(self.tmp.name, 'x.docx'))  # 沒有 Office 也沒有其他程式可開時，只記錄不報錯

    # ---- 從檔案總管拖曳加入照片
    def _make_images(self, names):
        paths = []
        for n in names:
            path = os.path.join(self.tmp.name, n)
            Image.new('RGB', (80, 60), 'blue').save(path)
            paths.append(path)
        return paths

    def test_delete_key_removes_selected_photo_and_selects_the_next(self):
        cat, table, (a, b, c) = self._photos()
        self.app.photo_captions[cat][b] = '要刪的'
        self._show_photos_tab()
        table.set_selection([1])
        self.app._on_photo_select(cat, table)
        table.focus_force()
        self.root.update()
        table.event_generate('<Delete>')
        self.root.update()
        self.assertEqual(self.app.photo_files[cat], [a, c])
        self.assertNotIn(b, self.app.photo_captions[cat])
        self.assertEqual(table.selected_indexes(), (1,))  # 原位置換成下一張
        self.assertEqual(self.app.photo_preview.current, (cat, c))
        table.event_generate('<Delete>')
        table.event_generate('<Delete>')
        self.root.update()
        self.assertEqual(self.app.photo_files[cat], [])
        table.event_generate('<Delete>')  # 空表格按 Delete 不出錯
        self.root.update()

    def test_default_draft_name_carries_the_duty_name(self):
        from datetime import datetime
        now = datetime(2026, 10, 14, 20, 40)
        data = {'start_year': '115', 'start_month': '10', 'start_day': '1', 'end_year': '115', 'end_month': '10',
                'end_day': '2', 'location': '排雲/山莊'}
        self.assertEqual(g.default_draft_name(data, now), '1151001-02_排雲山莊協勤回報單草稿_1014_2040')
        self.assertEqual(g.default_draft_name(dict(data, location=''), now), '1151001-02_巡查回報單草稿_1014_2040')
        self.assertEqual(g.default_draft_name(dict(data, start_day='abc'), now), '排雲山莊協勤回報單草稿_1014_2040')
        self.assertEqual(g.default_draft_name({}, now), '巡查回報單草稿_1014_2040')
        # 人員姓名：只取第一位，多位加「等」
        self.assertEqual(g.default_draft_name(dict(data, personnel='王小明'), now),
                         '1151001-02_排雲山莊協勤回報單草稿_王小明_1014_2040')
        self.assertEqual(g.default_draft_name(dict(data, personnel='王小明、李小華 陳大文'), now),
                         '1151001-02_排雲山莊協勤回報單草稿_王小明等_1014_2040')
        self.assertEqual(g.default_output_name(dict(data, personnel='李/小華, 王小明')), '1151001-02_排雲山莊協勤回報單_李等')
        self.assertEqual(g.name_label(''), '')
        self.assertEqual(g.name_label('John Smith, Mary Jane'), 'John Smith等')

    def test_save_draft_dialog_suggests_name_with_duty_name(self):
        self.fill_dates()
        self.app.location.insert(0, '排雲山莊')
        seen = {}

        def ask(**kw):
            seen.update(kw)
            return ''

        with mock.patch.object(filedialog, 'asksaveasfilename', ask):
            self.app.save_draft()
        self.assertIn('排雲山莊協勤回報單草稿_', seen['initialfile'])
        self.assertTrue(seen['initialfile'].endswith('.json'))

    def test_collect_dropped_images_filters_and_expands_folders(self):
        a, b = self._make_images(['b.JPG', 'a.png'])
        txt = os.path.join(self.tmp.name, 'note.txt')
        open(txt, 'w').write('x')
        folder = os.path.join(self.tmp.name, 'sub')
        os.mkdir(folder)
        for n in ('z.jpg', 'm.bmp', 'skip.docx'):
            open(os.path.join(folder, n), 'wb').write(b'')
        images, ignored = g.collect_dropped_images([a, txt, folder, b])
        # 整批依檔名排序（不分大小寫），與拖進來的順序、所在資料夾無關
        self.assertEqual(images, sorted([a, b, os.path.join(folder, 'm.bmp'), os.path.join(folder, 'z.jpg')], key=lambda x: os.path.basename(x).lower()))
        self.assertEqual([os.path.basename(x) for x in images], ['a.png', 'b.JPG', 'm.bmp', 'z.jpg'])
        self.assertEqual(ignored, 1)  # 只有直接拖進來的 note.txt 算略過

    def test_dropping_images_adds_them_to_the_current_category_and_shows_photo_tab(self):
        files = self._make_images(['p1.jpg', 'p2.png'])
        self.app.photo_notebook.select(self.app.photo_pages[g.PHOTO_CATEGORIES[2]])
        self.assertEqual(self.app.handle_dropped_files(files), 2)
        self.assertEqual(self.app.photo_files[g.PHOTO_CATEGORIES[2]], files)
        self.assertEqual(self.app.notebook.index(self.app.notebook.select()), self.app.notebook.index(self.app._tab_by_title(g.TAB_PHOTOS)))
        self.assertEqual(self.app.listboxes[g.PHOTO_CATEGORIES[2]].count(), 2)
        self.assertEqual(self.app.handle_dropped_files(files), 0)  # 重複不再加入
        self.assertEqual(len(self.app.photo_files[g.PHOTO_CATEGORIES[2]]), 2)
        self.assertIn('已加入過', self.app.status_label.cget('text'))

    def test_dropping_non_images_adds_nothing(self):
        txt = os.path.join(self.tmp.name, 'a.txt')
        open(txt, 'w').write('x')
        self.assertEqual(self.app.handle_dropped_files([txt]), 0)
        self.assertTrue(all(not v for v in self.app.photo_files.values()))
        self.assertIn('沒有圖片', self.app.status_label.cget('text'))

    def test_dropping_on_a_sub_tab_header_adds_to_that_category(self):
        files = self._make_images(['q.jpg'])
        self._show_photos_tab()
        sub = self.app.photo_notebook
        target = 3
        left, right, top, bottom = g.tab_geometry(sub, target)
        sx, sy = sub.winfo_rootx() + (left + right) // 2, sub.winfo_rooty() + (top + bottom) // 2
        self.app.handle_dropped_files(files, sx, sy)
        self.assertEqual(self.app.photo_files[g.PHOTO_CATEGORIES[target]], files)
        self.assertEqual(self.app._current_photo_category(), g.PHOTO_CATEGORIES[target])

    def test_dropped_images_are_sorted_naturally_without_touching_existing_order(self):
        existing = self._make_images(['z_old.jpg', 'a_old.jpg'])
        cat = g.PHOTO_CATEGORIES[0]
        self.app.photo_files[cat].extend(existing)
        new = self._make_images(['IMG_10.jpg', 'img_2.jpg', 'IMG_1.jpg'])
        self.app.photo_notebook.select(self.app.photo_pages[cat])
        self.app.handle_dropped_files([new[0], new[1], new[2]])
        names = [os.path.basename(f) for f in self.app.photo_files[cat]]
        self.assertEqual(names, ['z_old.jpg', 'a_old.jpg', 'IMG_1.jpg', 'img_2.jpg', 'IMG_10.jpg'])

    def test_busy_indicator_runs_while_generating_and_stops_afterwards(self):
        self.fill_dates()
        out = os.path.join(self.tmp.name, 'busy.docx')
        started = {}
        real_start = self.app._start_busy

        def spy(message):
            real_start(message)
            started['bar'] = bool(self.app.busy_bar.place_info())
            started['cursor'] = str(self.root.cget('cursor'))
            started['text'] = self.app.status_label.cget('text')

        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out),                 mock.patch.object(self.app, '_start_busy', spy),                 mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertTrue(started['bar'])
        self.assertEqual(started['cursor'], 'watch')
        self.assertIn('正在產生', started['text'])
        self.assertFalse(self.app.busy_bar.place_info())
        self.assertEqual(str(self.root.cget('cursor')), '')
        self.assertIn('SUCCESS', self.app.status_label.cget('text'))  # 最後的狀態文字沒有被計時文字蓋掉

    @unittest.skipUnless(sys.platform == 'win32', '需要 Windows')
    def test_real_wm_dropfiles_message_adds_photos(self):
        """用真正的 WM_DROPFILES 訊息（跟檔案總管拖曳時送出的一樣）送進視窗。"""
        import ctypes
        from ctypes import wintypes
        self.assertIsNotNone(self.app._file_drop, '沒有成功啟用拖放')
        files = self._make_images(['r1.jpg', 'r2.jpg'])
        payload = ''.join(f + '\0' for f in files) + '\0'
        data = payload.encode('utf-16-le')
        header = ctypes.sizeof(ctypes.c_uint32) + 8 + 4 + 4  # DROPFILES：pFiles、POINT、fNC、fWide
        kernel32, user32 = ctypes.windll.kernel32, ctypes.windll.user32
        kernel32.GlobalAlloc.restype = ctypes.c_void_p
        kernel32.GlobalLock.restype = ctypes.c_void_p
        kernel32.GlobalLock.argtypes = [ctypes.c_void_p]
        kernel32.GlobalUnlock.argtypes = [ctypes.c_void_p]
        handle = kernel32.GlobalAlloc(0x0002 | 0x0040, header + len(data))
        ptr = kernel32.GlobalLock(handle)
        ctypes.memmove(ptr, header.to_bytes(4, 'little'), 4)
        ctypes.memmove(ptr + 16, (1).to_bytes(4, 'little'), 4)  # fWide = TRUE
        ctypes.memmove(ptr + header, data, len(data))
        kernel32.GlobalUnlock(handle)
        user32.GetParent.restype = wintypes.HWND
        user32.PostMessageW.argtypes = [wintypes.HWND, wintypes.UINT, wintypes.WPARAM, wintypes.LPARAM]
        self.root.update()
        hwnd = user32.GetParent(self.root.winfo_id()) or self.root.winfo_id()
        # 用 PostMessage（非同步）：訊息是在主迴圈裡被派送的，和檔案總管實際拖曳時一樣。
        # 以前用 SendMessage（直接呼叫）測不出「在訊息處理中碰 tkinter 會讓程式當掉」的問題。
        user32.PostMessageW(hwnd, 0x0233, handle, 0)
        deadline = time.time() + 5
        while time.time() < deadline and not any(self.app.photo_files[c] for c in g.PHOTO_CATEGORIES):
            self.root.update()
            time.sleep(0.05)
        added = [f for cat in g.PHOTO_CATEGORIES for f in self.app.photo_files[cat]]
        self.assertEqual(sorted(added), sorted(files))

    # ---- 照片旋轉（預覽區的按鈕；只記錄角度，不改原檔，報告與草稿都會套用）
    def _select_first(self, cat, table):
        self._show_photos_tab()
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        self.root.update()

    def test_rotate_buttons_disabled_until_a_photo_is_selected(self):
        self.assertEqual(str(self.app.rotate_left_btn.cget('state')), 'disabled')
        self.assertEqual(str(self.app.rotate_right_btn.cget('state')), 'disabled')
        cat, table, (a,) = self._photos(names=('r.jpg',))
        self._select_first(cat, table)
        self.assertEqual(str(self.app.rotate_left_btn.cget('state')), 'normal')
        self.assertEqual(str(self.app.rotate_right_btn.cget('state')), 'normal')

    def test_rotating_swaps_preview_orientation_and_keeps_original_file(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200))
        self._select_first(cat, table)
        preview = self.app.photo_preview
        self.assertGreater(preview.image.width(), preview.image.height())      # 原本橫式
        self.app.rotate_right_btn.invoke()
        self.assertEqual(self.app.photo_rotations[cat][a], 90)
        self.assertGreater(preview.image.height(), preview.image.width())      # 轉 90° 變直式
        self.app.rotate_right_btn.invoke()
        self.assertEqual(self.app.photo_rotations[cat][a], 180)
        self.assertGreater(preview.image.width(), preview.image.height())
        self.app.rotate_left_btn.invoke()
        self.app.rotate_left_btn.invoke()
        self.assertNotIn(a, self.app.photo_rotations[cat])                     # 轉回 0° 就不再記錄
        self.app.rotate_left_btn.invoke()
        self.assertEqual(self.app.photo_rotations[cat][a], 270)                # 逆時針 90° = 順時針 270°
        with Image.open(a) as original:
            self.assertEqual(original.size, (400, 200))                        # 原檔沒被改動

    def test_rotation_direction_is_clockwise(self):
        im = Image.new('RGB', (4, 2), 'white')
        im.putpixel((0, 0), (255, 0, 0))  # 左上角是紅色
        turned = g.rotate_clockwise(im, 90)
        self.assertEqual(turned.size, (2, 4))
        self.assertEqual(turned.getpixel((1, 0)), (255, 0, 0))  # 順時針 90° 後，左上角跑到右上角
        self.assertEqual(g.normalize_rotation(-90), 270)
        self.assertEqual(g.normalize_rotation(450), 90)

    def test_preview_top_is_level_with_the_sub_tab_strip(self):
        self._show_photos_tab()
        sub, preview = self.app.photo_notebook, self.app.photo_preview
        self.assertLessEqual(abs(preview.winfo_rooty() - sub.winfo_rooty()), 4,
                             '預覽框上緣應該和左邊小分頁列的上緣切齊，不要留空白')

    def test_preview_has_no_filename_label_and_uses_more_space(self):
        self._show_photos_tab()
        right, preview = self.app.photo_right_panel, self.app.photo_preview
        self.assertFalse([w for w in preview.winfo_children() if isinstance(w, ttk.Label)])
        # 預覽框幾乎占滿框架（只剩很薄的邊）
        self.assertLessEqual(preview.winfo_height() - preview.box.winfo_height(), 8)

    def test_rotation_is_applied_in_the_report(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200))
        self.app.photo_rotations[cat][a] = 90
        src, size = g.prepare_image(a, 90)
        self.assertEqual(size, (200, 400))
        with Image.open(src) as rotated:
            self.assertEqual(rotated.size, (200, 400))
        self.assertEqual(g.prepare_image(a)[1], (400, 200))
        out = os.path.join(self.tmp.name, 'rot.docx')
        self.fill_dates()
        data, error = self.app.collect_data()
        self.assertIsNone(error)
        g.build_report(TEMPLATE, data, {c: list(v) for c, v in self.app.photo_files.items()}, {}, out,
                       rotations={c: dict(v) for c, v in self.app.photo_rotations.items()})
        doc = Document(out)
        shapes = doc.inline_shapes
        self.assertGreaterEqual(len(shapes), 1)
        self.assertGreater(shapes[0].height, shapes[0].width)  # 報告裡這張變成直式

    def test_rotation_survives_draft_roundtrip_and_is_dropped_with_the_photo(self):
        cat, table, (a, b, c) = self._photos()
        self.app.photo_rotations[cat][b] = 270
        draft = self.app.build_draft()
        self.assertEqual(draft['rotations'][cat], {b: 270})
        self.assertTrue(self.app.is_dirty())
        self.app.apply_draft(draft)
        self.assertEqual(self.app.photo_rotations[cat], {b: 270})
        table.set_selection([1])
        self.app.remove_photos(cat, table)
        self.assertEqual(self.app.photo_rotations[cat], {})
        self.app.photo_rotations[cat][a] = 90
        self.app.clear_photos(cat, table)
        self.assertEqual(self.app.photo_rotations[cat], {})

    def test_old_draft_without_rotations_still_loads(self):
        cat, table, (a,) = self._photos(names=('o.jpg',))
        draft = self.app.build_draft()
        draft.pop('rotations')
        self.assertEqual(self.app.apply_draft(draft), [])
        self.assertEqual(self.app.photo_rotations[cat], {})

    # ---- 檔名左側的 (橫)/(直)
    def test_filename_column_shows_landscape_or_portrait(self):
        cat = g.PHOTO_CATEGORIES[0]
        table = self.app.listboxes[cat]
        wide, tall, exif, square = self._make_images(['w.jpg', 't.jpg', 'e.jpg', 's.jpg'])
        Image.new('RGB', (100, 300), 'red').save(tall)
        im = Image.new('RGB', (300, 100), 'red')
        ex = im.getexif()
        ex[0x0112] = 6  # 轉正後變直式
        im.save(exif, exif=ex)
        Image.new('RGB', (100, 100), 'red').save(square)
        self.app._add_photo_files(cat, [wide, tall, exif, square])
        self.assertEqual([table.row_text(i) for i in range(4)],
                         ['● (橫) w.jpg', '● (直) t.jpg', '● (直) e.jpg', '● (橫) s.jpg'])  # 正方形算橫式（跟報告一致）
        bad = os.path.join(self.tmp.name, 'bad.jpg')
        open(bad, 'wb').close()
        self.app._add_photo_files(cat, [bad])
        self.assertEqual(table.row_text(4), '● (？) bad.jpg')

    def test_rotating_a_photo_flips_its_orientation_tag_in_the_table(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200))
        self._select_first(cat, table)
        self.assertEqual(table.row_text(0), '● (橫) r.jpg')
        self.app.rotate_right_btn.invoke()
        self.assertEqual(table.row_text(0), '● (直) r.jpg')
        self.app.rotate_right_btn.invoke()
        self.assertEqual(table.row_text(0), '● (橫) r.jpg')

    # ---- 拖曳加入：安插在滑鼠放開的位置
    def _drop_point(self, table, index, upper=True):
        self.root.update()
        box = table.bbox(str(index))
        return (table.winfo_rootx() + 20, table.winfo_rooty() + box[1] + (box[3] // 4 if upper else box[3] * 3 // 4))

    def test_drop_on_upper_half_of_a_row_inserts_before_it(self):
        cat, table, (a, b, c) = self._photos()
        new = self._make_images(['n1.jpg', 'n2.jpg'])
        self._show_photos_tab()
        x, y = self._drop_point(table, 1, upper=True)
        self.app.handle_dropped_files(new, x, y)
        self.assertEqual(self.app.photo_files[cat], [a, new[0], new[1], b, c])
        self.assertEqual(table.selected_indexes(), (1, 2))  # 剛放進來的照片被選取

    def test_drop_on_lower_half_of_a_row_inserts_after_it(self):
        cat, table, (a, b, c) = self._photos()
        new = self._make_images(['n1.jpg'])
        self._show_photos_tab()
        x, y = self._drop_point(table, 1, upper=False)
        self.app.handle_dropped_files(new, x, y)
        self.assertEqual(self.app.photo_files[cat], [a, b, new[0], c])

    def test_drop_on_table_heading_inserts_at_the_top_and_below_rows_appends(self):
        cat, table, (a, b, c) = self._photos()
        first, last = self._make_images(['n1.jpg', 'n2.jpg'])
        self._show_photos_tab()
        self.app.handle_dropped_files([first], table.winfo_rootx() + 20, table.winfo_rooty() + 5)  # 表頭
        self.assertEqual(self.app.photo_files[cat][0], first)
        self.root.update()
        x, y = table.winfo_rootx() + 20, table.winfo_rooty() + table.winfo_height() - 3  # 最後一列下面的空白
        self.app.handle_dropped_files([last], x, y)
        self.assertEqual(self.app.photo_files[cat][-1], last)

    def test_drop_outside_the_table_appends_to_the_end(self):
        cat, table, (a, b, c) = self._photos()
        new = self._make_images(['n1.jpg'])
        self._show_photos_tab()
        x = self.app.photo_right_panel.winfo_rootx() + 30  # 放在右邊預覽區
        y = self.app.photo_right_panel.winfo_rooty() + 30
        self.app.handle_dropped_files(new, x, y)
        self.assertEqual(self.app.photo_files[cat], [a, b, c, new[0]])

    def test_drop_with_duplicates_inserts_only_new_ones_at_the_position(self):
        cat, table, (a, b, c) = self._photos()
        new = self._make_images(['n1.jpg'])
        self._show_photos_tab()
        x, y = self._drop_point(table, 0, upper=True)
        self.app.handle_dropped_files([b, new[0]], x, y)
        self.assertEqual(self.app.photo_files[cat], [new[0], a, b, c])

    # ---- 填寫框可上下調整高度
    def test_every_text_box_has_a_resize_grip_that_changes_its_height(self):
        self.assertEqual(set(self.app.text_grips), set(self.app.text_fields))
        key = 'road_status'
        widget, grip = self.app.text_fields[key], self.app.text_grips[key]
        self.app.notebook.select(self.app._tab_by_title(g.TAB_ROAD))
        self.root.update()
        before = int(widget.cget('height'))
        line = tkfont_linespace(self.root, widget)
        grip.event_generate('<Button-1>', x=5, y=2, rootx=100, rooty=100)
        grip.event_generate('<B1-Motion>', x=5, y=2, rootx=100, rooty=100 + 3 * line)
        self.root.update()
        self.assertEqual(int(widget.cget('height')), before + 3)
        grip.event_generate('<B1-Motion>', x=5, y=2, rootx=100, rooty=100 - 20 * line)
        self.root.update()
        self.assertEqual(int(widget.cget('height')), g.TEXT_MIN_LINES)  # 不會小於下限
        grip.event_generate('<B1-Motion>', x=5, y=2, rootx=100, rooty=100 + 500 * line)
        self.root.update()
        self.assertEqual(int(widget.cget('height')), g.TEXT_MAX_LINES)  # 不會大於上限

    def test_resizing_a_text_box_updates_the_tab_scroll_region(self):
        self.app.notebook.select(self.app._tab_by_title(g.TAB_ROAD))
        self.root.update()
        canvas, inner, scrollbar = self.app._scroll_tabs[g.TAB_ROAD]
        widget = self.app.text_fields['road_status']
        self.app.set_text_height(widget, g.TEXT_MAX_LINES)
        self.root.update()
        self.assertGreater(inner.winfo_reqheight(), canvas.winfo_height())
        self.assertTrue(scrollbar.winfo_manager())  # 內容變高，出現捲動軸
        self.app.set_text_height(widget, g.TEXT_MIN_LINES)
        self.root.update()
        self.assertEqual(int(widget.cget('height')), g.TEXT_MIN_LINES)

    def test_visitor_photo_categories_renamed_and_old_drafts_still_load(self):
        self.assertEqual(g.PHOTO_CATEGORIES[3:5], ['環境清理情形', '遊客宣導與違法情形'])
        self.assertEqual(g.PHOTO_SHORT_NAMES['環境清理情形'], '環境清理')
        self.assertEqual(g.PHOTO_SHORT_NAMES['遊客宣導與違法情形'], '遊客')
        tabs = [self.app.photo_notebook.tab(t, 'text') for t in self.app.photo_notebook.tabs()]
        self.assertTrue(any(x.startswith('遊客 ') for x in tabs))
        self.assertFalse(any('遊客違法' in x for x in tabs))
        [a, b] = self._make_images(['x.jpg', 'y.jpg'])
        draft = self.app.build_draft()
        draft['photos'] = {'環境清理及遊客宣導情形': [a], '遊客違法情形': [b]}
        draft['captions'] = {'環境清理及遊客宣導情形': {a: '清垃圾'}, '遊客違法情形': {b: '亂丟'}}
        draft['rotations'] = {'遊客違法情形': {b: 90}}
        self.assertEqual(self.app.apply_draft(draft), [])
        self.assertEqual(self.app.photo_files['環境清理情形'], [a])
        self.assertEqual(self.app.photo_files['遊客宣導與違法情形'], [b])
        self.assertEqual(self.app.photo_captions['遊客宣導與違法情形'], {b: '亂丟'})
        self.assertEqual(self.app.photo_rotations['遊客宣導與違法情形'], {b: 90})

    # ---- 照片裁切（只記錄範圍，不改原檔；預覽、報告、草稿都會套用）
    def test_normalize_crop_rejects_bad_tiny_and_full_ranges(self):
        self.assertIsNone(g.normalize_crop(None))
        self.assertIsNone(g.normalize_crop('abc'))
        self.assertIsNone(g.normalize_crop((0, 0, 1, 1)))                 # 整張＝不裁切
        self.assertIsNone(g.normalize_crop((0.2, 0.2, 0.21, 0.9)))        # 太窄
        self.assertEqual(g.normalize_crop((0.6, 0.9, 0.2, 0.1)), (0.2, 0.1, 0.6, 0.9))  # 左右上下顛倒會自動排好
        self.assertEqual(g.normalize_crop((-1, 0.1, 0.5, 2)), (0.0, 0.1, 0.5, 1.0))      # 超出範圍會夾在 0~1

    def test_crop_coordinates_convert_back_and_forth_with_rotation(self):
        crop = (0.1, 0.2, 0.5, 0.9)
        for angle in (0, 90, 180, 270):
            shown = g.crop_to_display(crop, angle)
            back = g.crop_from_display(shown, angle)
            for a, b in zip(back, crop):
                self.assertAlmostEqual(a, b, places=9, msg=f'{angle}°')
        # 順時針 90°：原圖左上角的框，轉過去之後在畫面右上角
        left, top, right, bottom = g.crop_to_display((0.0, 0.0, 0.5, 0.25), 90)
        self.assertAlmostEqual(right, 1.0)
        self.assertAlmostEqual(top, 0.0)
        self.assertAlmostEqual(left, 0.75)
        self.assertAlmostEqual(bottom, 0.5)

    def test_prepare_image_crops_before_rotating(self):
        path = os.path.join(self.tmp.name, 'c.png')
        im = Image.new('RGB', (400, 200), 'white')
        for x in range(0, 100):
            for y in range(0, 200):
                im.putpixel((x, y), (255, 0, 0))       # 左邊 1/4 是紅色
        im.save(path)
        src, size = g.prepare_image(path, 0, (0.0, 0.0, 0.25, 1.0))
        self.assertEqual(size, (100, 200))
        with Image.open(src) as out:
            r, g_, b_ = out.getpixel((50, 100))
            self.assertTrue(r > 240 and g_ < 15 and b_ < 15)         # 裁出來的全是紅色（JPEG 壓縮可能差 1~2）
        src, size = g.prepare_image(path, 90, (0.0, 0.0, 0.25, 1.0))
        self.assertEqual(size, (200, 100))                            # 裁完再旋轉，寬高互換
        self.assertEqual(g.prepare_image(path)[1], (400, 200))        # 沒給裁切就是原樣
        with Image.open(path) as original:
            self.assertEqual(original.size, (400, 200))               # 原檔沒被改動

    def test_crop_button_disabled_until_a_photo_is_selected(self):
        self.assertEqual(str(self.app.crop_btn.cget('state')), 'disabled')
        cat, table, (a,) = self._photos(names=('r.jpg',))
        self._select_first(cat, table)
        self.assertEqual(str(self.app.crop_btn.cget('state')), 'normal')

    def _crop_with(self, selection, cat, table):
        """用假的對話框流程：開啟裁切對話框，把選取範圍設成 selection（畫面比例），按確定。"""
        def fake_run(dialog):
            dialog.set_selection(*selection)
            dialog.ok()
            return dialog.result
        with mock.patch.object(g.CropDialog, 'run', fake_run):
            self.app.crop_btn.invoke()
        self.root.update()

    def test_cropping_updates_preview_table_and_keeps_original_file(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200))
        self._select_first(cat, table)
        self.assertEqual(table.row_text(0), '● (橫) r.jpg')
        self._crop_with((0.0, 0.0, 0.25, 1.0), cat, table)             # 只留左邊 1/4：100×200 → 直式
        self.assertEqual(self.app.photo_crops[cat][a], (0.0, 0.0, 0.25, 1.0))
        self.assertEqual(table.row_text(0), '● (直) ✂ r.jpg')
        preview = self.app.photo_preview
        self.assertGreater(preview.image.height(), preview.image.width())
        with Image.open(a) as original:
            self.assertEqual(original.size, (400, 200))
        self._crop_with((0.0, 0.0, 1.0, 1.0), cat, table)              # 選取整張＝還原
        self.assertNotIn(a, self.app.photo_crops[cat])
        self.assertEqual(table.row_text(0), '● (橫) r.jpg')

    def test_cropped_photos_are_marked_in_the_table_and_button_has_no_ellipsis(self):
        self.assertEqual(self.app.crop_btn.cget('text'), '✂ 裁切')
        cat, table, (a, b) = self._photos(names=('a.jpg', 'b.jpg'), size=(400, 200))
        self._select_first(cat, table)
        self._crop_with((0.0, 0.0, 0.25, 1.0), cat, table)
        self.assertEqual(table.row_text(0), '● (直) ✂ a.jpg')   # 裁切過的有 ✂
        self.assertEqual(table.row_text(1), '● (橫) b.jpg')      # 沒裁切的沒有
        self._crop_with((0.0, 0.0, 1.0, 1.0), cat, table)
        self.assertEqual(table.row_text(0), '● (橫) a.jpg')      # 還原後記號消失

    def test_cancelling_the_crop_dialog_changes_nothing(self):
        cat, table, (a,) = self._photos(names=('r.jpg',))
        self._select_first(cat, table)

        def fake_run(dialog):
            dialog.set_selection(0.1, 0.1, 0.5, 0.5)
            dialog.cancel()
            return dialog.result
        with mock.patch.object(g.CropDialog, 'run', fake_run):
            self.app.crop_btn.invoke()
        self.assertEqual(self.app.photo_crops[cat], {})

    def test_crop_dialog_converts_selection_with_the_current_rotation(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200))
        self.app.photo_rotations[cat][a] = 90
        self._select_first(cat, table)
        # 畫面（轉 90° 後是 200×400 直式）上選「右上角的 1/2 寬 × 1/4 高」，對應原圖「左上角 1/4 寬 × 1/2 高」
        self._crop_with((0.5, 0.0, 1.0, 0.25), cat, table)
        l, t, r, b = self.app.photo_crops[cat][a]
        self.assertAlmostEqual(l, 0.0)
        self.assertAlmostEqual(t, 0.0)
        self.assertAlmostEqual(r, 0.25)
        self.assertAlmostEqual(b, 0.5)
        # 之後把旋轉改回 0°，裁切範圍仍是同一塊原圖區域
        self.app.rotate_left_btn.invoke()
        self.assertEqual(self.app.photo_crops[cat][a], (l, t, r, b))

    def test_crop_dialog_reopens_with_the_saved_selection_and_mouse_drag_works(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200))
        self.app.photo_crops[cat][a] = (0.25, 0.25, 0.75, 0.75)
        dialog = g.CropDialog(self.root, a, 0, self.app.photo_crops[cat][a], self.app.palette, self.app.font_family)
        try:
            dialog.deiconify()  # 沒顯示出來的元件收不到滑鼠事件
            dialog.update()
            for got, want in zip(dialog.selection(), (0.25, 0.25, 0.75, 0.75)):
                self.assertAlmostEqual(got, want, places=2)
            self.assertEqual(dialog.kept_size(), (200, 100))
            pad = dialog.PAD
            # 在框外按下並拖曳，畫一個新的框（左上 10%,10% 到 40%,60%）
            x0, y0 = pad + 0.10 * dialog.view_w, pad + 0.10 * dialog.view_h
            x1, y1 = pad + 0.40 * dialog.view_w, pad + 0.60 * dialog.view_h
            dialog.canvas.event_generate('<ButtonPress-1>', x=int(x0), y=int(y0))
            dialog.canvas.event_generate('<B1-Motion>', x=int(x1), y=int(y1))
            dialog.canvas.event_generate('<ButtonRelease-1>', x=int(x1), y=int(y1))
            dialog.update()
            for got, want in zip(dialog.selection(), (0.10, 0.10, 0.40, 0.60)):
                self.assertAlmostEqual(got, want, delta=0.02)
            # 拖曳框內部可以移動，不會超出照片
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            dialog.canvas.event_generate('<ButtonPress-1>', x=int(cx), y=int(cy))
            dialog.canvas.event_generate('<B1-Motion>', x=int(cx + 5000), y=int(cy))
            dialog.canvas.event_generate('<ButtonRelease-1>', x=int(cx + 5000), y=int(cy))
            dialog.update()
            left, top, right, bottom = dialog.selection()
            self.assertAlmostEqual(right, 1.0, delta=0.01)
            self.assertAlmostEqual(right - left, 0.30, delta=0.02)
            # 只點一下（框太小）不會把框弄不見
            dialog.canvas.event_generate('<ButtonPress-1>', x=pad + 3, y=pad + 3)
            dialog.canvas.event_generate('<ButtonRelease-1>', x=pad + 3, y=pad + 3)
            dialog.update()
            self.assertAlmostEqual(dialog.selection()[2] - dialog.selection()[0], 0.30, delta=0.02)
            dialog.select_whole()
            dialog.ok()
            self.assertIsNone(dialog.result)   # 選整張＝不裁切
        finally:
            if dialog.winfo_exists():
                dialog.destroy()

    def test_crop_survives_draft_roundtrip_and_is_dropped_with_the_photo(self):
        cat, table, (a, b, c) = self._photos()
        self.app.photo_crops[cat][b] = (0.1, 0.2, 0.6, 0.9)
        draft = self.app.build_draft()
        self.assertEqual(draft['crops'][cat], {b: [0.1, 0.2, 0.6, 0.9]})
        self.assertTrue(self.app.is_dirty())
        import json
        draft = json.loads(json.dumps(draft))      # 真的經過 JSON 存檔再讀回來
        self.app.apply_draft(draft)
        self.assertEqual(self.app.photo_crops[cat], {b: (0.1, 0.2, 0.6, 0.9)})
        table.set_selection([1])
        self.app.remove_photos(cat, table)
        self.assertEqual(self.app.photo_crops[cat], {})
        self.app.photo_crops[cat][a] = (0.1, 0.1, 0.5, 0.5)
        self.app.clear_photos(cat, table)
        self.assertEqual(self.app.photo_crops[cat], {})

    def test_old_draft_without_crops_and_garbage_crops_still_load(self):
        cat, table, (a,) = self._photos(names=('o.jpg',))
        draft = self.app.build_draft()
        draft.pop('crops')
        self.assertEqual(self.app.apply_draft(draft), [])
        self.assertEqual(self.app.photo_crops[cat], {})
        draft['crops'] = {cat: {a: 'bad'}}
        self.assertEqual(self.app.apply_draft(draft), [])
        self.assertEqual(self.app.photo_crops[cat], {})

    def test_crop_is_applied_in_the_report(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200))
        out = os.path.join(self.tmp.name, 'crop.docx')
        self.fill_dates()
        data, error = self.app.collect_data()
        self.assertIsNone(error)
        files = {c: list(v) for c, v in self.app.photo_files.items()}
        g.build_report(TEMPLATE, data, files, {}, out, crops={cat: {a: (0.0, 0.0, 0.25, 1.0)}})
        shapes = Document(out).inline_shapes
        self.assertGreaterEqual(len(shapes), 1)
        self.assertGreater(shapes[0].height, shapes[0].width)  # 裁成 100×200 之後是直式

    # ---- 把照片拖到其他小分頁
    def _tab_point(self, index):
        """第 index 個小分頁標籤中央的螢幕座標。"""
        sub = self.app.photo_notebook
        left, right, top, bottom = g.tab_geometry(sub, index)
        return sub.winfo_rootx() + (left + right) // 2, sub.winfo_rooty() + (top + bottom) // 2

    def _press_row(self, src, src_table):
        """在選取的那一列按下滑鼠，開始拖曳。"""
        row = src_table.selected_indexes()[0]
        self.root.update()
        box = src_table.bbox(str(row))
        self.app._drag_start(src, src_table, self._Ev(box[1] + box[3] // 2))

    def _setup_two_categories(self):
        src, dst = g.PHOTO_CATEGORIES[0], g.PHOTO_CATEGORIES[2]
        _, src_table, (a, b, c) = self._photos(cat=src)
        [x, y] = self._make_images(['x.jpg', 'y.jpg'])
        self.app._add_photo_files(dst, [x, y])
        self._show_photos_tab()
        return src, dst, src_table, self.app.listboxes[dst], (a, b, c), (x, y)

    def test_dragging_over_another_sub_tab_switches_to_it_automatically(self):
        src, dst, src_table, dst_table, (a, b, c), (x, y) = self._setup_two_categories()
        src_table.set_selection([1])
        self._press_row(src, src_table)
        px, py = self._tab_point(2)
        self.app._drag_motion(src, src_table, self._Ev(5, px, py))
        self.root.update()
        self.assertEqual(self.app._current_photo_category(), dst)
        self.assertEqual(self.app.photo_files[src], [a, b, c])  # 還沒放開，沒有移動
        px, py = self._tab_point(0)
        self.app._drag_motion(src, src_table, self._Ev(5, px, py))   # 移回原本的分頁標籤
        self.root.update()
        self.assertEqual(self.app._current_photo_category(), src)
        self.app._drag_end(src, src_table, self._Ev(5, px, py))

    def test_drop_on_another_tab_table_inserts_at_the_release_position(self):
        src, dst, src_table, dst_table, (a, b, c), (x, y) = self._setup_two_categories()
        self.app.photo_captions[src][b] = '山屋'
        self.app.photo_rotations[src][b] = 90
        self.app.photo_crops[src][b] = (0.1, 0.1, 0.9, 0.9)
        src_table.set_selection([1])
        self._press_row(src, src_table)
        px, py = self._tab_point(2)
        self.app._drag_motion(src, src_table, self._Ev(5, px, py))
        self.root.update()
        box = dst_table.bbox('1')   # 目標表格第 2 列的上半部＝插在它前面
        rx, ry = dst_table.winfo_rootx() + 30, dst_table.winfo_rooty() + box[1] + box[3] // 4
        self.app._drag_motion(src, src_table, self._Ev(5, rx, ry))
        self.assertTrue(self.app._drag_marker.winfo_manager())          # 顯示「會插在這裡」的橫線
        self.app._drag_end(src, src_table, self._Ev(5, rx, ry))
        self.root.update()
        self.assertEqual(self.app.photo_files[src], [a, c])
        self.assertEqual(self.app.photo_files[dst], [x, b, y])
        self.assertEqual(self.app.photo_captions[dst], {b: '山屋'})    # 說明、旋轉、裁切跟著走
        self.assertEqual(self.app.photo_rotations[dst], {b: 90})
        self.assertEqual(self.app.photo_crops[dst], {b: (0.1, 0.1, 0.9, 0.9)})
        self.assertEqual(self.app.photo_captions[src], {})
        self.assertEqual(self.app.photo_rotations[src], {})
        self.assertEqual(dst_table.selected_indexes(), (1,))              # 選取剛移過來的照片
        self.assertEqual(self.app.photo_preview.current, (dst, b))
        self.assertFalse(self.app._drag_marker.winfo_manager())
        self.assertEqual(src_table.count(), 2)

    def test_drop_on_a_tab_label_appends_to_the_end_of_that_category(self):
        src, dst, src_table, dst_table, (a, b, c), (x, y) = self._setup_two_categories()
        src_table.set_selection([0])
        self._press_row(src, src_table)
        px, py = self._tab_point(2)
        self.app._drag_motion(src, src_table, self._Ev(5, px, py))
        self.app._drag_end(src, src_table, self._Ev(5, px, py))
        self.assertEqual(self.app.photo_files[dst], [x, y, a])
        self.assertEqual(self.app.photo_files[src], [b, c])

    def test_drop_in_empty_space_cancels_and_returns_to_the_original_tab(self):
        src, dst, src_table, dst_table, (a, b, c), (x, y) = self._setup_two_categories()
        src_table.set_selection([0])
        self._press_row(src, src_table)
        px, py = self._tab_point(2)
        self.app._drag_motion(src, src_table, self._Ev(5, px, py))
        self.root.update()
        nowhere = (self.root.winfo_rootx() + 2, self.root.winfo_rooty() + self.root.winfo_height() - 3)
        self.app._drag_motion(src, src_table, self._Ev(5, *nowhere))
        self.app._drag_end(src, src_table, self._Ev(5, *nowhere))
        self.root.update()
        self.assertEqual(self.app.photo_files[src], [a, b, c])
        self.assertEqual(self.app.photo_files[dst], [x, y])
        self.assertEqual(self.app._current_photo_category(), src)

    def test_dropping_a_photo_already_in_the_target_category_does_not_duplicate_it(self):
        src, dst, src_table, dst_table, (a, b, c), (x, y) = self._setup_two_categories()
        self.app.photo_files[dst].append(a)
        self.app._render_list(dst)
        self.assertFalse(self.app.transfer_photo(src, a, dst, 0))
        self.assertEqual(self.app.photo_files[src], [a, b, c])
        self.assertEqual(self.app.photo_files[dst], [x, y, a])
        self.assertIn('已經有這張照片', self.app.status_label.cget('text'))

    def test_transfer_marks_draft_dirty_and_updates_tab_marks(self):
        src, dst, src_table, dst_table, (a, b, c), (x, y) = self._setup_two_categories()
        self.app._saved_signature = self.app._signature()
        self.assertTrue(self.app.transfer_photo(src, a, dst, None))
        self.assertTrue(self.app.is_dirty())
        self.assertEqual(self.app.missing_caption_count(dst), 3)
        self.assertEqual(self.app.build_draft()['photos'][dst], [x, y, a])

    def test_font_is_iansui(self):
        self.assertEqual(g.FONT_FAMILY, 'Iansui')
        self.assertEqual(g.FONT_FILES, ('Iansui-Regular.ttf',))
        self.assertEqual(g.FONT_LICENSE, 'OFL-Iansui.txt')
        self.assertEqual(self.app.font_family, 'Iansui')
        old = [n for n in os.listdir(g.FONT_DIR) if any(k in n for k in ('LXGW', 'Cactus', 'Chocolate', 'Chiron'))]
        self.assertEqual(old, [])  # 舊字型已移除

    def test_font_line_height_is_compact_enough_for_photo_tabs(self):
        """字型的行高不能是一般字型的兩倍（之前試過的字型就是這樣），否則服勤照片分頁一定要捲動。"""
        import tkinter.font as tkfont
        font = tkfont.nametofont('TkDefaultFont', root=self.root)
        self.assertLessEqual(font.metrics('linespace'), 30)

    def test_only_tab_titles_and_page_headings_are_bold(self):
        """只有分頁標題與各頁標題用粗體；一般介面字型與欄位標籤維持一般字重。"""
        import tkinter.font as tkfont
        for name in ('TkDefaultFont', 'TkTextFont', 'TkFixedFont', 'TkMenuFont'):
            self.assertEqual(tkfont.nametofont(name, root=self.root).cget('weight'), 'normal')
        for style_name, weight in (('Heading.TLabel', 'bold'), ('TNotebook.Tab', 'bold'),
                                   ('Sub.TNotebook.Tab', 'normal'), ('Field.TLabel', 'normal')):
            font = tkfont.Font(root=self.root, font=style_lookup(self.root, style_name, 'font'))
            self.assertEqual(font.cget('weight'), weight, style_name)
        for w in self._all_widgets(self.root):
            if isinstance(w, ttk.Label) and str(w.cget('style')) == 'Field.TLabel':
                self.assertNotIn('bold', str(w.cget('font')))

    def test_numbered_section_titles_use_heading_style(self):
        """一～八大項的標題（藍色粗體）都用標題樣式；「1、」這類小項用欄位標籤樣式。"""
        found = {}
        for w in self._all_widgets(self.root):
            if isinstance(w, ttk.Label):
                text = str(w.cget('text'))
                if text[:2] in ('一、', '二、', '三、', '四、', '五、', '六、', '七、', '八、', '九、'):
                    found[text[0]] = str(w.cget('style'))
                elif text[:2] in ('1、', '2、', '3、', '4、'):
                    self.assertEqual(str(w.cget('style')), 'Field.TLabel', text)
        self.assertEqual(sorted(found), sorted('一二三四五六七八九'))
        self.assertEqual(set(found.values()), {'Heading.TLabel'}, found)

    def test_headings_are_distinguished_by_size_and_color(self):
        import tkinter.font as tkfont
        palette = g.THEMES[g.DEFAULT_THEME]
        heading = tkfont.Font(root=self.root, font=style_lookup(self.root, 'Heading.TLabel', 'font'))
        field = tkfont.Font(root=self.root, font=style_lookup(self.root, 'Field.TLabel', 'font'))
        body = tkfont.nametofont('TkDefaultFont', root=self.root)
        self.assertGreater(abs(heading.cget('size')), abs(field.cget('size')))
        self.assertGreater(abs(field.cget('size')), abs(body.cget('size')))
        self.assertEqual(style_lookup(self.root, 'Heading.TLabel', 'foreground'), palette['heading_fg'])
        self.assertEqual(style_lookup(self.root, 'Field.TLabel', 'foreground'), palette['field_fg'])

    def test_all_labels_use_the_shared_label_styles(self):
        styles = {str(w.cget('style')) for w in self._all_widgets(self.root)
                  if isinstance(w, ttk.Label) and w.cget('text').startswith(('1. 填表', '三、', '九、服勤照片', '二、設施'))}
        self.assertEqual(styles, {'Field.TLabel', 'Heading.TLabel'})

    # ---- 產生前檢查未填欄位
    def _fill_everything(self):
        self.fill_dates()
        self.app.personnel.insert(0, '王小明')
        self.app.location.insert(0, '排雲山莊')
        for w in self.app.text_fields.values():
            w.insert('1.0', '內容')
        for i, cat in enumerate(g.PHOTO_CATEGORIES):
            p = os.path.join(self.tmp.name, f'p{i}.jpg')
            Image.new('RGB', (40, 30)).save(p)
            self.app.photo_files[cat].append(p)
            self.app.photo_captions[cat][p] = '說明'
            self.app._render_list(cat)

    def test_blank_fields_listed_per_tab(self):
        self.fill_dates()
        data, _ = self.app.collect_data()
        blanks = self.app.find_blank_fields(data)
        self.assertEqual(list(blanks), ['基本資料', '道路/步道', '設施', '自然資源/環境維護/遊客管理',
                                        '遊客意見/心得建議/疑問解答', '服勤照片'])
        self.assertEqual(blanks['基本資料'], ['協勤人員姓名', '協勤地點'])
        self.assertEqual(len(blanks['道路/步道']), 3)
        self.assertEqual(len(blanks['設施']), 4)
        self.assertEqual(len(blanks['自然資源/環境維護/遊客管理']), 3)
        self.assertEqual(len(blanks['遊客意見/心得建議/疑問解答']), 3)
        self.assertEqual(len(blanks['服勤照片']), 6)  # 六個分類都還沒加入照片

    def test_partially_filled_form_lists_only_the_blank_ones(self):
        self._fill_everything()
        self.app.text_fields['road_grass'].delete('1.0', tk.END)
        self.app.location.delete(0, tk.END)
        self.app.clear_photos('其他', self.app.listboxes['其他'])
        data, _ = self.app.collect_data()
        blanks = self.app.find_blank_fields(data)
        self.assertEqual(blanks, {
            '基本資料': ['協勤地點'],
            '道路/步道': ['2、道路、步道雜草情形'],
            '服勤照片': ['其他（尚未加入照片）'],
        })

    def test_whitespace_only_counts_as_blank(self):
        self._fill_everything()
        self.app.text_fields['nature'].delete('1.0', tk.END)
        self.app.text_fields['nature'].insert('1.0', '   \n  ')
        data, _ = self.app.collect_data()
        self.assertIn('三、自然資源狀況', self.app.find_blank_fields(data)['自然資源/環境維護/遊客管理'])

    def test_generate_warns_about_blanks_with_default_yes_and_proceeds(self):
        self.fill_dates()
        self.app.personnel.insert(0, '王小明')
        out = os.path.join(self.tmp.name, 'o.docx')
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out), \
                mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertEqual(len(self.confirm_calls), 1)
        title, message, kwargs = self.confirm_calls[0]
        self.assertEqual(title, '有欄位尚未填寫')
        self.assertEqual(kwargs.get('default'), 'yes')   # 預設「是」
        self.assertEqual(kwargs.get('icon'), 'warning')
        self.assertIn('【基本資料】', message)
        self.assertIn('協勤地點', message)
        self.assertNotIn('協勤人員姓名', message)         # 已填的不會列出
        self.assertIn('【服勤照片】', message)
        self.assertIn('仍要產生報告嗎？', message)
        self.assertTrue(os.path.exists(out))             # 回答「是」就照常產生

    def test_answering_no_cancels_and_jumps_to_first_blank_tab(self):
        self.fill_dates()
        saved = []
        with mock.patch.object(messagebox, 'askyesno', lambda *a, **k: False), \
                mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: saved.append(1) or ''):
            self.app.generate_report()
        self.assertEqual(saved, [])  # 沒有進到存檔步驟
        self.assertEqual(self.app.notebook.select(), self.app._tab_by_title('基本資料'))
        self.assertEqual(str(self.app.generate_btn['state']), 'normal')
        self.assertEqual(self.pdf_calls, [])

    def test_no_warning_when_everything_is_filled(self):
        self._fill_everything()
        out = os.path.join(self.tmp.name, 'o.docx')
        with mock.patch.object(filedialog, 'asksaveasfilename', lambda **k: out), \
                mock.patch.object(g, 'open_file', lambda path: None):
            self.app.generate_report()
            self._wait_until_done()
        self.assertEqual(self.confirm_calls, [])
        self.assertTrue(os.path.exists(out))

    def test_date_errors_are_reported_before_blank_check(self):
        self.app.generate_report()  # 協勤日期還是空的
        self.assertEqual(self.shown[-1][0], 'err')
        self.assertEqual(self.confirm_calls, [])

    def test_primary_and_danger_buttons_keep_3d_bevel(self):
        """產生報告/清除按鈕要和其他按鈕一樣有凸起的立體邊框（亮邊 + 暗邊，且暗邊比底色深）。"""
        for style_name in ('Primary.TButton', 'Danger.TButton', 'TButton'):
            bg = style_lookup(self.root, style_name, 'background')
            light = style_lookup(self.root, style_name, 'lightcolor')
            dark = style_lookup(self.root, style_name, 'darkcolor')
            with self.subTest(style=style_name):
                self.assertNotEqual(light, bg)
                self.assertNotEqual(dark, bg)
                self.assertGreater(self._luminance(light), self._luminance(dark))
        for name in g.THEMES:
            self.app.apply_theme(name)
            for style_name, key in (('Primary.TButton', 'primary'), ('Danger.TButton', 'danger')):
                with self.subTest(theme=name, style=style_name):
                    fill = g.THEMES[name][key]
                    self.assertEqual(style_lookup(self.root, style_name, 'lightcolor'), g.BEVEL_LIGHT)
                    self.assertLess(self._luminance(style_lookup(self.root, style_name, 'darkcolor')),
                                    self._luminance(fill))

    def test_mix_color(self):
        self.assertEqual(g.mix_color('#ffffff', '#000000', 0.25), '#bfbfbf')
        self.assertEqual(g.mix_color('#102030', '#102030', 0.5), '#102030')
        self.assertEqual(g.mix_color('#000000', '#ffffff', 1), '#ffffff')

    def test_window_title(self):
        self.assertEqual(g.APP_TITLE, '玉山保育志工回報單生成器')
        self.assertEqual(self.root.title(), f'{g.APP_TITLE}－保育巡查回報單')          # 版本不放標題，放在右下角
        self.assertNotIn('v0', self.root.title())

    def test_version_is_a_small_gray_label_in_the_corner(self):
        label = self.app.version_label_widget
        self.assertEqual(label.cget('text'), g.version_label())
        self.assertRegex(label.cget('text'), r'^v\d+\.\d+(\.\d+)?-[0-9a-f]{7}$')   # 版本由 git tag 推算，不寫死
        self.assertEqual(str(label.cget('foreground')), g.VERSION_COLOR)
        self.assertEqual(label.winfo_class(), 'TLabel')
        self.assertEqual(label.master, self.app.theme_box.master)                       # 和「外觀」同一排的最右邊

    def test_build_info_overrides_git_when_packaged(self):
        import types
        fake = types.ModuleType('yscv_build_info')
        fake.LABEL = 'v1.0-abcdef0'
        with mock.patch.dict(sys.modules, {'yscv_build_info': fake}):
            self.assertEqual(g.version_label(), 'v1.0-abcdef0')
            self.assertNotEqual(g.version_label(self.tmp.name), 'v1.0-abcdef0')           # 指定資料夾時不用烙入的值

    def test_version_label_shows_tag_commits_since_tag_and_sha(self):
        with mock.patch.object(g, 'git_describe', return_value=('v0.5', 0, 'abc1234')):      # 剛好在 tag 上
            self.assertEqual(g.version_label(), 'v0.5-abc1234')
        with mock.patch.object(g, 'git_describe', return_value=('v0.5', 12, 'abc1234')):     # tag 之後第 12 個 commit
            self.assertEqual(g.version_label(), 'v0.5.12-abc1234')
        with mock.patch.object(g, 'git_describe', return_value=None), \
                mock.patch.object(g, 'git_revision', return_value='abc1234'):                # 還沒有 tag：用程式裡的版本
            self.assertEqual(g.version_label(), f'v{g.VERSION}-abc1234')
        with mock.patch.object(g, 'git_describe', return_value=None), \
                mock.patch.object(g, 'git_revision', return_value=''):                       # 不是 git 倉庫：只有版本
            self.assertEqual(g.version_label(), f'v{g.VERSION}')
        self.assertEqual(g.git_revision(self.tmp.name), '')            # 暫存資料夾不是 git 倉庫
        self.assertIsNone(g.git_describe(self.tmp.name))
        self.assertEqual(g.version_label(self.tmp.name), f'v{g.VERSION}')
        self.assertRegex(g.git_revision(), r'^[0-9a-f]{7,}$')            # 本專案是 git 倉庫
        self.assertRegex(g.version_label(), r'^v[\d.]+-[0-9a-f]{7}$|^v[\d.]+\.\d+-[0-9a-f]{7}$')

    def test_version_label_with_a_real_git_repository(self):
        import subprocess
        repo = os.path.join(self.tmp.name, 'repo')
        os.makedirs(repo)

        def git(*args):
            subprocess.run(['git', '-c', 'user.name=t', '-c', 'user.email=t@example.com', *args], cwd=repo,
                           check=True, capture_output=True)

        git('init', '-q')
        git('commit', '-q', '--allow-empty', '-m', 'one')
        self.assertEqual(g.version_label(repo), f'v{g.VERSION}-{g.git_revision(repo)}')        # 還沒有 tag
        git('tag', '-a', 'v1.2', '-m', 'v1.2')
        self.assertEqual(g.version_label(repo), f'v1.2-{g.git_revision(repo)}')                 # 剛好在 tag 上
        git('commit', '-q', '--allow-empty', '-m', 'two')
        git('commit', '-q', '--allow-empty', '-m', 'three')
        self.assertEqual(g.version_label(repo), f'v1.2.2-{g.git_revision(repo)}')               # tag 之後第 2 個 commit
        git('tag', '-a', 'v1.3', '-m', 'v1.3')
        self.assertEqual(g.version_label(repo), f'v1.3-{g.git_revision(repo)}')                 # 新 tag 之後歸零

    def test_date_entries_are_centered(self):
        for name in ('year', 'month', 'day', 'start_year', 'start_month', 'start_day',
                     'end_year', 'end_month', 'end_day'):
            with self.subTest(name=name):
                self.assertEqual(str(getattr(self.app, name).cget('justify')), 'center')

    def test_selected_tab_is_not_smaller_than_other_tabs(self):
        import re
        style = ttk.Style(self.root)
        numbers = lambda value: [int(n) for n in re.findall(r'\d+', str(value))]
        normal = numbers(style.lookup('TNotebook.Tab', 'padding', []))
        active = numbers(style.lookup('TNotebook.Tab', 'padding', ['selected']))
        self.assertEqual(len(normal), 2)
        self.assertEqual(len(active), 2)
        self.assertGreaterEqual(active[0], normal[0])
        self.assertGreaterEqual(active[1], normal[1])

    def test_selected_tab_is_taller_and_not_narrower_on_screen(self):
        nb = self.app.notebook
        self.root.update()
        count = len(nb.tabs())
        self.assertEqual(count, 6)
        for i in range(count):
            nb.select(i)
            self.root.update()
            selected = g.tab_geometry(nb, i)
            self.assertIsNotNone(selected, f'量不到分頁 {i}')
            others = [g.tab_geometry(nb, j) for j in range(count) if j != i]
            self.assertTrue(all(o is not None for o in others))
            # 被選取的分頁比其他分頁高（頂端更高）
            self.assertLess(selected[2], min(o[2] for o in others), f'分頁 {i} 被選取時沒有比其他分頁突出')
            # 選取時的寬度不應比沒選取時窄
            nb.select((i + 1) % count)
            self.root.update()
            idle = g.tab_geometry(nb, i)
            self.assertGreaterEqual(selected[1] - selected[0], idle[1] - idle[0], f'分頁 {i} 選取時比沒選取時窄')

    def test_notebook_tabs_are_visually_separated(self):
        style = ttk.Style(self.root)
        self.assertEqual(style.theme_use(), 'clam')
        active = style.lookup('TNotebook.Tab', 'background', ['selected'])
        idle = style.lookup('TNotebook.Tab', 'background', [])
        palette = g.THEMES[g.DEFAULT_THEME]
        self.assertEqual(active, palette['tab_active'])
        self.assertEqual(idle, palette['tab_idle'])
        self.assertNotEqual(active, idle)
        self.assertEqual(style.lookup('TNotebook.Tab', 'foreground', ['selected']), palette['tab_active_fg'])
        self.assertGreaterEqual(int(style.lookup('TNotebook.Tab', 'borderwidth')), 2)

    def test_dropdown_offers_solid_circle_instead_of_small_dot(self):
        w = self.app.text_fields['nature']
        menu = self._bullet_menu(w)
        labels = [menu.entrycget(i, 'label') for i in range(menu.index('end') + 1) if menu.type(i) == 'command']
        self.assertTrue(any(lbl.startswith('●') and '實心圓形' in lbl for lbl in labels))
        self.assertFalse(any('•' in lbl or '圓點' in lbl for lbl in labels))
        idx = [i for i in range(menu.index('end') + 1) if menu.type(i) == 'command'
               and menu.entrycget(i, 'label').startswith('●')][0]
        w.insert('1.0', 'a')
        w.tag_add(tk.SEL, '1.0', 'end-1c')
        menu.invoke(idx)
        self.assertEqual(w.get('1.0', 'end-1c'), '● a')

    def test_enter_keeps_the_chosen_symbol(self):
        w = self.app.text_fields['nature']
        w.insert('1.0', '■ 第一項')
        w.mark_set(tk.INSERT, 'end-1c')
        self.app._on_return(w)
        self.assertEqual(w.get('1.0', 'end-1c'), '■ 第一項\n■ ')

    def test_enter_continues_list_and_ends_on_empty_item(self):
        w = self.app.text_fields['nature']
        w.insert('1.0', '1. 第一項')
        w.mark_set(tk.INSERT, 'end-1c')
        self.assertEqual(self.app._on_return(w), 'break')
        self.assertEqual(w.get('1.0', 'end-1c'), '1. 第一項\n2. ')
        self.assertEqual(self.app._on_return(w), 'break')  # 空項目再按 Enter：移除編號、結束清單
        self.assertEqual(w.get('1.0', 'end-1c'), '1. 第一項\n')
        self.assertIsNone(self.app._on_return(w))  # 一般文字維持預設行為

    def test_list_markers_flow_into_report(self):
        self.fill_dates()
        w = self.app.text_fields['nature']
        w.insert('1.0', '山羌\n水鹿')
        w.tag_add(tk.SEL, '1.0', 'end-1c')
        self.app.apply_list(w, g.toggle_numbers)
        data, err = self.app.collect_data()
        self.assertIsNone(err)
        self.assertEqual(data['nature'], '1. 山羌\n2. 水鹿')

    def test_gui_fonts_are_enlarged_by_two(self):
        import tkinter.font as tkfont
        probe = tk.Tk()
        base = tkfont.nametofont('TkDefaultFont', root=probe).cget('size')
        probe.destroy()
        scaled = tkfont.nametofont('TkDefaultFont', root=self.root).cget('size')
        step = 2 if base > 0 else -2
        self.assertEqual(scaled, base + step)
        self.assertEqual(g.GUI_FONT_DELTA, 2)

    def test_gui_fonts_use_configured_weight(self):
        import tkinter.font as tkfont
        self.assertIn(g.GUI_FONT_WEIGHT, ('normal', 'bold'))
        for name in ('TkDefaultFont', 'TkTextFont', 'TkFixedFont', 'TkMenuFont'):
            with self.subTest(font=name):
                self.assertEqual(tkfont.nametofont(name, root=self.root).cget('weight'), g.GUI_FONT_WEIGHT)

    def test_scale_gui_fonts_can_keep_normal_weight(self):
        import tkinter.font as tkfont
        probe = tk.Tk()
        try:
            g.scale_gui_fonts(probe, 0, None, 'normal')
            self.assertEqual(tkfont.nametofont('TkDefaultFont', root=probe).cget('weight'), 'normal')
        finally:
            probe.destroy()

    def test_font_files_bundled_with_license(self):
        for name in g.FONT_FILES + (g.FONT_LICENSE,):
            self.assertTrue(os.path.exists(os.path.join(g.FONT_DIR, name)), name)

    def test_gui_uses_bundled_font_family(self):
        import tkinter.font as tkfont
        self.assertEqual(self.app.font_family, g.FONT_FAMILY)
        for name in ('TkDefaultFont', 'TkTextFont', 'TkFixedFont'):
            self.assertEqual(tkfont.nametofont(name, root=self.root).cget('family'), g.FONT_FAMILY)

    def test_font_family_falls_back_when_not_available(self):
        with mock.patch.object(g, 'load_private_fonts', lambda: False), \
                mock.patch.object(g.tkfont, 'families', lambda root=None: ['Arial']):
            self.assertEqual(g.pick_font_family(self.root), 'Arial')
        with mock.patch.object(g, 'load_private_fonts', lambda: False), \
                mock.patch.object(g.tkfont, 'families', lambda root=None: ['Arial', g.FALLBACK_FONT_FAMILY]):
            self.assertEqual(g.pick_font_family(self.root), g.FALLBACK_FONT_FAMILY)

    # ---- 照片預覽與拖曳排序
    class _Ev:
        def __init__(self, y, x_root=0, y_root=0):
            self.y = y
            self.x_root = x_root
            self.y_root = y_root

    def _photos(self, names=('a.jpg', 'b.jpg', 'c.jpg'), cat=None, size=(400, 200), exif=None):
        cat = cat or g.PHOTO_CATEGORIES[0]
        table = self.app.listboxes[cat]
        paths = []
        for n in names:
            p = os.path.join(self.tmp.name, n)
            im = Image.new('RGB', size, 'red')
            if exif:
                ex = im.getexif()
                ex[0x0112] = exif
                im.save(p, exif=ex)
            else:
                im.save(p)
            self.app.photo_files[cat].append(p)
            paths.append(p)
        self.app._render_list(cat)
        return cat, table, paths

    def _drag(self, cat, table, start, *targets):
        """模擬從 start 列按下，依序拖到 targets 列，最後放開。"""
        row = {'i': start}
        with mock.patch.object(table, 'identify_row', lambda y: str(row['i']) if row['i'] is not None else ''):
            self.app._drag_start(cat, table, self._Ev(5))
            for target in targets:
                row['i'] = target
                self.app._drag_motion(cat, table, self._Ev(5))
            self.app._drag_end()

    def test_drag_reorders_photos_and_listbox(self):
        cat, table, (a, b, c) = self._photos()
        table.set_selection([0])
        self._drag(cat, table, 0, 1, 2)
        self.assertEqual(self.app.photo_files[cat], [b, c, a])
        self.assertEqual([table.row_text(i) for i in range(3)], ['● (橫) b.jpg', '● (橫) c.jpg', '● (橫) a.jpg'])
        self.assertEqual(table.selected_indexes(), (2,))
        self._drag(cat, table, 2, 0)
        self.assertEqual(self.app.photo_files[cat], [a, b, c])

    def test_drag_order_persists_in_draft_and_marks_dirty(self):
        cat, table, (a, b, c) = self._photos()
        self.app._saved_signature = self.app._signature()
        self._drag(cat, table, 2, 0)
        self.assertEqual(self.app.build_draft()['photos'][cat], [c, a, b])
        self.assertTrue(self.app.is_dirty())

    def test_drag_in_blank_area_does_nothing(self):
        cat, table, (a, b, c) = self._photos()
        row = {'i': None}  # 按在表格最後一列下方的空白處：identify_row 回傳空字串
        with mock.patch.object(table, 'identify_row', lambda y: ''):
            self.app._drag_start(cat, table, self._Ev(500))
            self.assertIsNone(self.app._drag)
            self.assertIsNone(self.app._drag_motion(cat, table, self._Ev(500)))
        self.assertEqual(self.app.photo_files[cat], [a, b, c])

    def test_drag_target_is_clamped(self):
        cat, lb, (a, b, c) = self._photos()
        self._drag(cat, lb, 0, 99)
        self.assertEqual(self.app.photo_files[cat], [b, c, a])

    def test_select_shows_preview_with_info(self):
        cat, table, (a, b, c) = self._photos(size=(400, 200))
        self.app.photo_captions[cat][a] = '山屋外觀'
        self.app._render_list(cat)
        table.set_selection([0])
        nb = self.app.notebook
        nb.select(self.app._tab_by_title(g.TAB_PHOTOS))  # 沒顯示出來的元件收不到事件，要先切到照片分頁
        self.root.update()
        table.event_generate('<<TreeviewSelect>>')  # 確認事件有綁到預覽
        self.root.update()
        preview = self.app.previews[cat]
        self.assertEqual(preview.current, (cat, a))
        self.assertIsNotNone(preview.image)
        bw, bh = preview.box_size()
        self.assertLessEqual(preview.image.width(), bw)
        self.assertLessEqual(preview.image.height(), bh)
        self.assertFalse([w for w in preview.winfo_children() if isinstance(w, ttk.Label)])  # 預覽區不再顯示檔名與直式／橫式，空間全部給圖片
        self.assertEqual(self.app.caption_var.get(), '山屋外觀')  # 說明顯示在預覽下方的輸入框
        self.assertEqual(str(self.app.caption_entry.cget('state')), 'normal')

    def test_preview_applies_exif_rotation(self):
        cat, table, (a,) = self._photos(names=('r.jpg',), size=(400, 200), exif=6)
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        preview = self.app.previews[cat]
        self.assertGreater(preview.image.height(), preview.image.width())  # 轉正後為直式

    def test_preview_of_unreadable_file_does_not_crash(self):
        cat = g.PHOTO_CATEGORIES[0]
        table = self.app.listboxes[cat]
        bad = os.path.join(self.tmp.name, 'bad.jpg')
        with open(bad, 'wb') as f:
            f.write(b'not an image')
        self.app.photo_files[cat].append(bad)
        self.app._render_list(cat)
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        preview = self.app.previews[cat]
        self.assertIsNone(preview.image)
        self.assertIn('bad.jpg', preview.picture.cget('text'))  # 錯誤訊息直接顯示在預覽框裡

    def test_preview_cleared_when_photo_removed_or_cleared(self):
        cat, table, (a, b, c) = self._photos()
        table.set_selection([1])
        self.app._on_photo_select(cat, table)
        self.assertEqual(self.app.previews[cat].current, (cat, b))
        table.set_selection([0])
        self.app.remove_photos(cat, table)  # 移除 a：預覽的是 b，不受影響
        self.assertEqual(self.app.previews[cat].current, (cat, b))
        table.set_selection([0])
        self.app.remove_photos(cat, table)  # 移除 b：預覽與說明輸入框都要清掉
        self.assertIsNone(self.app.previews[cat].current)
        self.assertIsNone(self.app.previews[cat].image)
        self.assertEqual(str(self.app.caption_entry.cget('state')), 'disabled')
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        self.assertEqual(self.app.previews[cat].current, (cat, c))
        self.app.clear_photos(cat, table)
        self.assertIsNone(self.app.previews[cat].current)

    def test_enter_in_caption_box_saves_and_jumps_to_next_photo(self):
        cat, table, (a, b, c) = self._photos()
        table.set_selection([0])
        self.app._on_photo_select(cat, table)
        self.app.caption_var.set('  新說明  ')
        self.assertEqual(self.app._caption_return(), 'break')
        self.assertEqual(self.app.photo_captions[cat][a], '新說明')       # 前後空白去掉
        self.assertEqual(table.item('0', 'values')[1], '新說明')
        self.assertEqual(table.selected_indexes(), (1,))                  # 跳到下一張
        self.assertEqual(self.app.previews[cat].current, (cat, b))
        self.assertEqual(self.app.caption_var.get(), '')                  # b 還沒有說明
        self.app.caption_var.set('第二張')
        self.app._caption_return()
        self.app.caption_var.set('第三張')
        self.app._caption_return()                                        # 最後一張：存起來，停在原地
        self.assertEqual(self.app.photo_captions[cat][c], '第三張')
        self.assertEqual(table.selected_indexes(), (2,))

    def test_all_categories_share_one_preview_and_caption_box(self):
        previews = {id(self.app.previews[c]) for c in g.PHOTO_CATEGORIES}
        self.assertEqual(len(previews), 1)

    def test_mousewheel_scrolls_tab_from_child_widget(self):
        self.root.geometry('900x300')
        self.root.update()
        nb = next(w for w in self.root.winfo_children()[0].winfo_children() if isinstance(w, ttk.Notebook))
        nb.select(1)  # 道路/步道
        self.root.update()
        tab = nb.nametowidget(nb.tabs()[1])
        canvas = next(w for w in tab.winfo_children() if isinstance(w, tk.Canvas))
        inner = canvas.nametowidget(canvas.winfo_children()[0])
        label = next(w for w in inner.winfo_children() if isinstance(w, ttk.Label))
        self.assertEqual(canvas.yview()[0], 0.0)
        label.event_generate('<MouseWheel>', delta=-120)
        self.root.update()
        self.assertGreater(canvas.yview()[0], 0.0)


if __name__ == '__main__':
    unittest.main()
